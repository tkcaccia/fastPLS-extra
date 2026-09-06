#!/usr/bin/env python3
"""Revise publication documents from source-checked algorithms and saved results.

This builder never fits a model or executes a comparator.
"""
from pathlib import Path
import csv
import json
import math
import statistics
import re
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

EXTRA = Path(__file__).resolve().parents[1]
ROOT = EXTRA.parent / "fastPLS_fresh_github"
BASE = ROOT / "artifacts/CMPB_rewrite_20260903_cycle138"
OUT = EXTRA / "manuscript/revision_20260905"
OUT.mkdir(parents=True, exist_ok=True)
PUB = ROOT / "publication_results/0.99.39/current_release"
OPT = ROOT / "benchmark_results/optimization_20260904"
SOURCES = {}

def rows(path):
    SOURCES[str(path)] = {"bytes": path.stat().st_size}
    with path.open() as stream:
        return list(csv.DictReader(stream))

def replace(p, text):
    p.clear()
    p.add_run(text)

def after(p, text, style="Body Text"):
    element = OxmlElement("w:p")
    p._p.addnext(element)
    result = Paragraph(element, p._parent)
    result.style = style
    result.add_run(text)
    return result

def remove(p):
    p._p.getparent().remove(p._p)

def formula(p, text):
    p.clear()
    obj = OxmlElement("m:oMathPara")
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = text
    run.append(content)
    equation.append(run)
    obj.append(equation)
    p._p.append(obj)

def font(size, bold=False):
    suffix = " Bold" if bold else ""
    return ImageFont.truetype(f"/System/Library/Fonts/Supplemental/Arial{suffix}.ttf", size)

def comparison_figure(records, path, gpu=False):
    """Three quantitative panels; different workloads never share cells."""
    datasets = ["breast", "metref", "cifar100"] if not gpu else ["cifar100"]
    methods = (["fastPLS_cpu_rsvd", "IKPLS_numpy_alg1", "IKPLS_numpy_alg2"]
               if not gpu else ["fastPLS_cuda_rsvd", "IKPLS_jax_cuda_alg1", "IKPLS_jax_cuda_alg2"])
    labels = ["fastPLS SIMPLS / rSVD", "IKPLS score-explicit", "IKPLS cross-product"]
    fields = (["accuracy", "median_total_sec", "median_peak_rss_mb"] if not gpu else
              ["accuracy", "median_cold_total_sec", "median_warm_total_sec"])
    titles = (["Accuracy (%)", "Total time (s)", "Peak process RSS (MiB)"] if not gpu else
              ["Accuracy (%)", "Time vs IKPLS first call (s)", "Time vs IKPLS repeat call (s)"])
    image = Image.new("RGB", (1800, 750 if gpu else 1030), "white")
    draw = ImageDraw.Draw(image)
    draw.text((25, 20), "CIFAR-100 GPU workflow comparison" if gpu else
              "Matched CPU workflows including IKPLS", font=font(37, True), fill="black")
    draw.text((25, 70), "Float64; 50 components; 10,000 held-out predictions" if gpu else
              "Float64; Breast 10, MetRef 22, CIFAR-100 50 components; three runs",
              font=font(25), fill="black")
    lookup = {(r["dataset"], r["implementation"]): r for r in records}
    colors = [(0, 114, 178), (0, 158, 115), (213, 94, 0)]
    for panel, (field, title) in enumerate(zip(fields, titles)):
        left = 40 + panel * 590
        draw.text((left, 128), chr(65 + panel) + "  " + title, font=font(28, True), fill="black")
        for di, dataset in enumerate(datasets):
            top = 185 + di * 245
            draw.text((left, top), {"breast":"Breast", "metref":"MetRef", "cifar100":"CIFAR-100"}[dataset],
                      font=font(26, True), fill="black")
            vals = [float(lookup[dataset, method][field]) for method in methods]
            maximum = 100 if field == "accuracy" else max(vals) * 1.16
            for mi, value in enumerate(vals):
                if field == "accuracy": value *= 100
                y = top + 47 + mi * 51
                width = 400 * value / maximum
                draw.rectangle((left, y, left + width, y + 28), fill=colors[mi])
                label = f"{value:.2f}" if field == "accuracy" else f"{value:.4g}"
                draw.text((left + width + 8, y - 2), label, font=font(23), fill="black")
    legend_y = image.height - 105
    for i, label in enumerate(labels):
        x = 30 + i * 590
        draw.rectangle((x, legend_y, x + 25, legend_y + 25), fill=colors[i])
        draw.text((x + 35, legend_y - 2), label, font=font(25), fill="black")
    note = ("IKPLS first/repeat calls retained separately; new fastPLS rows do not distinguish JIT states."
            if gpu else "Stored CPU comparison; not a rerun of the final rSVD-only source. Workflows differ in estimator and outputs.")
    draw.text((25, image.height - 49), note, font=font(23), fill="black")
    image.save(path, dpi=(280, 280))

cpu = [r for r in rows(PUB / "ikpls_cross_language_cpu/ikpls_cross_language_summary.csv")
       if "irlba" not in r["implementation"]]
gpu = [r for r in rows(ROOT / "benchmark_results/ikpls_cross_language_20260825/ikpls_cross_language_cuda_summary.csv")
       if r["dataset"] == "cifar100" and r["implementation"].startswith("IKPLS")]
new_gpu = [rows(OPT / f"cuda/cifar_public_cuda_nocopy_rep{i}.csv")[0] for i in (1, 2, 3)]
gpu_time = statistics.median(float(r["total_sec"]) for r in new_gpu)
gpu.append(dict(dataset="cifar100", implementation="fastPLS_cuda_rsvd",
                accuracy=new_gpu[0]["accuracy"], median_cold_total_sec=gpu_time,
                median_warm_total_sec=gpu_time))
comparison_figure(cpu, OUT / "Figure1B_IKPLS_CPU.png")
comparison_figure(gpu, OUT / "Figure2B_IKPLS_CUDA.png", gpu=True)
nmr = rows(OPT / "float_factor_scores_nmr_summary/summary.csv")
cuda_nmr = {k: rows(OPT / f"remote/nmr/cuda_simpls_{k}.csv") for k in (50,165)}

main = Document(BASE / "fastPLS_CMPB_main_0.99.39.docx")
supp = Document(BASE / "fastPLS_CMPB_supplement_0.99.39.docx")
m, s = list(main.paragraphs), list(supp.paragraphs)
assert m[63].text == "3.1 Comparison with independent R implementations"
assert s[50].text == "S2.6 Bundled IRLBA"

updates = {
9: "Background and objective: High-dimensional biomedical partial least-squares (PLS) workflows can be limited by sequential component extraction and storage. We developed fastPLS to accelerate SIMPLS-based modelling, particularly multivariate NMR prediction, through compiled execution and compact prediction.",
10: "Methods: Native randomized singular value decomposition (rSVD) supplies approximate directions for SIMPLS, PLS-SVD, orthogonal PLS and kernel PLS. Fresh candidate blocks, cached products and incremental prediction are combined with CPU, NVIDIA CUDA and Apple Metal routes. Numerical-reference comparisons are distinguished from end-to-end comparisons with R packages and Improved Kernel Partial Least Squares (IKPLS). Precision, component counts and output contracts are reported separately.",
11: "Results: On the matched 50-component CIFAR-100 CUDA workload, the updated fastPLS implementation required a median 0.667 s and achieved 70.94% accuracy, compared with 1.859-1.861 s for the first IKPLS GPU call and 70.95% accuracy. For 28,355-response NMR prediction, repeated 50-component CUDA SIMPLS-rSVD calls had a median time of 0.707 s and RMSD 0.000749. CPU and Metal NMR experiments showed that float32 reduced some memory requirements but did not uniformly improve speed or predictive error. Comparisons with independent R implementations demonstrated dataset- and output-dependent trade-offs.",
12: "Conclusions: fastPLS extends practical SIMPLS-based modelling through compact representations and backend-specific computation. Approximate directions and batched execution require numerical validation; performance advantages depend on matrix dimensions, retained outputs and hardware.",
18: "SIMPLS constructs sequential components without explicitly deflating the predictor matrix [11], whereas PLS-SVD derives directions from one cross-covariance decomposition. OPLS removes predictor variation unrelated to the response [12], and kernel PLS extends modelling through nonlinear kernels [13,14]. Iterative truncated decomposition [15] and randomized SVD [16] reduce unnecessary dense decomposition. Relevant software includes the R package pls [17] and CPU/GPU improved-kernel PLS (IKPLS) [18]. We compare their computational workflows while distinguishing estimator, output-storage and execution differences.",
22: "The public interfaces select the PLS family, component counts, backend and classification head; native rSVD is the only public SVD method. Single and nested cross-validation retain the same fold, tuning and prediction contracts. Unavailable accelerators produce an error rather than a CPU fallback. The numerical core is shared with the R adapters, while GPL-dependent iterative-SVD comparison tools are separated from the main package. Detailed route capabilities are given in Supplementary Table S1.",
23: "Let X have n rows and p predictors, and Y have q response columns. The symbol a indexes a component, A is the maximum retained count, C is the set of requested component counts, J is the number of classes, and K is the number of cross-validation folds. Lowercase k denotes retrieval cutoffs.",
24: "CPU operations use compiled C/C++ and BLAS/LAPACK. CUDA uses cuBLAS matrix products and cuSOLVER factorizations; Metal uses Apple GPU products with route-specific host stages. CPU scaling is evaluated with OpenBLAS thread settings verified at runtime. Timing boundaries distinguish first-call accelerator costs from repeat calls and include synchronization before stopping the timer. Unavailable backend requests produce an error.",
29: "PLS-SVD obtains approximate leading factors S ≈ U D Vᵀ once and reuses their prefixes. For a requested prefix a, Tₐ = XUₐ and the compact response map Hₐ solves (TₐᵀTₐ)Hₐ = DₐVₐᵀ. Prediction uses XnewUₐHₐ, with the training response mean restored. Thus singular directions are not themselves regression coefficients. The available rank is bounded by S and by J−1 for a centred J-class response.",
30: "The accelerated SIMPLS-family path retains score normalization, predictor-loading orthogonalization and cross-covariance deflation, but its direction calculation is approximate. At a refresh it computes one or several newly randomized candidates from the current operator. Eligible large CPU and CUDA classification fits consume up to eight candidates sequentially; the CPU path batches candidate scores and loadings as matrix products. Other routes refresh component-wise, including the massive-matrix rank-one profile. Reusing a candidate within a block is not equivalent to solving afresh after every deflation, and exact de Jong estimator preservation is not asserted for this approximate batched path.",
34: "Algorithm 1. Randomized SIMPLS-family execution. A fresh candidate block is generated at each refresh and consumed sequentially. With block width one this reduces to a component-wise refresh. The updates define the implemented procedure, not an assertion of equality with an exact dominant-direction SIMPLS estimator.",
37: "Native rSVD forms a Gaussian range sketch, applies alternating operator products and orthonormalization, and decomposes a reduced matrix. Ordinary PLS-SVD controls are 32 oversampling directions and five power iterations. SIMPLS-family controls also account for response dimension and class density; the massive profile records 12/2 while using a fresh rank-one direction update. Oversampling metadata must therefore be read together with executed direction width. Failed checked CPU sketches can trigger wider, freshly initialized retries; unsuccessful recovery raises an error. Iterative-SVD references are comparison implementations, not a solver option in the current public package.",
45: "Regression produces continuous response estimates; classification decodes the multivariate response by argmax or fits LDA to PLS scores. LDA forms pooled covariance from score moments and solves (Σ + ρsI)wⱼ = μⱼ without inversion, where s is the positive finite mean diagonal or 1 otherwise. The sequence ρ = 10⁻⁸, 10⁻⁶, 10⁻⁵, 10⁻⁴, 10⁻³, 10⁻² starts at its first value and advances after an unsuccessful factorization or non-finite solve. Accuracy or balanced accuracy guides classification tuning; RMSD or fold-training-mean Q² guides regression. Full-data fitted R² is descriptive, not a tuning endpoint.",
54: "End-to-end comparisons include R PLS packages and IKPLS 6.1.2 [18]. IKPLS has a score-explicit formulation and a cross-product formulation based on predictor and predictor-response products. It is not a SIMPLS estimator reference. Within each panel, precision, preprocessing, split, component count and retained prediction output were matched. Minimum-output and ordinary-public-object R comparisons remain distinct. IKPLS CPU measurements use Breast, MetRef and CIFAR-100 at 10, 22 and 50 components; they must not be pooled with the different component grids in the broader R panel. Stored comparator measurements were reused without rerunning external packages, and software-build provenance accompanies the result tables.",
57: "2.7 Large-scale case-study protocols",
61: "We first compare complete PLS software workflows, including R implementations and IKPLS, and then examine accelerator execution. NMR provides the principal high-response biomedical case study. ImageNet embeddings serve as a separate computational stress test of foundation-model representations. Numerical-reference evidence, approximate-direction agreement and predictive performance are interpreted separately.",
63: "3.1 Comparison with independent implementation",
64: "The broad R-package panel and matched IKPLS panel answer complementary questions (Figure 1). The R panel compares ordinary PLS-DA workflows, including argmax and latent-score LDA, at dataset-specific component counts. Its fixed-control IRLBA rows describe the comparison build, not a solver available in the current rSVD-only package. The matched IKPLS panel instead fixes Breast, MetRef and CIFAR-100 at 10, 22 and 50 components. These panels cannot establish a universal ranking across all datasets or output policies.",
70: "In the stored float64 CPU comparison, IKPLS required 0.000562-0.000674 s on Breast, 0.00267-0.00271 s on MetRef and 0.165-0.720 s on CIFAR-100. The corresponding fastPLS rSVD comparison-build times were 0.005, 0.072 and 4.176 s. Accuracy was identical on Breast (94.29%); IKPLS and fastPLS obtained 75.0% and 77.0% on MetRef and 70.95% and 70.83% on CIFAR-100. These stored CPU measurements precede completion of the shared-core migration and are not evidence for the speed of the final rSVD-only source. They show why IKPLS is a relevant comparator rather than implying universal fastPLS superiority.",
72: "3.2 GPU performance and internal acceleration",
73: "The updated matched CIFAR-100 CUDA SIMPLS-rSVD workload took 0.666, 0.667 and 0.668 s at 50 components, with accuracy 70.94%. Stored IKPLS GPU first-call medians were 1.859 and 1.861 s for its score-explicit and cross-product formulations; their repeat-call medians were 1.480 and 1.552 s, and accuracy was 70.95%. The corresponding timing ratios favor fastPLS by approximately 2.8 for first calls and 2.2-2.3 for repeat calls (Figure 2B). The 0.01-percentage-point accuracy difference is not interpreted as a predictive improvement. These are different PLS implementations, and first-call/JIT costs are distinguished from repeated-call costs.",
74: "Within fastPLS, accelerator benefit depends on the numerical workload rather than sample count alone. Matrix products, component count, precision and data movement determine whether CUDA or Metal improves runtime. CPU/accelerator ratios and paired predictive differences are therefore presented separately. GPU IKPLS measurements in the matched panel do not establish its behavior on NMR or million-row ImageNet; the available large-case IKPLS extension used CPU float32.",
78: "3.3 Multivariate NMR prediction",
80: "The deposited 165-component PLS-SVD workflow remains a distinct implementation comparator. New three-seed NMR measurements show that SIMPLS-rSVD float64 required median total times of 11.181 s on CPU and 4.799 s on Metal at 165 components, with RMSD 0.000954 and 0.000937. The updated CUDA repeat-call experiment required 1.797 s and gave RMSD 0.000940 at seed 123. These campaigns have different timing replication structures and are reported separately, rather than pooled into a nominally identical three-run benchmark. The original common-165-component comparison and spectra remain in Figure 3; updated precision/seed measurements are tabulated in Supplementary Section S14.",
81: "At 50 components, the new three-seed SIMPLS-rSVD experiment gave median RMSD 0.000751 in 2.976 s on CPU and 0.000750 in 1.908 s on Metal. The CUDA repeat-call median was 0.707 s with RMSD 0.000749. Float32 CPU SIMPLS used a smaller baseline-corrected peak RSS increment (256 versus 455 MiB) but took 4.747 s rather than 2.976 s. In contrast, float32 Metal PLS-SVD at 50 components took 0.679 s versus 6.657 s for float64, with nearly identical RMSD around 0.000740. Precision therefore changes storage and execution costs in route-specific ways; it is not a guarantee of faster fitting.",
86: "3.4 Foundation-model embedding feasibility",
94: "The main implementation advances are compiled sequential updates, reuse of deflation products, batched candidate geometry, compact latent prediction and backend-specific rSVD. They reduce different costs: matrix batching improves arithmetic intensity, latent factors avoid large coefficient paths, and implicit products avoid explicit cross-covariance storage. Approximate candidate blocks retain supervised deflation but do not in general reproduce independent exact leading-direction solves. Their predictive and subspace behavior must therefore be evaluated directly.",
95: "The IKPLS comparison illustrates that optimized cross-product algorithms can be substantially faster on CPU for some score dimensions and output contracts, whereas the updated CUDA fastPLS workflow was faster in the matched CIFAR-100 panel. Neither observation establishes a universal software ranking. Multithreaded BLAS can accelerate covariance and prediction products, although sequential recurrences and small matrices limit scaling. The separate LDA experiment improved total time at 500 scores by approximately twofold with LAPACK and by a further twofold with four OpenBLAS threads; this does not establish an equivalent whole-PLS speedup.",
96: "Native rSVD is the public solver, and its accuracy depends on spectrum, rank, retained components and seed. Independent dense or iterative reference analyses remain useful for sensitive scientific conclusions. Float32 reduces representation size but needs route-specific metric and memory validation, particularly in ill-conditioned problems. The newer LDA LAPACK float32 candidate remains experimental because a singular-covariance test showed weight disagreement despite small residuals; candidate results are not described as production behavior.",
98: "fastPLS extends SIMPLS-based biomedical modelling through compact prediction, reusable numerical state, approximate candidate extraction and CPU/GPU execution. The NMR experiments demonstrate practical high-response modelling, while independent R and IKPLS comparisons identify workload-dependent advantages and limitations. The current public package uses native rSVD; numerical-reference comparisons and experimental alternatives remain separate from its supported production paths.",
110: "The fastPLS development source is available at https://github.com/tkcaccia/fastPLS, with Bioconductor submission at https://github.com/Bioconductor/BiocContributions/issues/84. Shared native headers implement reusable CPU numerical stages. The companion project fastPLS-extra separates GPL-dependent comparison code, manuscript sources and benchmark evidence from the main package. Source ownership and third-party distribution permissions are being audited before an MIT package release; no completed package-wide relicensing is asserted here. Per-analysis source and result files are identified in the supplementary provenance material.",
}
for index, text in updates.items(): replace(m[index], text)
for index in (55,56,69): remove(m[index])
replace(m[68], "Figure 1A. Independent R-package classification workflows at dataset-specific component counts. Panels retain the reported accuracy, total fitting-plus-prediction time and complete-process RSS. Fixed-control fastPLS/IRLBA rows belong to the comparison build rather than the current public solver. Different output contracts and prediction heads are workflow comparisons. Figure 1B adds the separate matched IKPLS panel; its different component counts must not be pooled with panel A.")
p = after(m[68], "", "Normal")
p.add_run().add_picture(str(OUT / "Figure1B_IKPLS_CPU.png"), width=Inches(6.25))
after(p, "Figure 1B. Matched float64 CPU software comparison including IKPLS 6.1.2. Breast, MetRef and CIFAR-100 use 10, 22 and 50 components, respectively. Bars show recorded medians from three runs; exact repetitions and dispersion are retained in the supplementary tables. These stored CPU results have not been relabelled as a rerun of the final rSVD-only source.", "Caption")
p = after(m[77], "", "Normal")
p.add_run().add_picture(str(OUT / "Figure2B_IKPLS_CUDA.png"), width=Inches(6.25))
after(p, "Figure 2B. Updated fastPLS CUDA SIMPLS-rSVD versus stored IKPLS JAX-CUDA results on the matched 50-component float64 CIFAR-100 workload. fastPLS uses oversampling 32, power 5 and seed 123. IKPLS first-call and repeat-call medians are distinct; the same fastPLS repeat-call median is shown beside each and is not a measurement of JAX compilation. Timing includes fitting and held-out prediction. Both implementations have approximately 70.95% accuracy.", "Caption")

algorithm = [
"Input: centred X and Y, maximum component count A and requested prefixes C.",
"1. Form S₀ = XᵀY or its implicit products; allocate compact factors R, Q and orthogonal basis V.",
"2. At each refresh choose block width b (one, or up to eight on eligible CPU/CUDA classification routes) and draw new randomized candidates D from the current deflated operator.",
"3. Where supported, compute candidate geometry XD and Xᵀ(XD) together. Consume the candidate columns sequentially until the block is exhausted.",
"4. For candidate rₐ, set tₐ = Xrₐ; divide both by ‖tₐ‖₂. Compute pₐ = Xᵀtₐ and qₐ = Yᵀtₐ, using cached products when available.",
"5. Orthogonalize vₐ = pₐ − VVᵀpₐ, repeat the projection where configured, and normalize vₐ.",
"6. Compute hₐ = vₐᵀSₐ₋₁ once; update Sₐ = Sₐ₋₁ − vₐhₐ, or the corresponding implicit state.",
"7. Append rₐ, qₐ and vₐ. If requested, update Bₐ = Bₐ₋₁ + rₐqₐᵀ and Ŷₐ = Ŷₐ₋₁ + tₐqₐᵀ; retain only requested prefixes.",
"8. Refresh after the block; stop at A or numerical breakdown. Predict through XnewRₐQₐᵀ and restore the training mean.",
"Candidate reuse is an approximate algorithmic choice, not an exact dominant-direction equivalence claim."]
main.tables[0].cell(0,0).text = "\n".join(algorithm)

sup_updates = {
10: "The public model families are PLS-SVD, SIMPLS, OPLS and kernel PLS, with argmax and LDA classification. Native rSVD is the only public SVD solver. IRLBA comparison integration is separated into the GPL companion, which is not a dependency of fastPLS. Stored IRLBA evidence concerns reference implementations rather than a current public option.",
13: "Notation: X has n samples and p predictors; Y has q responses. Component index a, maximum A and requested-prefix set C are distinct from J classes and K validation folds. Lowercase k is reserved for retrieval cutoffs.",
17: "PLS-SVD requests the leading valid rank once. Let S ≈ UDVᵀ and Tₐ = XUₐ. The cached-score-Gram path solves (TₐᵀTₐ)Hₐ = DₐVₐᵀ and predicts XnewUₐHₐ, restoring the training response mean. The noncached path computes the latent response cross-product explicitly. Approximate singular factors can make these two right-hand sides differ numerically; they are not assumed exactly interchangeable. Rank is bounded by the cross-covariance and by J−1 for a centred J-class response.",
25: updates[30], 27: "At a refresh, the current operator generates fresh candidates. The CPU classification batch rule requires more than one and at most 2,048 response columns, at least four requested components and n p q ≥ 5×10⁸; width is bounded by eight, remaining components and matrix dimensions. CUDA adds at least 5,000 samples and 50 requested components. Actual low-level route availability and diagnostics determine the executed rule. Regression and ordinary component-wise routes do not reuse a multi-column block.",
28: "The deflation product hₐ = vₐᵀSₐ₋₁ is evaluated once and reused in Sₐ = Sₐ₋₁ − vₐhₐ.",
33: "For a candidate block D, the CPU path can compute XD and Xᵀ(XD) as matrix products before sequential acceptance. Each candidate's score and projection are normalized consistently. Loadings and response products can reuse the cached initial cross-covariance. This batches geometry, not independent model fits.",
36: "At each refresh compute b freshly initialized randomized candidates from the current deflated state, where b is one except for an eligible batched route. Consume the candidates in order and refresh after the block. A within-block candidate is not recomputed as the exact dominant direction of the newly deflated state.",
39: "Evaluate hₐ = vₐᵀSₐ₋₁ once and update Sₐ = Sₐ₋₁ − vₐhₐ.",
44: "OPLS first centres/scales X using training statistics and centres Y. For each orthogonal filter step obtain a predictive direction w, set t = Xw and p = Xᵀt/(tᵀt), then normalize wₒ = p − w(wᵀp)/(wᵀw). Compute tₒ = Xwₒ, pₒ = Xᵀtₒ/(tₒᵀtₒ), and replace X by X − tₒpₒᵀ. Retain wₒ and pₒ for held-out filtering and fit SIMPLS to the filtered predictors. The R CPU float64 filter currently retains its dense leading-direction callback; the native standalone wrapper uses rSVD. Therefore rSVD-only public dispatch does not mean every internal small decomposition is randomized. Accelerator filtering remains route dependent.",
48: "rSVD draws a Gaussian test matrix Ω, forms Z = MΩ, alternates products with Mᵀ and M with orthonormalization, and decomposes the reduced operator QᵀM. The ordinary controls are 32 oversampling directions and five power iterations; documented response/class profiles may increase them. The massive SIMPLS profile records 12/2 with a separately reported rank-one executed refresh. Failed checked CPU sketches can trigger four additional fresh, wider attempts with more iterations, subject to a conservative 256-MiB major-workspace budget; rejected recovery throws. This is not a bound on process RSS or proof of statistical equivalence.",
50: "S2.6 Independent iterative-SVD comparison",
51: "Augmented implicitly restarted Lanczos bidiagonalization [15] is retained in the separate GPL comparison project, not in the current fastPLS public solver list. Existing iterative-reference measurements are labelled by their generating build and are not reinterpreted as current native-rSVD measurements. Exact dense decompositions remain internal tools for suitable small or reduced problems.",
58: "Explicit and implicit rSVD routes use the same centred/scaled cross-covariance operator. Their activation and workspace rules are computational choices; they do not define different statistical preprocessing.",
67: "Set s = tr(Σ)/d when finite and positive, otherwise s = 1, where d is the score dimension. The implementation starts with ρ = 10⁻⁸ and tries 10⁻⁶, 10⁻⁵, 10⁻⁴, 10⁻³ and 10⁻² in that order after an unsuccessful factorization or non-finite triangular solve. It solves (Σ + ρsI)wⱼ = μⱼ by Cholesky and triangular solves; it does not first attempt an unregularized covariance or explicitly form an inverse. The prediction discriminant is given below.",
70: "CPU float64 LDA uses LAPACK Cholesky; non-Windows CPU float32 currently uses a reusable float32 workspace and triangular loops. CUDA uses cuBLAS products and cuSOLVER Dpotrf/Dpotrs or Spotrf/Spotrs. The float32 CUDA workspace persists between calls, while host label validation and model transfer remain explicit. Metal combines accelerated projection with host LDA. The experimental CPU float32 LAPACK alternative is evaluated separately and has not replaced all public float32 routes.",
134: "The current approximate solver must be characterized by both controls and refresh width. Fresh per-component sketches, massive rank-one refreshes and eligible CPU/CUDA candidate blocks are distinct numerical routes. Stored three-seed qualification counts concern the source/configuration that produced them; they do not automatically qualify a later batched or shared-core build. The current extraction tests establish preservation of tested values during refactoring, not a new all-dataset approximation qualification.",
139: "Ordinary R numeric data store float64 values; float::float32 stores four-byte values and selects supported float32 execution without a user precision switch. Representation size, process peak RSS and numerical accuracy are separate properties. CPU, CUDA and Metal support route-specific float32 operations, with host-assisted filtering, kernel and LDA stages identified explicitly. Windows uses a separate portable CPU route. Unavailable backends must error; no silent CPU substitution is allowed. The new NMR precision measurements below show both gains and regressions, so availability is not presented as a universal speed or memory advantage.",
150: "S12.1 Matched IKPLS comparison within the independent-software panel",
151: "The matched IKPLS comparison uses two Improved Kernel PLS formulations: score-explicit and cross-product. The CPU panel has float64 inputs, externally applied training centring, fixed one-hot responses, the same split and component counts, final held-out predictions and three repetitions. Breast, MetRef and CIFAR-100 use 10, 22 and 50 components. These differ from the broad R-package panel's component grid and cannot be pooled. Figure 1B reports the stored CPU measurements without attributing them to completion of the final shared-core migration. In the GPU extension, updated fastPLS CUDA CIFAR-100 measurements are compared with stored IKPLS JAX-CUDA first- and repeat-call measurements (Figure 2B). GPU IKPLS was not evaluated on the large NMR/ImageNet extension, which used CPU float32.",
153: "All IKPLS values are read from saved IKPLS 6.1.2 measurements; no external implementation was rerun during this revision. The accompanying machine-readable ledger identifies source tables, timing contracts and which fastPLS measurements were updated.",
}
for index,text in sup_updates.items(): replace(s[index],text)
# The number of classes is J; C denotes the requested component-prefix set.
for node in main._element.xpath('.//m:t'):
    if node.text == "C": node.text = "J"
# Correct class-count notation in retained equation objects, without disturbing layout.
for index in (55,56,66):
    for node in s[index]._p.xpath('.//m:t'):
        if node.text == "C": node.text = "J"
formula(s[69], "δⱼ(t) = tᵀwⱼ − ½ μⱼᵀwⱼ + log(nⱼ/n)")
p = after(s[46], "For the nonlinear kernels, Kᵢⱼ = exp(−γ‖xᵢ−xⱼ‖²) or (γxᵢᵀxⱼ + c)ᵈ. Training centring is Kc = HKH with H = I − 11ᵀ/n. A held-out kernel row subtracts its own mean and the stored training column means, then adds the training grand mean. Test data never determine training preprocessing or kernel means.")

def add_table_after(document, anchor, headers, records):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = document.tables[0].style
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for c,h in zip(table.rows[0].cells,headers): c.text=h
    for record in records:
        for c,value in zip(table.add_row().cells,record): c.text=str(value)
    anchor._p.addnext(table._tbl)
    return table

p = after(s[70], "S3.3 CPU LDA solver alternatives", "Heading 2")
p = after(p, "A separate synthetic score-stage experiment compared workspace and BLAS/LAPACK Cholesky in float32 and float64 with OpenBLAS 0.3.33 on the Apple M3. The same random double sequence generated both precisions; conversion preceded timing. Training sizes 1,000/5,000, score widths 10/50/200/500, ten classes and 400 held-out rows were tested at runtime-verified OpenBLAS settings 1/2/4. Each case had one warm-up and 15 timed fits per solver, giving 1,440 measured fits. Moments, solve and prediction were timed separately; PLS fitting and process startup were excluded.")
p = after(p, "At 5,000 training rows and 500 scores, LAPACK reduced float32 total LDA time from 37.44 to 17.20 ms on one thread and to 8.17 ms on four threads. Float64 times were 55.64, 33.41 and 14.54 ms. All compared labels agreed and accuracy was 82.75% in this synthetic workload. Maximum relative score differences were 7.73×10⁻⁷ and 1.44×10⁻¹⁵. A rank-deficient float32 covariance nevertheless gave approximately 1.96% weight disagreement with small backward residuals and equal regularization; the candidate is therefore not claimed universally interchangeable with the production solver.")
# Insert the new table before the existing NMR table and update later references.
for doc in (main, supp):
    for node in doc._element.xpath('.//w:t'):
        if node.text:
            node.text = re.sub(r'\bS(25|26|27|28)\b',
                               lambda m: 'S' + str(int(m.group(1)) + 1), node.text)
p = after(s[172], "Table S25. NMR precision and seed experiment. Values below use the same water-masked-predictor protocol and all 28,355 response columns, with seeds 123-125, three timing replicates per seed and three separate process-memory runs. They are distinct from the single-seed common-165-component illustration in Figure 3. Baseline-corrected RSS is process memory, not isolated allocation.")
table_rows=[]
for r in nmr:
    table_rows.append([r['family'].replace('plssvd','PLS-SVD').replace('simpls','SIMPLS'),r['backend'],r['precision'],r['ncomp'],f"{float(r['median_total_sec']):.3f}",f"{float(r['median_RMSD']):.6f}",f"{float(r['median_incremental_rss_mib']):.0f}"])
add_table_after(supp,p,["Family","Backend","Precision","A","Time (s)","RMSD","RSS Δ (MiB)"],table_rows)

# Preserve source evidence and make the incomplete release-level comparison explicit outside the paper.
ledger = {"sources":SOURCES,"unresolved":[
"Complete CPU/IKPLS and broad R comparisons have not been rerun with the final rSVD-only source.",
"Figure 1A and original Figure 3 retain their identified generating comparison builds.",
"Final package-wide MIT licensing and frozen-release benchmark qualification remain incomplete."]}
(OUT/'evidence_ledger.json').write_text(json.dumps(ledger,indent=2)+'\n')
for doc in (main,supp):
    # Keep portrait scientific figures and their captions together without
    # allowing a tall image to strand its caption on a nearly empty page.
    for shape in doc.inline_shapes:
        limit = Inches(7.1)
        if shape.height > limit:
            shape.width = int(shape.width * limit / shape.height)
            shape.height = limit
    for i, p in enumerate(doc.paragraphs):
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1
        if p.text.startswith('Figure '):
            following = doc.paragraphs[i+1] if i+1 < len(doc.paragraphs) else None
            if following is not None and following._p.xpath('.//w:drawing'):
                p.paragraph_format.keep_with_next = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size=Pt(8)
    for section in doc.sections:
        if section.page_width > section.page_height:
            section.page_width,section.page_height=section.page_height,section.page_width
        section.left_margin=Inches(.75);section.right_margin=Inches(.75)
main.save(OUT/'fastPLS_manuscript_revised.docx')
supp.save(OUT/'fastPLS_supplement_revised.docx')
print(OUT)
