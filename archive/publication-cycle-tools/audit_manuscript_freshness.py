"""Inventory manuscript evidence without executing any fitted-model code."""
from pathlib import Path
from docx import Document
import csv
import hashlib
import json
import re

ROOT = Path(__file__).resolve().parents[2]
PKG = ROOT / 'fastPLS_fresh_github'
EXTRA = ROOT / 'fastPLS-extra'
REV = EXTRA / 'manuscript/revision_20260905'
BASE = PKG / 'artifacts/CMPB_rewrite_20260903_cycle138'
OPT = PKG / 'benchmark_results/optimization_20260904'
OUT = EXTRA / 'manuscript/evidence_audit_20260905'
OUT.mkdir(parents=True, exist_ok=True)

def read(path):
    with path.open() as handle:
        return list(csv.DictReader(handle))

def write(name, records):
    with (OUT / name).open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)

def fingerprint(data):
    return hashlib.sha256(data).hexdigest()

def images(doc):
    return [doc.part.related_parts[x._inline.graphic.graphicData.pic.blipFill.blip.embed].blob
            for x in doc.inline_shapes]

def table_text(table):
    return '\n'.join('\t'.join(c.text for c in r.cells) for r in table.rows)

inventory = []
for name, base_name in [('manuscript','main'), ('supplement','supplement')]:
    doc = Document(REV / f'fastPLS_{name}_revised.docx')
    old = Document(BASE / f'fastPLS_CMPB_{base_name}_0.99.39.docx')
    old_images = {fingerprint(b) for b in images(old)}
    old_tables = {fingerprint(table_text(t).encode()) for t in old.tables}
    for i, data in enumerate(images(doc), 1):
        unchanged = fingerprint(data) in old_images
        inventory.append(dict(document=name, kind='embedded_figure', item=i,
            status='inherited_asset_not_regenerated' if unchanged else 'new_asset_selected_sources_only',
            detail='Image bytes identical to input manuscript; resizing is not a data update.' if unchanged
            else 'IKPLS panel: CPU uses earlier saved fastPLS rows; CUDA uses a selected newer workload.'))
    for i, table in enumerate(doc.tables, 1):
        same = fingerprint(table_text(table).encode()) in old_tables
        inventory.append(dict(document=name, kind='table', item=i,
            status='inherited_values' if same else 'changed_requires_source_mapping',
            detail=table_text(table).replace('\n',' ')[:180]))
    for p in doc.paragraphs:
        if re.match(r'^(Table S\d+|Figure (?:S\d+|\d+)[A-Za-z]?)\.', p.text):
            inventory.append(dict(document=name, kind='caption', item=p.text.split('.')[0],
                status='caption_inventoried_not_a_freshness_guarantee', detail=p.text))
write('all_tables_figures.csv', inventory)

old = read(OPT / 'float_factor_scores_nmr_summary/summary.csv')
new = read(OPT / 'cpu_prediction_prefix_v2/nmr_summary/summary.csv')
keys = ['family','backend','precision','ncomp','oversample','power','protocol_version',
        'water_columns_masked','response_columns_scored','seeds']
lookup = {tuple(r[k] for k in keys):r for r in old}
changes = []
for row in new:
    previous = lookup.get(tuple(row[k] for k in keys))
    changes.append({k: row[k] for k in keys} | {
        'draft_time_sec':previous['median_total_sec'] if previous else '',
        'newer_time_sec':row['median_total_sec'],
        'draft_RMSD':previous['median_RMSD'] if previous else '',
        'newer_RMSD':row['median_RMSD'],
        'newer_source':str(OPT / 'cpu_prediction_prefix_v2/nmr_summary/summary.csv'),
        'status':'newer_matched_protocol_evidence' if previous else 'new_component_count_not_in_draft_table'})
write('nmr_missed_updates.csv', changes)

families = {
 'external CPU / R comparison':'publication_cuda_fold_reuse_v1/external_simpls/external_simpls_timing_summary.csv',
 'selected CPU / CUDA':'publication_cuda_fold_reuse_v1/selected_backend/matched_cuda_summary.csv',
 'selected CPU / Metal':'publication_metal_fold_reuse_v1/selected_backend/matched_metal_summary.csv',
 'Metal component paths':'publication_metal_fold_reuse_v1/component_paths/component_path_summary.csv',
 'compiled CPU / CUDA CV':'publication_cuda_fold_reuse_v1/cv_compiled_vs_r_loop/cv_compiled_vs_r_loop_summary.csv',
 'compiled CPU / Metal CV':'publication_metal_fold_reuse_v1/cv_compiled_vs_r_loop/cv_compiled_vs_r_loop_summary.csv',
 'CPU prediction timing':'cpu_flash_summary/prediction_summary.csv',
 'CPU prediction memory':'cpu_flash_summary/memory_summary.csv',
 'NMR latest local endpoint panel':'cpu_prediction_prefix_v2/nmr_summary/summary.csv',
}
sources = []
for family, relative in families.items():
    path = OPT / relative
    data = read(path) if path.exists() else []
    sources.append(dict(analysis=family, source=str(path), rows=len(data),
        status='requires_contract_and_source_match_not_automatically_final_release',
        irlba_cells=sum('irlba' in str(v).lower() for r in data for v in r.values())))
write('candidate_source_inventory.csv', sources)

summary = {
 'all_material_current':False,
 'main_inherited_images':sum(r['document']=='manuscript' and r['status']=='inherited_asset_not_regenerated' for r in inventory),
 'supplement_inherited_images':sum(r['document']=='supplement' and r['status']=='inherited_asset_not_regenerated' for r in inventory),
 'inherited_supplement_tables':sum(r['document']=='supplement' and r['status']=='inherited_values' for r in inventory),
 'missing_newer_nmr_rows':len(changes),
 'notes':[
 'Package version 0.99.39 alone does not distinguish successive numerical sources.',
 'The five-component NMR PLS-SVD selection predates the corrected latent-weight scorer; later selection is 75.',
 'CPU/IKPLS final-source comparison is not established by the existing draft.',
 'Latest-code runtime preservation needs matched-current builds; extraction metric tests are not timing tests.',
 'No old fastPLS, IKPLS, or external model was executed by this audit.']}
(OUT/'summary.json').write_text(json.dumps(summary,indent=2)+'\n')
print(json.dumps(summary,indent=2))
