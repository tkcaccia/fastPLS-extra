#!/usr/bin/env python3

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SOURCE = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "documents/fastPLS_manuscript_figure3_figure4_0.99.66.docx"
)
FIGURE = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "figures/figure3_nmr_family_components.png"
)
OUTPUT = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "documents/fastPLS_manuscript_figure3_reference_figure4_0.99.66.docx"
)


REPLACEMENTS = {
    "Because the families use different component counts, family contrasts are descriptive; backend comparisons remain matched within each family.":
        "Because the families use different component counts, family contrasts are descriptive; backend comparisons remain matched within each family. The deposited fastsimpls implementation from the preceding NMR study was rerun as PLS-SVD with IRLBA at 165 components using the identical prepared split, predictor preprocessing and response metric. It is presented as a deposited-workflow reference rather than a matched backend comparison because implementation, solver, precision and component count differ.",
    "All routes used the same prepared split, float32 inputs, rSVD controls and seed; component count was fixed within each family.":
        "All current fastPLS routes used the same prepared split, float32 inputs, rSVD controls and seed; component count was fixed within each family. The deposited PLS-SVD/IRLBA workflow used float64 and 165 components. In its matched-data rerun, it required a median of 457.511 s and a baseline-corrected peak RSS increment of 3,747.97 MiB, with held-out RMSD 0.000785806. These values provide a reference for the deposited implementation but are not interpreted as an estimator-, precision- or component-matched speed ratio.",
}


CAPTION = (
    "Figure 3. NMR prediction and deposited PLS-SVD/IRLBA reference. Current "
    "fastPLS PLS-SVD was evaluated at 100 components and SIMPLS at 50 "
    "components using float32 rSVD; CPU, CUDA and Metal routes are shown for "
    "each family. The deposited fastsimpls PLS-SVD/IRLBA workflow used float64 "
    "and 165 components and is shown in gray. Panels compare median fitting-plus-"
    "prediction time with interquartile range over three isolated processes, "
    "held-out RMSD, baseline-corrected peak process RSS, CUDA device peak memory, "
    "per-spectrum RMSD and observed/predicted spectra for the held-out sample "
    "nearest the median Linux CPU SIMPLS error. The runtime panel uses a pseudo-"
    "logarithmic scale that includes zero so both the 457.511-s deposited value "
    "and current routes remain visible. CPU and CUDA were measured on the Intel/"
    "NVIDIA workstation, whereas Metal was measured on the Apple M3; cross-"
    "platform timings are not interpreted as hardware speed ratios. Input "
    "conversion is excluded, accelerator initialization and transfer are "
    "included, and all 28,355 response variables contribute to the metrics."
)


def replace_once(paragraph, old, new):
    if old not in paragraph.text:
        return False
    value = paragraph.text.replace(old, new)
    formatting = None
    if paragraph.runs:
        first = paragraph.runs[0]
        formatting = (first.bold, first.italic, first.underline)
    paragraph.clear()
    run = paragraph.add_run(value)
    if formatting is not None:
        run.bold, run.italic, run.underline = formatting
    return True


def main():
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError("The source manuscript or revised Figure 3 is missing.")
    document = Document(SOURCE)

    counts = {text: 0 for text in REPLACEMENTS}
    for paragraph in document.paragraphs:
        for old, new in REPLACEMENTS.items():
            if replace_once(paragraph, old, new):
                counts[old] += 1
    mismatches = [(text, count) for text, count in counts.items() if count != 1]
    if mismatches:
        raise RuntimeError("Prose replacement mismatch: " + repr(mismatches))

    caption_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.startswith("Figure 3.")
    )
    image_paragraph = document.paragraphs[caption_index - 1]
    image_paragraph.clear()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(FIGURE), width=Inches(6.65))

    caption_paragraph = document.paragraphs[caption_index]
    caption_paragraph.clear()
    caption_run = caption_paragraph.add_run(CAPTION)
    caption_run.italic = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
