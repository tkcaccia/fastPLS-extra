#!/usr/bin/env python3
"""Replace Figure S12 and its local methods/caption text in a supplement."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def paragraph_before(paragraph):
    element = paragraph._p.getprevious()
    while element is not None and element.tag != qn("w:p"):
        element = element.getprevious()
    if element is None:
        raise ValueError("No picture paragraph precedes Figure S12.")
    from docx.text.paragraph import Paragraph

    return Paragraph(element, paragraph._parent)


def replace_picture(paragraph, image_path):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.45))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--figure", required=True)
    args = parser.parse_args()

    document = Document(args.input)
    methods = find_paragraph(document, "NMR component selection used")
    methods.text = (
        "NMR component selection used only the predefined 1,200-spectrum training "
        "partition. The candidate grid was 1, 2, 3, 5, 8, 10, 25, 50, 75, 100, "
        "125, 150, 165, 175, 200, 250 and 300 components. Five fixed 80/20 inner "
        "splits used seeds 123, 456, 789, 1011 and 2027. For each family, the "
        "selected value was the smallest count whose mean validation RMSD did not "
        "exceed the minimum mean RMSD plus its standard error. Figure S12 reports a "
        "separate held-out analysis. Each family, route and component count was fitted "
        "on all 1,200 training spectra and evaluated against the fixed 321-spectrum "
        "test partition in three isolated processes using float32 inputs, rSVD seed "
        "123 and fastPLS 0.3. Complete fitting-plus-prediction time and "
        "baseline-corrected peak host RSS were monitored over the same interval. Mac "
        "CPU and Metal were measured on the Apple workstation; Linux CPU and CUDA "
        "were measured on the NVIDIA workstation. Test responses were not used for "
        "fitting or component selection. OPLS used one orthogonal component and "
        "kernel PLS used the linear kernel."
    )

    caption = find_paragraph(document, "Figure S12.")
    replace_picture(paragraph_before(caption), Path(args.figure))
    caption.text = (
        "Figure S12. NMR component-dependent prediction and computation using "
        "float32 inputs in fastPLS 0.3. PLS-SVD, SIMPLS-family, OPLS and linear "
        "kernel "
        "PLS were fitted only on the predefined 1,200-spectrum training partition "
        "and evaluated on the fixed 321-spectrum test partition. Rows report held-out "
        "RMSD, complete fitting-plus-prediction time and baseline-corrected peak host "
        "RSS. Mac CPU, Metal, Linux CPU and CUDA are distinguished by colour. Points "
        "are medians of three isolated-process runs with rSVD seed 123; ribbons show "
        "the interquartile range. Dotted lines mark component counts selected "
        "independently by training-only inner validation. Vertical scales are "
        "family- and measure-specific."
    )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    main()
