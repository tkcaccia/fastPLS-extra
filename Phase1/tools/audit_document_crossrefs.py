#!/usr/bin/env python3
"""Audit main-text citations to main and supplementary figures and tables."""

import re
import sys
from pathlib import Path

from docx import Document


def paragraph_texts(path):
    doc = Document(path)
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text.strip() for p in cell.paragraphs if p.text.strip())
    return texts


def narrative_paragraphs(path):
    """Return document paragraphs before the bibliography begins."""
    doc = Document(path)
    texts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.lower() == "references":
            break
        if text:
            texts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(
                    p.text.strip() for p in cell.paragraphs if p.text.strip()
                )
    return texts


def bibliography_numbers(path):
    doc = Document(path)
    in_references = False
    numbers = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.lower() == "references":
            in_references = True
            continue
        if not in_references:
            continue
        match = re.match(r"^(?:\[(\d+)\]|(\d+)\.)\s+", text)
        if match:
            numbers.append(int(match.group(1) or match.group(2)))
    return numbers


def captions(texts, kind):
    pattern = re.compile(rf"^(?:Supplementary\s+)?{kind}\s+(S?\d+)\.", re.I)
    return [(m.group(1).upper(), text) for text in texts if (m := pattern.match(text))]


def cited_labels(texts, kind):
    found = set()
    pattern = re.compile(
        rf"(?:Supplementary\s+)?{kind}s?\s+"
        r"((?:S?\d+)(?:\s*(?:-|\u2013|to|and|,)\s*S?\d+)*)",
        re.I,
    )
    for text in texts:
        for match in pattern.finditer(text):
            token = match.group(1).upper()
            prefix = "S" if "S" in token else ""
            nums = [int(value) for value in re.findall(r"\d+", token)]
            if re.search(r"(?:-|\u2013|TO)", token) and len(nums) == 2:
                nums = list(range(nums[0], nums[1] + 1))
            found.update(f"{prefix}{value}" for value in nums)
    return found


def cited_label_order(texts, kind):
    ordered = []
    seen = set()
    pattern = re.compile(
        rf"(?:Supplementary\s+)?{kind}s?\s+"
        r"((?:S?\d+)(?:\s*(?:-|\u2013|to|and|,)\s*S?\d+)*)",
        re.I,
    )
    for text in texts:
        for match in pattern.finditer(text):
            token = match.group(1).upper()
            prefix = "S" if "S" in token else ""
            nums = [int(value) for value in re.findall(r"\d+", token)]
            if re.search(r"(?:-|\u2013|TO)", token) and len(nums) == 2:
                nums = list(range(nums[0], nums[1] + 1))
            for value in nums:
                label = f"{prefix}{value}"
                if label not in seen:
                    seen.add(label)
                    ordered.append(label)
    return ordered


def numeric_citations(texts):
    out = set()
    for text in texts:
        for token in re.findall(r"\[([0-9,;\-\u2013 ]+)\]", text):
            for part in re.split(r"[,;]", token):
                nums = [int(value) for value in re.findall(r"\d+", part)]
                if re.search(r"[-\u2013]", part) and len(nums) == 2:
                    out.update(range(nums[0], nums[1] + 1))
                else:
                    out.update(nums)
    return out


def numeric_citation_order(texts):
    ordered = []
    seen = set()
    for text in texts:
        for token in re.findall(r"\[([0-9,;\-\u2013 ]+)\]", text):
            values = []
            for part in re.split(r"[,;]", token):
                nums = [int(value) for value in re.findall(r"\d+", part)]
                if re.search(r"[-\u2013]", part) and len(nums) == 2:
                    values.extend(range(nums[0], nums[1] + 1))
                else:
                    values.extend(nums)
            for value in values:
                if value not in seen:
                    seen.add(value)
                    ordered.append(value)
    return ordered


def supplementary_object_citation_locations(path):
    """Return supplementary figure/table citations with their current section."""
    document = Document(path)
    section = ""
    locations = []
    pattern = re.compile(r"Supplementary\s+(?:Figure|Table)s?\s+S\d+", re.I)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and text:
            section = text
        if pattern.search(text) and not re.match(
            r"^Supplementary\s+(?:Figure|Table)\s+S\d+\.", text, re.I
        ):
            locations.append((section, text))
    return locations


def main():
    main_path, supp_path = map(Path, sys.argv[1:3])
    main_texts = paragraph_texts(main_path)
    supp_texts = paragraph_texts(supp_path)
    main_narrative = [
        t for t in narrative_paragraphs(main_path)
        if not re.match(r"^(?:Supplementary\s+)?(?:Figure|Table)\s+S?\d+\.", t, re.I)
    ]

    main_figs = captions(main_texts, "Figure")
    main_tabs = captions(main_texts, "Table")
    supp_figs = captions(supp_texts, "Figure")
    supp_tabs = captions(supp_texts, "Table")
    cited_figs = cited_labels(main_narrative, "Figure")
    cited_tabs = cited_labels(main_narrative, "Table")
    first_figures = cited_label_order(main_narrative, "Figure")
    first_tables = cited_label_order(main_narrative, "Table")

    print("MAIN_FIGURES", [x[0] for x in main_figs])
    print("MAIN_TABLES", [x[0] for x in main_tabs])
    print("SUPP_FIGURES", [x[0] for x in supp_figs])
    print("SUPP_TABLES", [x[0] for x in supp_tabs])
    print("CITED_FIGURES", sorted(cited_figs, key=lambda x: (x.startswith("S"), int(x.lstrip("S")))))
    print("CITED_TABLES", sorted(cited_tabs, key=lambda x: (x.startswith("S"), int(x.lstrip("S")))))
    print("FIRST_FIGURE_CITATIONS", first_figures)
    print("FIRST_TABLE_CITATIONS", first_tables)
    print("MISSING_MAIN_FIGURES", [x[0] for x in main_figs if x[0] not in cited_figs])
    print("MISSING_MAIN_TABLES", [x[0] for x in main_tabs if x[0] not in cited_tabs])
    print("MISSING_SUPP_FIGURES", [x[0] for x in supp_figs if x[0] not in cited_figs])
    print("MISSING_SUPP_TABLES", [x[0] for x in supp_tabs if x[0] not in cited_tabs])
    available_figs = {x[0] for x in main_figs + supp_figs}
    available_tabs = {x[0] for x in main_tabs + supp_tabs}
    print(
        "CITED_BUT_ABSENT_FIGURES",
        sorted(cited_figs - available_figs,
               key=lambda x: (x.startswith("S"), int(x.lstrip("S")))),
    )
    print(
        "CITED_BUT_ABSENT_TABLES",
        sorted(cited_tabs - available_tabs,
               key=lambda x: (x.startswith("S"), int(x.lstrip("S")))),
    )
    print(
        "SUPP_FIGURE_SEQUENCE",
        [x[0] for x in supp_figs],
        "expected",
        [f"S{i}" for i in range(1, len(supp_figs) + 1)],
    )
    print(
        "SUPP_TABLE_SEQUENCE",
        [x[0] for x in supp_tabs],
        "expected",
        [f"S{i}" for i in range(1, len(supp_tabs) + 1)],
    )

    bibliography = bibliography_numbers(main_path)
    cited_numbers = numeric_citations(main_narrative + supp_texts)
    print("BIBLIOGRAPHY", bibliography)
    print("CITED_NUMBERS", sorted(cited_numbers))
    print("FIRST_REFERENCE_CITATIONS", numeric_citation_order(main_narrative))
    print("UNCITED_REFERENCES", sorted(set(bibliography) - cited_numbers))
    print("MISSING_REFERENCES", sorted(cited_numbers - set(bibliography)))

    print("MAIN_SUPPLEMENTARY_REFERENCE_PARAGRAPHS")
    for text in main_narrative:
        if "Supplementary" in text:
            print("-", text)

    expected_supp_figures = [f"S{i}" for i in range(1, len(supp_figs) + 1)]
    expected_supp_tables = [f"S{i}" for i in range(1, len(supp_tabs) + 1)]
    expected_main_figures = [str(i) for i in range(1, len(main_figs) + 1)]
    expected_main_tables = [str(i) for i in range(1, len(main_tabs) + 1)]
    observed_main_figures = [label for label, _ in main_figs]
    observed_main_tables = [label for label, _ in main_tabs]
    observed_supp_figures = [label for label, _ in supp_figs]
    observed_supp_tables = [label for label, _ in supp_tabs]
    cited_main_figures = [label for label in first_figures if not label.startswith("S")]
    cited_main_tables = [label for label in first_tables if not label.startswith("S")]
    cited_supp_figures = [label for label in first_figures if label.startswith("S")]
    cited_supp_tables = [label for label in first_tables if label.startswith("S")]
    supplementary_locations = supplementary_object_citation_locations(main_path)
    failures = []
    if observed_main_figures != expected_main_figures:
        failures.append(
            "main figure captions are not consecutive: "
            f"{observed_main_figures}"
        )
    if observed_main_tables != expected_main_tables:
        failures.append(
            "main table captions are not consecutive: "
            f"{observed_main_tables}"
        )
    if cited_main_figures != expected_main_figures:
        failures.append(
            "main figures are not first cited in order: "
            f"{cited_main_figures}"
        )
    if cited_main_tables != expected_main_tables:
        failures.append(
            "main tables are not first cited in order: "
            f"{cited_main_tables}"
        )
    if observed_supp_figures != expected_supp_figures:
        failures.append(
            "supplementary figure captions are not consecutive: "
            f"{observed_supp_figures}"
        )
    if observed_supp_tables != expected_supp_tables:
        failures.append(
            "supplementary table captions are not consecutive: "
            f"{observed_supp_tables}"
        )
    if cited_supp_figures != expected_supp_figures:
        failures.append(
            "supplementary figures are not first cited in order: "
            f"{cited_supp_figures}"
        )
    if cited_supp_tables != expected_supp_tables:
        failures.append(
            "supplementary tables are not first cited in order: "
            f"{cited_supp_tables}"
        )
    misplaced_supplementary = [
        (section, text) for section, text in supplementary_locations
        if not section.startswith("3.")
    ]
    if misplaced_supplementary:
        failures.append(
            "supplementary figure/table citations occur outside Results: "
            + "; ".join(
                f"{section or '[no section]'}: {text[:80]}"
                for section, text in misplaced_supplementary
            )
        )
    missing = (
        [f"main figure {label}" for label, _ in main_figs if label not in cited_figs]
        + [f"main table {label}" for label, _ in main_tabs if label not in cited_tabs]
        + [f"supplementary figure {label}" for label, _ in supp_figs if label not in cited_figs]
        + [f"supplementary table {label}" for label, _ in supp_tabs if label not in cited_tabs]
    )
    if missing:
        failures.append("uncited objects: " + ", ".join(missing))
    absent = sorted((cited_figs - available_figs) | (cited_tabs - available_tabs))
    if absent:
        failures.append("citations without objects: " + ", ".join(absent))
    if sorted(cited_numbers - set(bibliography)):
        failures.append("numeric citations without references")
    if failures:
        raise SystemExit("Cross-reference audit failed:\n- " + "\n- ".join(failures))
    print("Cross-reference audit passed.")


if __name__ == "__main__":
    main()
