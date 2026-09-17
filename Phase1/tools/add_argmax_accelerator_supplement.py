#!/usr/bin/env python3
"""Insert the argmax-only CPU/accelerator figure into the supplement."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches


def insert_before(anchor, paragraph) -> None:
    anchor._p.addprevious(paragraph._p)


def add_text_before(document, anchor, text, style=None, page_break=False):
    paragraph = document.add_paragraph(style=style)
    if page_break:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run(text)
    insert_before(anchor, paragraph)
    return paragraph


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("figure_png", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()

    document = Document(args.input_docx)
    anchors = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "S10. ImageNet stress test"
    ]
    if len(anchors) != 1:
        raise RuntimeError("Expected one S10. ImageNet stress test anchor")
    anchor = anchors[0]

    add_text_before(
        document,
        anchor,
        "Argmax classification across accelerator backends",
        style="Heading 2",
        page_break=True,
    )
    add_text_before(
        document,
        anchor,
        (
            "Figure S18 isolates the argmax classification workflows from the "
            "complete CPU/accelerator analysis in Figure 2. The same ten datasets, "
            "float32 inputs, component counts and three fresh-process repetitions "
            "were retained for PLS-SVD, SIMPLS, OPLS and linear kernel PLS. CPU and "
            "CUDA were paired on the Intel/NVIDIA workstation; CPU and Metal were "
            "paired on the Apple M3. No row was removed because of its predictive "
            "metric difference; cross-platform accuracy agreement is reported "
            "separately in Figure S13 and Table S9."
        ),
    )
    add_text_before(
        document,
        anchor,
        (
            "Accelerator initialization, transfer and synchronization overhead "
            "dominated the smaller classification tasks. CUDA provided the largest "
            "runtime benefit for the million-sample ImageNet/DINOv2 workload and "
            "improved three of four CIFAR-100 family routes. Metal improved three "
            "ImageNet family routes but did not reduce runtime consistently on the "
            "remaining workloads. Incremental host-memory ratios describe measured "
            "process growth and do not represent isolated device workspace."
        ),
    )

    figure_paragraph = document.add_paragraph()
    figure_paragraph.alignment = 1
    figure_paragraph.add_run().add_picture(str(args.figure_png), width=Inches(6.45))
    insert_before(anchor, figure_paragraph)
    add_text_before(
        document,
        anchor,
        (
            "Figure S18. Float32 argmax classification across CPU and accelerator "
            "backends. Panels A and B report CPU/accelerator total-runtime ratios; "
            "values above one favour CUDA or Metal. Panels C and D report "
            "accelerator/CPU baseline-corrected host-RSS ratios; values below one "
            "indicate lower host-memory growth for the accelerator route. Total "
            "time includes fitting, held-out prediction, accelerator initialization, "
            "transfer and synchronization. Cells are medians of three fresh "
            "processes using fastPLS 0.99.65. CPU/CUDA measurements used an Intel "
            "Core i7-13700 and NVIDIA GeForce RTX 5060 Ti; CPU/Metal measurements "
            "used an Apple M3. All completed argmax routes are shown irrespective "
            "of metric difference."
        ),
    )

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)


if __name__ == "__main__":
    main()
