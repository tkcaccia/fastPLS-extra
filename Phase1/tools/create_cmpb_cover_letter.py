#!/usr/bin/env python3
"""Create the cover letter for the fastPLS CMPB submission."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = Path(
    "/Users/stefano/Desktop/manuscripts/fastPLS/revision_20261006/"
    "fastPLS_CMPB_cover_letter.docx"
)


def set_cell_border(cell, colour: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")
        node.set(qn("w:color"), colour)


def add_paragraph(document, text="", *, bold_prefix=None, space_after=7):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.04
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)
    return paragraph


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.04

    properties = document.core_properties
    properties.title = "Cover letter for the fastPLS CMPB submission"
    properties.subject = "Submission to Computer Methods and Programs in Biomedicine"
    properties.author = "Stefano Cacciatore and Leonardo Tenori"

    date = add_paragraph(document, "6 October 2026", space_after=10)
    date.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    add_paragraph(document, "Editor", space_after=0)
    add_paragraph(document, "Computer Methods and Programs in Biomedicine", space_after=0)
    add_paragraph(document, "Elsevier", space_after=10)

    subject = add_paragraph(
        document,
        "Re: Submission of “fastPLS: accelerated SIMPLS for high-dimensional biomedical data”",
        space_after=10,
    )
    subject.runs[0].bold = True

    add_paragraph(document, "Dear Editor,", space_after=8)

    add_paragraph(
        document,
        "On behalf of Dupe Ojo, Alessia Vignoli, Stefano Cacciatore and Leonardo Tenori, "
        "we submit the manuscript entitled “fastPLS: accelerated SIMPLS for high-dimensional "
        "biomedical data” for consideration as a research article in Computer Methods and "
        "Programs in Biomedicine.",
    )

    add_paragraph(
        document,
        "The study addresses a practical limitation of partial least-squares analysis in modern "
        "biomedical applications. Sequential component extraction, repeated validation and dense "
        "coefficient paths can become prohibitive when predictors, observations or response "
        "dimensions are large. fastPLS combines compiled PLS-SVD, a SIMPLS-family estimator, "
        "orthogonal PLS and kernel PLS with reusable matrix products, compact latent prediction, "
        "leakage-free sufficient-statistics cross-validation and optional CUDA execution. The "
        "contribution is therefore an implementation and validation study rather than a claim of "
        "a new PLS estimator.",
    )

    add_paragraph(
        document,
        "We evaluated the implementation on 13 prepared datasets spanning metabolomics, cancer "
        "omics, single-cell measurements, multivariate drug response, image embeddings and "
        "spectral regression. The software had the lowest or tied-lowest median runtime among "
        "completed R workflows in every reported classification and regression task. In the "
        "million-observation ImageNet/DINOv2 experiment, fastPLS completed in 63.93 s, compared "
        "with 180.09 s for IKPLS, while the other tested R workflows did not complete. This "
        "experiment is presented as a foundation-model embedding analysis and computational "
        "feasibility demonstration, not as biomedical validation.",
    )

    add_paragraph(
        document,
        "The principal biomedical demonstration reconstructs 28,355 NMR response intensities "
        "from 13,000-bin spectra using a fixed 1,200-spectrum training set and 321-spectrum test "
        "set. CUDA SIMPLS-family execution required 0.468 s, versus 438.862 s for the deposited "
        "workflow, while also reducing held-out RMSD and memory use. Response-wise and "
        "intensity-stratified analyses are included so that the global error does not conceal "
        "local spectral failures.",
    )

    add_paragraph(
        document,
        "The performance evidence is accompanied by dense numerical comparisons, CPU/CUDA "
        "agreement checks, exact algebraic identities verified in Lean, explicit reporting of "
        "failed or resource-limited routes, and a grouped nested-validation implementation that "
        "keeps related observations together. Source code, benchmark scripts, formal proofs and "
        "the information needed to reproduce the reported workflows are publicly available.",
    )

    add_paragraph(
        document,
        "We believe the manuscript fits the journal because it presents a formal computing "
        "method, an openly available biomedical software implementation and empirical evidence "
        "on realistic biomedical prediction problems. The manuscript is original, is not under "
        "consideration elsewhere, and has been approved by all authors. The authors declare no "
        "competing interests, and the study received no specific grant from funding agencies in "
        "the public, commercial or not-for-profit sectors.",
    )

    add_paragraph(
        document,
        "Thank you for considering our manuscript. We would be pleased to provide any additional "
        "information required during editorial assessment.",
        space_after=10,
    )

    add_paragraph(document, "Yours sincerely,", space_after=9)

    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.3)
    table.columns[1].width = Inches(3.3)
    for cell in table.rows[0].cells:
        set_cell_border(cell)
        cell.vertical_alignment = 1

    left, right = table.rows[0].cells
    left.width = Inches(3.3)
    right.width = Inches(3.3)
    left.text = ""
    right.text = ""

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Stefano Cacciatore")
    r.bold = True
    p = left.add_paragraph("Bioinformatics Unit\nInternational Centre for Genetic Engineering and Biotechnology\nCape Town, South Africa\nstefano.cacciatore@icgeb.org")
    p.paragraph_format.space_after = Pt(0)

    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Leonardo Tenori")
    r.bold = True
    p = right.add_paragraph("CERM and Department of Chemistry “Ugo Schiff”\nUniversity of Florence\nSesto Fiorentino, Italy\nleonardo.tenori@unifi.it")
    p.paragraph_format.space_after = Pt(0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
