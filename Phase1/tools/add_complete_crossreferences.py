#!/usr/bin/env python3
"""Add complete main-text cross-references to current supplementary evidence."""

from pathlib import Path

from docx import Document


ROOT = Path("/Users/stefano/Documents/GPUPLS")
DOCS = ROOT / "manuscript_revision_0.99.66" / "documents"
MAIN_IN = DOCS / "fastPLS_manuscript_author_KODAMA_citations_0.99.66.docx"
SUPP_IN = DOCS / "fastPLS_supplement_figure2_metal_0.99.66.docx"
MAIN_OUT = DOCS / "fastPLS_manuscript_ordered_crossreferences_0.99.66.docx"
SUPP_OUT = DOCS / "fastPLS_supplement_ordered_crossreferences_0.99.66.docx"


REPLACEMENTS = {
    "The classification comparison used CCLE, CIFAR-100 [24], GTEx v8, MetRef, Retina, Tabula Muris, TCGA-BRCA, TCGA-HNSC methylation and TCGA Pan-Cancer. CBMC CITE-seq and PRISM supplied multivariate regression workloads, and NMR was analysed separately. Dataset construction, preprocessing, dimensions, acquisition and redistribution restrictions are reported in Supplementary Tables S1 and S2. The same prepared train/test split, component count and precision were used within each paired comparison.":
        "The classification comparison used CCLE, CIFAR-100 [24], GTEx v8, MetRef, Retina, Tabula Muris, TCGA-BRCA, TCGA-HNSC methylation and TCGA Pan-Cancer. CBMC CITE-seq and PRISM supplied multivariate regression workloads, and NMR was analysed separately. Dataset construction, preprocessing, dimensions, acquisition and redistribution restrictions are reported in Supplementary Tables S1 and S2. The same prepared train/test split, component count and precision were used within each paired comparison. Component-dependent predictive performance, fitting-plus-prediction time and host-memory paths are reported for the non-NMR datasets and NMR in Supplementary Figures S1-S12.",
    "The independent-implementation comparison used an Ubuntu 22.04 Linux workstation with an Intel Core i7-13700 processor and 32 GiB RAM. Every implementation used one effective CPU thread.":
        "The independent-implementation comparison used an Ubuntu 22.04 Linux workstation with an Intel Core i7-13700 processor and 32 GiB RAM. Every implementation used one effective CPU thread. Dominant operations and residency, numerical validation, the SIMPLS implementation ablation and the capabilities of compared implementations are reported in Supplementary Tables S3-S6.",
    "Host-memory results are baseline-corrected complete-process RSS increments; CUDA device allocation is reported separately in the Supplementary material.":
        "Host-memory results are baseline-corrected complete-process RSS increments; paired metrics, runtime, host memory, CUDA device allocation and executed residency are reported in Supplementary Tables S7 and S8.",
    "The matched Mac CPU/Metal comparison is reported separately in Supplementary Figure S19.":
        "The matched Mac CPU/Metal comparison is reported separately in the supplementary analysis.",
    "Metal was evaluated in float32 only because Apple GPUs do not provide the native float64 arithmetic required by this route.":
        "Metal was evaluated in float32 only because Apple GPUs do not provide the native float64 arithmetic required by this route (Supplementary Table S9 and Figure S13).",
    "This comparison evaluates solver cost within the same fastPLS execution design; IRLBA is not part of the public package.":
        "This comparison evaluates solver cost within the same fastPLS execution design; IRLBA is not part of the public package (Supplementary Table S10 and Figure S14). Training-only component settings and their associations with predictive and computational metrics are reported in Supplementary Tables S11 and S12.",
    "It is presented as a deposited-workflow reference rather than a matched backend comparison because implementation, solver, precision and component count differ.":
        "It is presented as a deposited-workflow reference rather than a matched backend comparison because implementation, solver, precision and component count differ. The training-only NMR selection decisions and fixed-component implementation comparison are reported in Supplementary Tables S13 and S14.",
    "A same-code ablation disabled one optimization at a time to quantify its route-dependent contribution.":
        "A same-code ablation disabled one optimization at a time to quantify its route-dependent contribution. CPU thread-scaling results are reported in Supplementary Table S15 and Figures S15 and S16; the public precision and backend capability matrix is reported in Supplementary Table S16.",
    "The ImageNet/DINOv2 analysis was a separate single-run 1,000-component boundary stress test: fastPLS SIMPLS-LDA achieved accuracy 0.8094 in 63.3 s, whereas IKPLS achieved accuracy 0.7999 in 197.1 s with an argmax prediction rule.":
        "The ImageNet/DINOv2 analysis was a separate single-run 1,000-component boundary stress test: fastPLS SIMPLS-LDA achieved accuracy 0.8094 in 63.3 s, whereas IKPLS achieved accuracy 0.7999 in 197.1 s with an argmax prediction rule. Detailed independent-Python results and the matched CUDA software comparison are reported in Supplementary Table S17 and Figure S17.",
    "Full metric values and signed differences are reported in Supplementary Table S9 and Figure S13.":
        "Full metric values and signed differences are reported in Supplementary Table S9 and Figure S13. The corresponding argmax-only accelerator analysis and the complete CPU/Metal comparison are shown in Supplementary Figures S18 and S19.",
    "Metal completed 52 of 52 pairs and was faster in 3 (Supplementary Figure S19).":
        "Metal completed 52 of 52 pairs and was faster in 3.",
    "At 100 PLS-SVD components, float32 fitting plus prediction required a median of 6.128 s on the Linux CPU and 0.500 s on CUDA, a 12.3-fold same-workstation runtime ratio.":
        "At 100 PLS-SVD components, float32 fitting plus prediction required a median of 6.128 s on the Linux CPU and 0.500 s on CUDA, a 12.3-fold same-workstation runtime ratio (Figure 3; Supplementary Table S14).",
    "The ImageNet/DINOv2 experiment used fastPLS 0.99.66 to fit float32 SIMPLS on 1,000,000 training embeddings and evaluate 281,167 held-out embeddings over 50–1,000 requested components with argmax and LDA.":
        "The ImageNet/DINOv2 experiment used fastPLS 0.99.66 to fit float32 SIMPLS on 1,000,000 training embeddings and evaluate 281,167 held-out embeddings over 50–1,000 requested components with argmax and LDA (Supplementary Table S18).",
    "Future work will expand validation on other GPU architectures, strengthen route-specific numerical diagnostics and refine automatic dispatch using controlled multi-factor scaling experiments.":
        "Future work will expand validation on other GPU architectures, strengthen route-specific numerical diagnostics and refine automatic dispatch using controlled multi-factor scaling experiments. The complete reproducibility, platform and CPU numerical-library summary is provided in Supplementary Table S19.",
}


def replace_fragment(paragraph, old, new):
    current = paragraph.text
    if old not in current:
        return False
    updated = current.replace(old, new)
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(updated)
    else:
        runs[0].text = updated
        for run in runs[1:]:
            run.text = ""
    return True


def main():
    doc = Document(MAIN_IN)
    applied = {old: False for old in REPLACEMENTS}
    for paragraph in doc.paragraphs:
        for old, new in REPLACEMENTS.items():
            if replace_fragment(paragraph, old, new):
                applied[old] = True
    missing = [old for old, done in applied.items() if not done]
    if missing:
        raise RuntimeError("Unmatched replacement fragments:\n" + "\n---\n".join(missing))
    doc.save(MAIN_OUT)

    supplement = Document(SUPP_IN)
    citation_fixed = False
    for table in supplement.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_fragment(
                        paragraph,
                        "Study-specific prepared matrices associated with Vignoli et al. [9]",
                        "Study-specific prepared matrices associated with Vignoli et al. [7]",
                    ):
                        citation_fixed = True
    if not citation_fixed:
        raise RuntimeError("The supplementary NMR citation was not found")
    supplement.save(SUPP_OUT)
    print(MAIN_OUT)
    print(SUPP_OUT)


if __name__ == "__main__":
    main()
