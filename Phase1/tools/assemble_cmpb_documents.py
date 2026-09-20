#!/usr/bin/env python3
"""Insert one audited CMPB campaign's tables and figures into the DOCX files."""

from __future__ import annotations

import argparse
import csv
from copy import deepcopy
import json
import math
from pathlib import Path
import re
from statistics import median

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
GRID = "D9D9D9"


TABLE_SPECS = {
    "S1": {
        "file": "TableS1_numerical_validation.csv",
        "old_caption": None,
        "caption": (
            "Table S1. Numerical validation of the evaluated implementation. "
            "Completed denotes "
            "calculations that returned a result; meeting the stated numerical "
            "tolerances is reported separately from completion."
        ),
        "columns": [
            ("validation_block", "Validation block"),
            ("route", "Route"),
            ("attempted", "Attempted"),
            ("completed", "Completed"),
            ("key_numerical_result", "Numerical result"),
            ("failures", "Failures"),
        ],
        "widths": (1.00, 1.20, 0.58, 0.62, 2.65, 0.55),
        "font": 6.7,
    },
    "S2": {
        "file": "TableS2_formal_invariants.csv",
        "old_caption": "Table S1. Algebraic invariants",
        "caption": (
            "Table S2. Algebraic invariants checked by Lean. Every listed theorem "
            "was accepted by the pinned Lean 4 and Mathlib proof checker."
        ),
        "columns": [
            ("lean_theorem", "Lean theorem"),
            ("implementation_area", "Implementation area"),
            ("checked_statement", "Checked statement"),
            ("lake_build_status", "Build status"),
        ],
        "widths": (1.28, 1.40, 3.30, 0.62),
        "font": 6.8,
    },
    "S3": {
        "file": "TableS3_independent_method_settings.csv",
        "old_caption": "Table S2. Independent PLS software",
        "caption": (
            "Table S3. Independent PLS software, functions and parameter settings "
            "used in Figure 1."
        ),
        "columns": None,
        "widths": None,
        "font": 6.3,
    },
    "S4": {
        "file": "TableS4_timing_memory_measurement.csv",
        "old_caption": "Table S3. Runtime and complete-process memory",
        "caption": (
            "Table S4. Runtime and complete-process memory measurement used for "
            "the R and Python comparisons."
        ),
        "columns": None,
        "widths": None,
        "font": 6.8,
    },
    "S5": {
        "file": "TableS5_fit_prediction_cpu_cuda.csv",
        "old_caption": "Table S4. Fitting and prediction time",
        "caption": (
            "Table S5. Median fitting-and-prediction time and predictive metric "
            "for the selected float32 CPU/CUDA comparisons. CPU/CUDA ratios above "
            "one favour CUDA."
        ),
        "columns": [
            ("dataset", "Dataset"), ("family", "Family"),
            ("components", "A"), ("cpu_time_seconds", "CPU s"),
            ("cuda_time_seconds", "CUDA s"),
            ("cpu_cuda_time_ratio", "CPU/CUDA"),
            ("metric", "Metric"), ("cpu_metric", "CPU metric"),
            ("cuda_metric", "CUDA metric"), ("status", "Status"),
        ],
        "widths": (0.90, 0.72, 0.35, 0.55, 0.55, 0.63, 0.55, 0.65, 0.68, 0.62),
        "font": 5.8,
        "cell_margin": 24,
        "paragraph_after": 0.4,
    },
    "S6": {
        "file": "TableS6_memory_cpu_cuda.csv",
        "old_caption": "Table S5. Host-memory increase",
        "caption": (
            "Table S6. Median process-memory measurements for the selected "
            "float32 CPU/CUDA comparisons. Host values are baseline and "
            "baseline-corrected peak RSS; device memory is reported separately."
        ),
        "columns": [
            ("dataset", "Dataset"), ("family", "Family"),
            ("components", "A"),
            ("cpu_baseline_rss_mib", "CPU base MiB"),
            ("cuda_baseline_rss_mib", "CUDA base MiB"),
            ("cpu_incremental_peak_rss_mib", "CPU incr. MiB"),
            ("cuda_incremental_peak_rss_mib", "CUDA incr. MiB"),
            ("cuda_cpu_increment_ratio", "CUDA/CPU incr."),
            ("cuda_device_peak_mib", "Device MiB"),
            ("status", "Status"),
        ],
        "widths": (0.88, 0.70, 0.34, 0.62, 0.65, 0.65, 0.68, 0.70, 0.62, 0.56),
        "font": 5.7,
        "cell_margin": 24,
        "paragraph_after": 0.4,
    },
    "S7": {
        "file": "TableS7_imagenet_simpls_component_path.csv",
        "old_caption": "Table S7. ImageNet/DINOv2 predictions",
        "caption": (
            "Table S7. ImageNet/DINOv2 SIMPLS-family component path. One maximal "
            "1,000-component model supplied all prefixes; top-1 and top-5 "
            "accuracy were evaluated on the fixed held-out partition."
        ),
        "columns": [
            ("classifier", "Classifier"),
            ("ncomp_requested", "A"),
            ("top1_accuracy", "Top-1"),
            ("top5_accuracy", "Top-5"),
            ("fit_predict_time_sec", "Fit + predict s"),
            ("top5_prediction_time_sec", "Top-5 s"),
            ("total_time_sec", "Total s"),
            ("process_peak_rss_mb", "Host peak MiB"),
            ("gpu_peak_mb", "Device peak MiB"),
        ],
        "widths": (0.80, 0.40, 0.55, 0.55, 0.82, 0.60, 0.62, 0.78, 0.78),
        "font": 6.2,
    },
    "S8": {
        "file": "TableS8_benchmark_platform.csv",
        "old_caption": "Table S8. Software interfaces, hardware",
        "caption": (
            "Table S8. Evaluated fastPLS release and Linux CPU, OpenBLAS and "
            "CUDA platform used for the CMPB performance experiments."
        ),
        "columns": [("item", "Item"), ("configuration", "Configuration")],
        "widths": (1.55, 5.05),
        "font": 7.1,
    },
}


FIGURE_S14_CAPTION = (
    "Figure S14. Localized error in held-out NMR spectral prediction. Panels "
    "summarize response-wise RMSD across the 28,355 predicted intensities and "
    "error stratified by observed signal intensity, complementing the global "
    "and per-spectrum RMSD in Figure 3."
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise RuntimeError(f"No rows in {path}")
    return rows


def paragraph_starting(document, start: str):
    matches = [p for p in document.paragraphs if p.text.startswith(start)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {start!r}; found {len(matches)}")
    return matches[0]


def next_table_element(paragraph):
    element = paragraph._p.getnext()
    while element is not None:
        if element.tag == qn("w:tbl"):
            return element
        if element.tag == qn("w:p") and "".join(element.itertext()).strip():
            raise RuntimeError(f"No table immediately follows {paragraph.text!r}")
        element = element.getnext()
    raise RuntimeError(f"No table follows {paragraph.text!r}")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, value: int = 50) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(
    table,
    widths: tuple[float, ...] | None,
    font_size: float,
    cell_margin: int = 50,
    paragraph_after: float = 1.2,
) -> None:
    table.style = "Table Grid"
    table.autofit = widths is None
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for column_index, cell in enumerate(row.cells):
            if widths is not None:
                cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, cell_margin)
            set_cell_fill(
                cell,
                BLUE if row_index == 0 else
                PALE_BLUE if row_index % 2 == 0 else WHITE,
            )
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(paragraph_after)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column_index in (0, 1)
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size + (0.2 if row_index == 0 else 0))
                    run.font.bold = row_index == 0
                    run.font.color.rgb = (
                        RGBColor(255, 255, 255) if row_index == 0
                        else RGBColor(0, 0, 0)
                    )


def humanize(name: str) -> str:
    return name.replace("_", " ").strip().capitalize()


def build_table(
    document,
    rows,
    columns,
    widths,
    font_size,
    cell_margin=50,
    paragraph_after=1.2,
):
    if columns is None:
        keys = list(rows[0])
        columns = [(key, humanize(key)) for key in keys]
    table = document.add_table(rows=1, cols=len(columns))
    for cell, (_, label) in zip(table.rows[0].cells, columns):
        cell.text = label
    for row in rows:
        cells = table.add_row().cells
        for cell, (key, _) in zip(cells, columns):
            cell.text = row.get(key, "")
    style_table(table, widths, font_size, cell_margin, paragraph_after)
    return table


def replace_table_after_caption(document, caption, table) -> None:
    old_table = next_table_element(caption)
    old_table.addprevious(table._tbl)
    old_table.getparent().remove(old_table)


def insert_caption_before(document, anchor, text: str):
    paragraph = document.add_paragraph(text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def replace_caption(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def normalize_shared_notation(document) -> None:
    """Normalize terminology in prose, captions, and algorithm cells."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        if table.rows:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        revised = paragraph.text.replace("kernel-PLS", "kernel PLS")
        revised = revised.replace(
            "Current-release numerical validation",
            "Numerical validation of the evaluated implementation",
        )
        revised = revised.replace("SIMPLS family", "SIMPLS-family")
        revised = revised.replace(
            "Classification compares SIMPLS-LDA with IKPLS and argmax",
            "Classification compares fastPLS SIMPLS-family with LDA against "
            "IKPLS with argmax",
        )
        if revised != paragraph.text:
            replace_caption(paragraph, revised)


def normalize_supplement_algorithms(document) -> None:
    """Align supplementary OPLS and kernel PLS algorithms with the C++ core."""
    opls_steps = (
        "Input training predictors X, responses Y or compact class labels z, "
        "requested predictive component-count set C with A = max(C), and O "
        "orthogonal components. Estimate predictor centres and scales from X; "
        "centre Y, or form the equivalent label-aware class cross-product "
        "without constructing a one-hot response.",
        "Standardize X to X(1) and form the current predictor-response "
        "cross-covariance S(1) = X(1)ᵀY(c). For classification, obtain S(1) "
        "from class counts and class-wise predictor sums.",
        "For orthogonal component o = 1,...,O, obtain a unit predictive "
        "direction w(o) from the leading left direction of S(o), using the "
        "configured randomized solver and component-specific seed.",
        "Compute t(o) = X(o)w(o) and p(o) = X(o)ᵀt(o) / [t(o)ᵀt(o)]. If the "
        "score norm is non-finite or zero, stop orthogonal filtering at the "
        "completed prefix; the predictive model remains usable.",
        "Remove the part of p(o) parallel to w(o): w⊥(o) = p(o) - "
        "w(o)[w(o)ᵀp(o)]/[w(o)ᵀw(o)], then normalize w⊥(o). If no nonzero "
        "orthogonal direction remains, stop at the completed prefix.",
        "Compute t⊥(o) = X(o)w⊥(o) and p⊥(o) = X(o)ᵀt⊥(o) / "
        "[t⊥(o)ᵀt⊥(o)]. If the score norm is invalid, stop at the completed "
        "prefix. Otherwise store w⊥(o) and p⊥(o), then update "
        "X(o+1) = X(o) - t⊥(o)p⊥(o)ᵀ.",
        "Refresh S(o+1) from the deflated predictors. An eligible "
        "sufficient-statistics route performs the algebraically corresponding "
        "updates to XᵀX and XᵀY instead of materializing each deflated "
        "predictor matrix. Repeat Steps 3-7 until O components are complete "
        "or no further orthogonal direction is estimable.",
        "Fit one maximal A-component SIMPLS-family predictive path to the "
        "filtered predictors X(O+1) and centred response. Store the completed "
        "orthogonal filter, predictive latent factors, preprocessing statistics "
        "and response mean.",
        "For new predictors X(new), apply the stored training centre and scale. "
        "In extraction order, compute t⊥(o,new) = X(new,o)w⊥(o) and update "
        "X(new,o+1) = X(new,o) - t⊥(o,new)p⊥(o)ᵀ.",
        "For each requested a in C, predict from X(new,O+1) with min(a,e), "
        "where e is the number of effective predictive directions, and restore "
        "the response mean. Prefixes above e repeat the last estimable result; "
        "if e = 0, use the training mean or class prior. Argmax uses "
        "reconstructed response scores, whereas LDA is fitted only from the "
        "corresponding training scores.",
    )
    kernel_steps = (
        "Input training predictors X, responses Y or compact class labels z, "
        "requested component-count set C with A = max(C), and a linear, "
        "radial-basis or polynomial kernel. Estimate predictor centres and "
        "scales from X and the response mean from Y; classification uses the "
        "equivalent label-aware response construction.",
        "Standardize X to X(c). For the linear kernel, set Z = X(c) and "
        "continue to Step 5 without forming an n by n Gram matrix.",
        "For a nonlinear kernel, form K(i,j) from standardized training rows: "
        "exp[-γ||x(i)-x(j)||²] for the radial-basis kernel or "
        "[γx(i)ᵀx(j)+c]ᵈ for the polynomial kernel. Store X(c) as the training "
        "reference.",
        "Compute the training row means, column means and grand mean of K, and "
        "double-centre it as K(c)(i,j) = K(i,j) - rowmean(i) - colmean(j) + "
        "grandmean. Set Z = K(c) and store the training column means and grand "
        "mean for prediction.",
        "Form S = ZᵀY(c), or its label-aware equivalent, and fit one maximal "
        "A-component SIMPLS-family path to Z. Retain compact latent factors, "
        "the response mean and preprocessing statistics. For requested a in C, "
        "use min(a,e), where e is the effective direction count; repeat the "
        "last estimable prefix above e, or retain the training-mean or class-"
        "prior path when e = 0.",
        "For new predictors X(new), apply the stored training centre and scale. "
        "With a linear kernel, set Z(new) = X(new,c).",
        "With a nonlinear kernel, compute the cross-kernel K(new) between "
        "X(new,c) and the stored training reference. Centre each new row using "
        "its own mean together with the stored training column means and "
        "training grand mean; set Z(new) to this centred cross-kernel.",
        "For each requested a in C, predict from Z(new) with the retained "
        "SIMPLS-family prefix and restore the response mean. Argmax uses the "
        "reconstructed response scores; LDA is fitted only from the "
        "corresponding training scores.",
        "During cross-validation, repeat all preprocessing and nonlinear Gram "
        "centring within each training fold. The linear route can use the "
        "leakage-free sufficient-statistics algorithm; the nonlinear route "
        "cannot reuse a full-data centred Gram matrix.",
    )
    found_opls = False
    found_kernel = False
    for table in document.tables:
        if len(table.columns) != 2 or len(table.rows) < 2:
            continue
        if table.rows[0].cells[0].text.strip() != "Step":
            continue
        first = table.rows[1].cells[1].text
        if len(table.rows) == 11 and "orthogonal components" in first:
            for row_index, text in enumerate(opls_steps, start=1):
                replace_caption(table.rows[row_index].cells[1].paragraphs[0], text)
            found_opls = True
        elif len(table.rows) == 10 and "kernel" in first.lower():
            for row_index, text in enumerate(kernel_steps, start=1):
                replace_caption(table.rows[row_index].cells[1].paragraphs[0], text)
            found_kernel = True
    if not found_opls:
        raise RuntimeError("Could not locate the supplementary OPLS algorithm")
    if not found_kernel:
        raise RuntimeError("Could not locate the supplementary kernel PLS algorithm")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def finite(value):
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def show(value, digits=3):
    number = finite(value)
    if number is None:
        return "not available"
    return f"{number:.{digits}f}"


def percent(value, digits=2):
    number = finite(value)
    if number is None:
        return "not available"
    return f"{100 * number:.{digits}f}%"


def row_for(rows, **criteria):
    matches = [
        row for row in rows
        if all(str(row.get(key, "")) == str(value)
               for key, value in criteria.items())
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one narrative row for {criteria}; found {len(matches)}"
        )
    return matches[0]


def replace_started(document, start: str, text: str) -> None:
    replace_caption(paragraph_starting(document, start), text)


def replace_abstract_results(document, text: str) -> None:
    paragraphs = list(document.paragraphs)
    start = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text.startswith("Results:")
    )
    end = next(
        index for index, paragraph in enumerate(paragraphs[start + 1:], start + 1)
        if paragraph.text.startswith("Conclusions:")
    )
    replace_caption(paragraphs[start], text)
    for paragraph in paragraphs[start + 1:end]:
        remove_paragraph(paragraph)


def replace_picture_before_caption(caption, path: Path, width) -> None:
    element = caption._p.getprevious()
    inspected = 0
    while element is not None and inspected < 12:
        if element.tag == qn("w:p"):
            if element.xpath(".//w:drawing") or element.xpath(".//w:pict"):
                from docx.text.paragraph import Paragraph
                picture_paragraph = Paragraph(element, caption._parent)
                properties = picture_paragraph._p.pPr
                picture_paragraph.clear()
                if properties is not None and picture_paragraph._p.pPr is None:
                    picture_paragraph._p.insert(0, deepcopy(properties))
                picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture_paragraph.add_run().add_picture(str(path), width=width)
                return
        element = element.getprevious()
        inspected += 1
    raise RuntimeError(f"No picture paragraph precedes {caption.text!r}")


def usable_width(document):
    section = document.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def insert_after(reference_element, new_element) -> None:
    reference_element.addnext(new_element)


def replace_main_table(document, table_path: Path) -> None:
    caption = paragraph_starting(document, "Table 1.")
    rows = read_csv(table_path)
    columns = [
        ("dataset", "Dataset"), ("task_type", "Task"),
        ("n_train", "Training n"), ("n_test", "Test n"),
        ("p", "p"), ("q", "q"),
    ]
    table = build_table(
        document, rows, columns,
        (1.60, 1.15, 0.82, 0.82, 0.70, 0.70), 7.2,
    )
    replace_table_after_caption(document, caption, table)


def replace_main_figures(document, figures: Path) -> None:
    width = usable_width(document)
    for number in range(1, 5):
        image = figures / f"Figure{number}.png"
        if not image.is_file():
            raise FileNotFoundError(image)
        caption = paragraph_starting(document, f"Figure {number}.")
        replace_picture_before_caption(caption, image, width)


def replace_supplement_tables(document, tables: Path) -> None:
    old_captions = {
        key: paragraph_starting(document, spec["old_caption"])
        for key, spec in TABLE_SPECS.items() if spec["old_caption"]
    }
    anchor = old_captions["S2"]
    first = TABLE_SPECS["S1"]
    first_caption = insert_caption_before(document, anchor, first["caption"])
    first_table = build_table(
        document, read_csv(tables / first["file"]), first["columns"],
        first["widths"], first["font"],
    )
    anchor._p.addprevious(first_table._tbl)

    for label in ("S2", "S3", "S4", "S5", "S6", "S7", "S8"):
        spec = TABLE_SPECS[label]
        caption = old_captions[label]
        table = build_table(
            document, read_csv(tables / spec["file"]), spec["columns"],
            spec["widths"], spec["font"], spec.get("cell_margin", 50),
            spec.get("paragraph_after", 1.2),
        )
        replace_table_after_caption(document, caption, table)
        replace_caption(caption, spec["caption"])

    # Keep each caption attached to the table that follows it.
    for caption in [first_caption, *old_captions.values()]:
        p_pr = caption._p.get_or_add_pPr()
        if p_pr.find(qn("w:keepNext")) is None:
            p_pr.append(OxmlElement("w:keepNext"))


def add_component_grid_text(document, tables: Path) -> None:
    grid_rows = read_csv(tables / "component_selection_candidate_grids.csv")
    evidence = read_csv(tables / "component_selection_evidence.csv")
    grid_text = "; ".join(
        f"{row['dataset'].replace('_', ' ')}: {row['components']}"
        for row in grid_rows
    )
    nmr = [row for row in evidence if row["dataset"] == "nmr"]
    family_labels = {
        "plssvd": "PLS-SVD",
        "simpls": "SIMPLS-family",
        "opls": "OPLS",
        "kernel_pls": "kernel PLS",
    }
    decisions = "; ".join(
        f"{family_labels.get(row['family'], row['family'])} selected "
        f"{row['selected_ncomp']} from "
        f"eligible values {row['eligible_ncomp']}"
        for row in nmr
    )
    text = (
        "Component-selection grids. A denotes the retained component count, "
        "and randomized singular value decomposition is abbreviated rSVD. "
        "All component counts were selected within "
        "the following evaluated training-only grids: " + grid_text + ". "
        "For NMR, five paired 80/20 training-only splits used the smallest "
        "component count whose mean validation RMSD was within one standard "
        "error of the minimum; " + decisions + "."
    )
    if any(p.text.startswith("Component-selection grids.") for p in document.paragraphs):
        return
    figure_caption = paragraph_starting(document, "Figure S1.")
    picture = figure_caption._p.getprevious()
    paragraph = document.add_paragraph(text)
    picture.addprevious(paragraph._p)


def normalize_main_prose_and_references(document) -> None:
    """Apply wording and citation fixes that must survive final assembly."""
    for paragraph in document.paragraphs:
        original_text = paragraph.text
        text = original_text
        if text == "fastPLS: accelerated SIMPLS for high-dimensional biomedical data":
            text = (
                "fastPLS: accelerated SIMPLS-family computation for "
                "high-dimensional biomedical data"
            )
        if text.startswith("Background and objective: Partial least squares"):
            text = (
                "Background and objective: Partial least squares (PLS) is "
                "widely used to predict biomedical outcomes from correlated "
                "measurements. Its computational cost increases when models "
                "retain many components or predict thousands of responses. "
                "We developed fastPLS to reduce the time and storage required "
                "for sequential PLS computation."
            )
        if text.startswith(
            "fastPLS was benchmarked on every prepared task"
        ):
            text = (
                "fastPLS was benchmarked on every prepared task using the "
                "component count specified for that dataset and PLS family. "
                "Classification used the LDA prediction head in the main "
                "software comparison, whereas regression returned continuous "
                "predictions. Independent R and Python implementations were "
                "run on the same Intel Core i7-13700 workstation with 32 GiB "
                "RAM, the same training and test observations, and one "
                "effective CPU thread. Each program used float32 where its "
                "public interface supported it. Unneeded fitted responses, "
                "loadings and variance summaries were omitted when the "
                "implementation allowed this. Replication was method and "
                "scale specific: every non-ImageNet fastPLS row and every "
                "ordinary IKPLS or scikit-learn row used ten fresh processes; "
                "every independent R method-dataset pair used up to three "
                "bounded fresh-process attempts, including NMR and ImageNet. "
                "When the first attempt failed or reached the time or memory "
                "limit, repetitions two and three were not run and were "
                "recorded explicitly; the "
                "fastPLS ImageNet and Python NMR/ImageNet routes used one. "
                "Unsupported routes, resource failures and timeouts were "
                "retained rather than replaced by extrapolated values."
            )
        if text.startswith("Classification was evaluated using either"):
            text = (
                "Classification was evaluated using either the largest "
                "predicted class-indicator response, termed argmax, or linear "
                "discriminant analysis (LDA) fitted to the PLS scores. For a "
                "score matrix T with retained dimension a and G classes, let "
                "n(g) and μ(g) denote the training count and score mean for "
                "class g. The pooled covariance is "
                "Σ = [TᵀT - ∑ over g n(g)μ(g)μ(g)ᵀ] / "
                "max(1, n - G). The implementation solves "
                "(Σ + λI)w(g) = μ(g) by Cholesky factorization and triangular "
                "solves. It sets s = trace(Σ) / a, or s = 1 when this value "
                "is not finite and positive, and tries λ = ρs for "
                "ρ = 10⁻⁸, 10⁻⁶, 10⁻⁵, 10⁻⁴, 10⁻³ and 10⁻², "
                "increasing ρ only after a failed factorization. Prediction "
                "uses δ(g | t) = tᵀw(g) - 0.5 μ(g)ᵀw(g) + "
                "log[n(g) / n]. "
                "Regression directly predicts the continuous responses."
            )
        if text.startswith("Figure 1. Prediction, runtime and memory"):
            text = (
                "Figure 1. Prediction, runtime and memory for PLS software on "
                "an Intel Core i7-13700 workstation with one effective CPU "
                "thread. fastPLS uses its SIMPLS-family estimator with LDA for "
                "classification and continuous prediction for regression. "
                "A and B: test accuracy and RMSD. C and D: fitting and "
                "prediction time. E and F: absolute peak process RSS. Time and "
                "memory share color scales across classification and "
                "regression. Cells report medians of successful runs; "
                "replication counts and retained failures are specified in "
                "Methods and Supplementary Table S4."
            )
        if text.startswith("Figure 2. CPU and CUDA performance"):
            text = (
                "Figure 2. CPU and CUDA performance on the same Intel Core "
                "i7-13700/RTX 5060 Ti workstation using float32 inputs. A: CPU "
                "time divided by CUDA time for fitting and fixed-test "
                "prediction; values above 1 favour CUDA. B: CUDA host-RSS "
                "increment divided by the CPU increment; values below 1 favour "
                "CUDA. Panels A and B use argmax for classification. C and D: "
                "complete 10-fold cross-validation time divided by one "
                "full-training fit plus fixed-test prediction on Linux CPU and "
                "CUDA, respectively. Classification uses LDA in both the "
                "numerator and denominator, and the LDA head is refitted in "
                "each training fold; regression uses continuous predictions. "
                "Panels C and D retain the dataset-specific component counts, "
                "folds, rSVD controls and seeds. Five fresh processes were used "
                "except for ImageNet, which used one. TO denotes the 1,800-s "
                "execution limit. Timings include initialization, transfers and "
                "synchronization. Host-RSS increments include numerical "
                "libraries, allocator pools, staging buffers and model "
                "workspaces; CUDA device allocation is reported separately in "
                "Supplementary Tables S5-S6."
            )
        if text.startswith("Figure 3. NMR spectral prediction"):
            text = (
                "Figure 3. NMR spectral prediction with fastPLS and the "
                "deposited PLS-SVD implementation. fastPLS used float32 rSVD "
                "at the predefined reporting points of 100 PLS-SVD and 50 "
                "SIMPLS-family components; these counts provide a matched "
                "workflow comparison and the 100-component PLS-SVD point is "
                "not the value selected by the training-only one-standard-error "
                "rule. The deposited workflow used float64 PLS-SVD with 165 "
                "components (grey). A: median fitting and prediction time and "
                "interquartile range from three fresh processes; the "
                "pseudo-logarithmic axis includes zero. B: test RMSD over all "
                "28,355 response intensities. C: baseline-corrected peak host "
                "RSS and CUDA device allocation. D: per-spectrum RMSD. E and F: "
                "observed and predicted spectra nearest the median CPU "
                "SIMPLS-family error, over 12-0 and 1.7-0.5 ppm. Component "
                "selection and the complete held-out paths are reported in "
                "Figure S12."
            )
        if text.startswith("34.") and "PMID:" not in text:
            text = text.rstrip(".") + ". [PMID: 38504018]."
        if text.startswith("36.") and "PMID:" not in text:
            text = text.rstrip(".") + ". [PMID: 38778098]."
        revised = text.replace(
            "ten-fold cross-validation", "10-fold cross-validation"
        ).replace(
            "Ten-fold cross-validation", "10-fold cross-validation"
        ).replace(
            "Ten-fold validation", "10-fold cross-validation"
        )
        revised = revised.replace(
            "UNI and its UNI-2 successor, and Prov-GigaPath, produce",
            "UNI [34], its officially released UNI2-h successor [35], and "
            "Prov-GigaPath [36] produce",
        )
        if revised != original_text:
            replace_caption(paragraph, revised)

    cv_methods_text = (
        "The timing comparison used one fixed, previously selected component "
        "count for each dataset and PLS family; it measured cross-validation "
        "execution rather than repeating component selection. In the 10-fold "
        "workflow, the PLS model and the LDA prediction head were fitted anew "
        "from the nine training folds and used to predict the remaining fold. "
        "The comparator fitted the same PLS family and LDA head once on the "
        "complete training partition and predicted the fixed test partition. "
        "Both classification workflows retained the complete LDA discriminant-"
        "score matrix, and both regression workflows retained continuous "
        "predictions, so that their output contracts were matched. Data loading "
        "and initial conversion were excluded; fold construction, training, "
        "prediction and assembly of out-of-fold outputs were included. CUDA "
        "times included initialization, transfers and synchronization."
    )
    if not any(
        p.text.startswith("The timing comparison used one fixed")
        for p in document.paragraphs
    ):
        algorithm_caption = paragraph_starting(
            document,
            "Algorithm 3. Leakage-free sufficient-statistics cross-validation.",
        )
        paragraph = document.add_paragraph(cv_methods_text)
        algorithm_caption._p.addprevious(paragraph._p)

    formal_discussion = (
        "Formal verification complements these numerical checks by "
        "establishing the exact-real identities used for implicit products, "
        "compact prediction, deflation, kernel centring and leakage-free "
        "subtraction of additive fold statistics. This reduces the risk of "
        "an algebraic transcription error in the optimized derivations, but "
        "it does not verify floating-point kernels, randomized approximation "
        "or end-to-end refinement of the C++ program; numerical and "
        "platform-specific tests therefore remain necessary."
    )
    if not any(
        p.text.startswith("Formal verification complements")
        for p in document.paragraphs
    ):
        anchor = paragraph_starting(
            document, "The numerical results also clarify"
        )
        paragraph = document.add_paragraph(formal_discussion)
        insert_after(anchor._p, paragraph._p)

    simpls_step_1 = (
        "Input requested component-count set C and let A = max(C). Centre or "
        "scale X with training statistics and centre Y. Form the initial "
        "cross-covariance S0 = XᵀY, or define operators that apply S0 and "
        "S0ᵀ without storing the p by q matrix. For classification, construct "
        "the required products from compact class labels rather than a dense "
        "one-hot Y."
    )
    simpls_step_8 = (
        "Let e be the number of effective directions. For each requested "
        "prefix a in C, use a* = min(a,e). When e > 0, predict with "
        "(Xnew R[a*])Q[a*]ᵀ after applying training preprocessing, then "
        "restore the training-response mean; prefixes above e repeat the "
        "last estimable prediction and coefficient path. When e = 0, retain "
        "every requested position and return the training-response mean for "
        "regression or the training class prior for classification. Construct "
        "fitted values and metrics only when requested."
    )
    simpls_step_6 = (
        "Orthogonalize p against the retained columns of V and normalize the "
        "result to obtain v. Compute h = vᵀS once and apply the rank-one "
        "deflation S <- S - vh; when G is cached, update G <- G - hᵀh and "
        "restore numerical symmetry."
    )
    for table in document.tables:
        if (
            len(table.rows) == 9
            and len(table.columns) == 2
            and table.rows[0].cells[0].text.strip() == "Step"
            and "initial cross-covariance" in table.rows[1].cells[1].text
        ):
            replace_caption(table.rows[1].cells[1].paragraphs[0], simpls_step_1)
            replace_caption(table.rows[6].cells[1].paragraphs[0], simpls_step_6)
            replace_caption(table.rows[8].cells[1].paragraphs[0], simpls_step_8)
            break

    plssvd_replacements = {
        1: (
            "Input requested component-count set C and let A = max(C). Centre "
            "or scale X and centre Y. Form S = XᵀY, or define the implicit "
            "products SΩ = Xᵀ(YΩ) and SᵀZ = Yᵀ(XZ). For classification, "
            "construct S from compact class labels without materializing a "
            "one-hot Y."
        ),
        2: (
            "Compute one rank-e randomized decomposition S ≈ UDVᵀ, where e "
            "is the smaller of A and the effective numerical rank. Retain U, V "
            "and the singular values in D."
        ),
        4: (
            "For each requested a in C, let a* = min(a,e) and solve "
            "H[1:a*,1:a*] Lₐ = D[1:a*,1:a*] by Cholesky factorization, with a "
            "general linear solve as the numerical fallback; do not invert H "
            "explicitly. If e = 0, retain the intercept-only path."
        ),
    }
    for table in document.tables:
        if (
            len(table.rows) == 7
            and len(table.columns) == 2
            and table.rows[0].cells[0].text.strip() == "Step"
            and "implicit products" in table.rows[1].cells[1].text
        ):
            for row_index, text in plssvd_replacements.items():
                replace_caption(
                    table.rows[row_index].cells[1].paragraphs[0], text
                )
            break

    cv_steps = (
        "Input X, response Y or compact class labels z, fixed K-fold map f, "
        "requested component-count set C, PLS family, scaling rule, prediction "
        "head, backend and seed. Validate dimensions and keep each user-"
        "supplied group within one fold.",
        "Before the fold loop, compute only eligible additive full-data "
        "statistics: predictor sums sx, squared sums qx and XᵀX; for "
        "regression, response sums sy and, when used, XᵀY or YYᵀ; for "
        "classification, class counts nc and class-wise predictor sums mc. Do "
        "not centre or scale with full-data means.",
        "For fold h = 1,...,K, let I_hold,h = {i : f(i) = h} and let "
        "I_train,h contain all remaining indices. Every fitted quantity below "
        "uses I_train,h; I_hold,h is used only to subtract additive "
        "contributions and to obtain held-out predictions.",
        "Recover training marginals by subtraction: sx,h = sx - "
        "∑(i ∈ I_hold,h) xᵢ and qx,h = qx - ∑(i ∈ I_hold,h) xᵢ²; "
        "for regression, sy,h = sy - ∑(i ∈ I_hold,h) yᵢ. For "
        "classification, subtract the "
        "holdout class counts and class-wise predictor sums.",
        "Compute predictor centres and scales and the response mean only from "
        "the recovered training statistics. Apply these training-fold values "
        "unchanged to I_hold,h. A class absent from I_train,h is not fitted or "
        "predicted, but its held-out observations remain in the evaluation.",
        "When cached, recover the raw training cross-product as full XᵀY "
        "minus X_holdᵀY_hold and the predictor Gram matrix as full XᵀX "
        "minus X_holdᵀX_hold; then centre and scale them with training-fold "
        "statistics. For labels, form the centred class cross-product from "
        "training class counts and class sums without a one-hot matrix.",
        "For an eligible wide response, extract YYᵀ[I_train,h,I_train,h]. "
        "Obtain each training row sum by subtracting entries against I_hold,h "
        "from the cached full row sum, and double-centre the principal "
        "submatrix using only training-fold quantities. Retain one triangle "
        "until a full matrix is required.",
        "Fit one maximal PLS-SVD, SIMPLS-family, OPLS or linear kernel PLS "
        "component path on I_train,h with seed + h and evaluate every requested "
        "a in C. Fit the OPLS filter inside I_train,h. For nonlinear kernel PLS, "
        "construct and centre the training Gram matrix and holdout-to-training "
        "cross-kernel within the fold; do not use the linear sufficient-"
        "statistics shortcut.",
        "Let eh be the number of effective directions in fold h. Evaluate each "
        "requested a with min(a,eh), repeating the last estimable prediction "
        "when a > eh. If eh = 0, use the training-response mean for regression "
        "or the training class prior for classification; a single-class "
        "training fold predicts that class. Record requested and effective "
        "counts and fallback status.",
        "Predict I_hold,h after applying the training-fold preprocessing. "
        "Argmax uses reconstructed response scores; LDA is fitted only from "
        "training scores or algebraically equivalent training-only moments. "
        "Write held-out predictions back to their original row positions.",
        "For single cross-validation, compute every candidate metric from the "
        "complete out-of-fold predictions and select within C using the "
        "declared rule. For nested cross-validation, perform Steps 2-10 on "
        "inner folds formed only within each outer-training partition, select "
        "within C, refit on that complete outer-training partition, and predict "
        "its outer holdout. Aggregate the outer predictions only after every "
        "outer fold is complete.",
    )
    for table in document.tables:
        if (
            len(table.rows) == 12
            and len(table.columns) == 2
            and table.rows[0].cells[0].text.strip() == "Step"
            and table.rows[1].cells[1].text.startswith("Input X, response Y")
        ):
            for row_index, text in enumerate(cv_steps, start=1):
                replace_caption(
                    table.rows[row_index].cells[1].paragraphs[0], text
                )
            break
    else:
        raise RuntimeError("Could not locate the cross-validation pseudocode table")

    if any(p.text.startswith("34.") for p in document.paragraphs):
        return
    reference_33 = paragraph_starting(document, "33.")
    entries = (
        "34.\tChen, R.J., et al., Towards a general-purpose foundation model "
        "for computational pathology. Nature Medicine, 2024. 30: p. 850-862. "
        "https://doi.org/10.1038/s41591-024-02857-3. [PMID: 38504018].",
        "35.\tMahmood Lab, UNI2-h: Pathology Foundation Model. Hugging Face "
        "model repository, 2025. https://huggingface.co/MahmoodLab/UNI2-h "
        "(accessed 19 September 2026).",
        "36.\tXu, H., et al., A whole-slide foundation model for digital "
        "pathology from real-world data. Nature, 2024. 630: p. 181-188. "
        "https://doi.org/10.1038/s41586-024-07441-w. [PMID: 38778098].",
    )
    anchor = reference_33._p
    for entry in entries:
        paragraph = document.add_paragraph(entry)
        insert_after(anchor, paragraph._p)
        anchor = paragraph._p


def rewrite_quantitative_narrative(document, facts: dict) -> None:
    """Replace quantitative prose with text derived from audited result files."""
    validation = facts["numerical_validation"]
    dense_rows = [
        row for row in validation
        if row.get("validation_block") == "Dense numerical reference"
    ]
    rsvd_rows = [
        row for row in validation
        if row.get("validation_block") == "rSVD qualification"
    ]
    estimator_rows = [
        row for row in validation
        if row.get("validation_block") == "Independent estimator reference"
    ]
    precision_rows = [
        row for row in validation
        if row.get("validation_block") == "float32 versus float64"
    ]
    if not (
        len(dense_rows) == 1 and len(rsvd_rows) == 2
        and len(estimator_rows) == 3 and len(precision_rows) == 2
    ):
        raise RuntimeError("The numerical-validation narrative contract changed")
    dense_attempted = sum(int(row["attempted"]) for row in dense_rows)
    rsvd_attempted = sum(int(row["attempted"]) for row in rsvd_rows)
    estimator_attempted = sum(int(row["attempted"]) for row in estimator_rows)
    precision_attempted = sum(int(row["attempted"]) for row in precision_rows)
    if any(int(row["failures"]) for row in validation):
        raise RuntimeError("A current numerical-validation block contains failures")
    if any(
        not re.match(r"^\d+/\d+ met tolerances", row["key_numerical_result"])
        for row in rsvd_rows + estimator_rows
    ):
        raise RuntimeError("Tolerance counts are missing from validation evidence")
    replace_started(
        document,
        "Numerical validation of the evaluated implementation",
        "Numerical validation preceded interpretation of the performance "
        f"experiments. The SIMPLS-family dense-reference panel completed "
        f"{dense_attempted} component-prefix comparisons. rSVD met the stated "
        f"tolerances in all {rsvd_attempted} comparisons: 174 on CPU and 174 "
        f"on CUDA. All {estimator_attempted} independent OPLS and nonlinear "
        f"kernel PLS cases also met their stated tolerances, and all "
        f"{precision_attempted} paired float32-versus-float64 routes completed. "
        "Supplementary Table S1 reports the prediction, label and metric "
        "differences and distinguishes execution completion from numerical "
        "tolerance attainment.",
    )
    figure1 = facts["figure1"]
    datasets = figure1["per_dataset"]
    image = datasets["imagenet"]
    nmr = datasets["nmr"]
    tabula = datasets["tabula"]
    image_fast = image["fastpls"]
    image_ikpls = image["ikpls"]
    nmr_fast = nmr["fastpls"]
    tabula_fast = tabula["fastpls"]
    tabula_ikpls = tabula["ikpls"]

    classification = [
        dataset for dataset, values in datasets.items()
        if values["fastpls"]
        and values["fastpls"].get("task_type") == "classification"
        and dataset != "imagenet"
    ]
    fastest = set(figure1["fastpls_fastest_completed_r_datasets"])
    fastest_classification = sum(dataset in fastest for dataset in classification)
    image_ratio = (
        finite(image_ikpls["total_sec"]) / finite(image_fast["total_sec"])
    )

    nmr_rows = facts["figure3_nmr"]

    def implementation_row(fragment: str):
        matches = [
            row for row in nmr_rows
            if fragment in row.get("implementation", "")
        ]
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one NMR implementation containing {fragment!r}; "
                f"found {len(matches)}"
            )
        return matches[0]

    deposited = implementation_row("Deposited PLS-SVD")
    nmr_plssvd_cpu = implementation_row("PLS-SVD\nCPU")
    nmr_plssvd_cuda = implementation_row("PLS-SVD\nCUDA")
    nmr_simpls_cpu = implementation_row("SIMPLS-family\nCPU")
    nmr_simpls_cuda = implementation_row("SIMPLS-family\nCUDA")
    deposited_time = finite(deposited["total_time_sec"])
    simpls_cuda_time = finite(nmr_simpls_cuda["total_time_sec"])
    plssvd_cuda_time = finite(nmr_plssvd_cuda["total_time_sec"])
    simpls_speed = deposited_time / simpls_cuda_time
    plssvd_speed = deposited_time / plssvd_cuda_time

    abstract = (
        f"Results: fastPLS had the lowest median fitting-and-prediction time "
        f"among completed R workflows on {fastest_classification} of "
        f"{len(classification)} classification datasets. The dense-reference "
        f"panel completed {dense_attempted} component-prefix comparisons, and "
        f"rSVD met the stated tolerances in all {rsvd_attempted} CPU/CUDA "
        f"comparisons. On ImageNet, with "
        f"1,000,000 training and 281,167 test embeddings, fastPLS completed in "
        f"{show(image_fast['total_sec'], 2)} s versus "
        f"{show(image_ikpls['total_sec'], 2)} s for IKPLS "
        f"({image_ratio:.2f}-fold faster). For NMR reconstruction, fastPLS "
        f"mapped 13,000 predictor bins to 28,355 response intensities from "
        f"1,200 training spectra and was the only compared implementation to "
        f"complete the matched large multivariate task. CUDA SIMPLS-family "
        f"execution required {simpls_cuda_time:.3f} s with a test RMSD of "
        f"{finite(nmr_simpls_cuda['RMSD']):.7f}, compared with "
        f"{deposited_time:.3f} s and RMSD "
        f"{finite(deposited['RMSD']):.7f} for the deposited workflow "
        f"({simpls_speed:.0f}-fold faster)."
    )
    replace_abstract_results(document, abstract)

    replace_started(
        document,
        "fastPLS had the lowest median fitting and prediction time",
        f"fastPLS had the lowest median fitting-and-prediction time among "
        f"completed R workflows on {fastest_classification} of the "
        f"{len(classification)} classification datasets (Figure 1). It also "
        "completed every regression task. The component-path analyses are "
        "reported in Supplementary Figures S1-S12. At the two largest scales, "
        "fastPLS was the only R implementation to complete: the remaining R "
        "workflows reached the predefined timeout or memory limit on ImageNet, "
        "NMR, or both.",
    )
    tabula_gain = (
        finite(tabula_fast["accuracy"]) - finite(tabula_ikpls["accuracy"])
    ) * 100
    balanced_gain = (
        finite(tabula_fast["balanced_accuracy"])
        - finite(tabula_ikpls["balanced_accuracy"])
    ) * 100
    replace_started(
        document,
        "The classification results also show",
        "The classification results also show the value of applying LDA after "
        "PLS score extraction. On Tabula Muris, fastPLS attained "
        f"{percent(tabula_fast['accuracy'])} accuracy and "
        f"{percent(tabula_fast['balanced_accuracy'])} balanced accuracy, "
        f"compared with {percent(tabula_ikpls['accuracy'])} and "
        f"{percent(tabula_ikpls['balanced_accuracy'])} for IKPLS with argmax. "
        f"The differences were {tabula_gain:.2f} and {balanced_gain:.2f} "
        "percentage points, respectively. Because the programs also differ in "
        "their PLS estimators and output policies, this comparison does not "
        "isolate LDA as the sole cause; it nevertheless demonstrates that the "
        "classifier fitted in latent score space can materially affect prediction.",
    )
    replace_started(
        document,
        "IKPLS and fastPLS had comparable practical runtimes",
        "IKPLS and fastPLS had comparable practical runtimes on many small and "
        "medium datasets, where interface calls, allocation and result "
        "construction represented a larger fraction of elapsed time. The "
        "difference became substantial for the million-sample ImageNet task. "
        f"At 1,000 components, fastPLS completed fitting and prediction in "
        f"{show(image_fast['total_sec'], 3)} s, compared with "
        f"{show(image_ikpls['total_sec'], 3)} s for IKPLS, a "
        f"{image_ratio:.2f}-fold advantage. Their top-1 accuracies were "
        f"{show(image_fast['accuracy'], 4)} and "
        f"{show(image_ikpls['accuracy'], 4)}, respectively; this predictive "
        "difference reflects both the estimator and the prediction head and is "
        "not an estimator-matched comparison.",
    )
    max_peak_dataset, max_peak = figure1["fastpls_max_absolute_peak_rss_mib"]
    replace_started(
        document,
        "The multivariate NMR problem exposed",
        "The multivariate NMR problem exposed a different scaling limit. "
        f"fastPLS predicted 28,355 response intensities from 13,000 predictor "
        f"bins in {show(nmr_fast['total_sec'], 3)} s, whereas IKPLS did not "
        "complete the requested calculation because its coefficient path "
        "exceeded available memory. fastPLS was therefore the only compared "
        "implementation to complete both NMR and ImageNet. Across the Figure 1 "
        f"workloads, its largest absolute process peak was {max_peak:.0f} MiB "
        f"on {max_peak_dataset.replace('_', ' ')}. These complete-process values "
        "include loaded matrices, runtimes and numerical libraries and should "
        "not be interpreted as isolated algorithmic workspace.",
    )

    figure2 = facts["figure2"]
    routes = figure2["cpu_cuda_routes"]
    paired = [
        row for row in routes
        if row.get("status_cpu") == row.get("status_accelerator") == "success"
        and finite(row.get("runtime_ratio")) is not None
    ]
    faster = [row for row in paired if finite(row["runtime_ratio"]) > 1]
    by_dataset = {}
    for row in paired:
        by_dataset.setdefault(row["dataset"], []).append(
            finite(row["runtime_ratio"])
        )
    accelerated_datasets = sorted({row["dataset"] for row in faster})
    ranges = "; ".join(
        f"{dataset.replace('_', ' ')} {min(by_dataset[dataset]):.2f}-"
        f"{max(by_dataset[dataset]):.2f}-fold"
        for dataset in accelerated_datasets
    )
    replace_started(
        document,
        "We tested CUDA backend in all",
        f"CUDA was evaluated in all {len(paired)} paired calculations, but was "
        "faster only when the workload was large enough to offset device "
        "initialization, transfer and synchronization (Figure 2A). It reduced "
        f"runtime in {len(faster)} pairs. The observed CPU/CUDA ranges among "
        f"datasets with at least one accelerated route were {ranges}. Routes "
        "below one remained CPU-favourable.",
    )

    cv_routes = figure2["cross_validation_routes"]
    cv_sentences = []
    cv_terminal_sentences = []
    for backend in ("cpu", "cuda"):
        values = [
            finite(row["cv_over_fit_predict"]) for row in cv_routes
            if row.get("backend") == backend
            and row.get("status_fit_predict") == row.get("status_cv") == "success"
            and finite(row.get("cv_over_fit_predict")) is not None
        ]
        over_ten = sum(value > 10 for value in values)
        cv_sentences.append(
            f"{backend.upper()} median {figure2['cv_over_fit_prediction'][backend]['median']:.2f}, "
            f"range {min(values):.2f}-{max(values):.2f}, with "
            f"{over_ten} of {len(values)} ratios above ten"
        )
        terminal = [
            row for row in cv_routes
            if row.get("backend") == backend
            and row.get("status_cv") in {"timeout", "error"}
        ]
        if terminal:
            timeout_count = sum(
                row.get("status_cv") == "timeout" for row in terminal
            )
            error_count = len(terminal) - timeout_count
            details = []
            if timeout_count:
                details.append(f"{timeout_count} reached the 1,800-s limit")
            if error_count:
                details.append(f"{error_count} failed")
            cv_terminal_sentences.append(
                f"On {backend.upper()}, " + " and ".join(details)
            )
    terminal_text = (
        " " + "; ".join(cv_terminal_sentences) + "."
        if cv_terminal_sentences else ""
    )
    replace_started(
        document,
        "Complete 10-fold cross-validation with LDA classification",
        "Complete 10-fold cross-validation with LDA classification was compared "
        "with one full-training fit plus fixed-test prediction (Figure 2C,D). "
        + "; ".join(cv_sentences)
        + "."
        + terminal_text
        + " The classifier and retained component "
        "count were held constant so that the comparison measured execution "
        "cost rather than a second component-selection procedure. LDA was still "
        "refitted within every training fold. Complete discriminant-score "
        "matrices were retained in both classification workflows, whereas the "
        "Figure 1 workflow returned only the outputs needed for its software "
        "comparison; their absolute times therefore have different output "
        "contracts. A ratio below one means that the fused compiled CV route "
        "completed faster than the score-heavy direct comparator; it does not "
        "mean that fewer than 10 folds were evaluated.",
    )
    replace_started(
        document,
        "CUDA increased the baseline-corrected host-RSS peak",
        f"CUDA increased the baseline-corrected host-RSS peak in "
        f"{figure2['cuda_host_memory_increase_count']} of the "
        f"{figure2['paired_routes']} paired calculations, with a median "
        f"CUDA/CPU ratio of {figure2['median_cuda_cpu_host_memory_ratio']:.2f}. "
        "On small tasks, context creation, numerical libraries, allocator pools "
        "and staging buffers dominated this increment. On the largest tasks, "
        "substantial storage moved from host memory to the CUDA device. These "
        "measurements describe complete runtime residency rather than isolated "
        "algorithmic workspace (Supplementary Tables S5-S6).",
    )

    selections = facts["component_selection"]
    nmr_plssvd_selection = row_for(selections, dataset="nmr", family="plssvd")
    nmr_simpls_selection = row_for(selections, dataset="nmr", family="simpls")
    replace_started(
        document,
        "The NMR application addresses",
        "The NMR application addresses a practical limitation of high-throughput "
        "metabolomics. Complementary NMR experiments emphasize different small- "
        "and macromolecular signals, and the earlier deposited workflow showed "
        "that PLS can reconstruct them from one acquired NOESY spectrum [7]. "
        "The present diffusion-edited task maps 13,000 NOESY bins to 28,355 "
        "intensities from 1,200 training spectra and is evaluated on 321 held-out "
        "spectra. Using five paired training-only splits, the one-standard-error "
        f"rule selected {nmr_plssvd_selection['selected_ncomp']} PLS-SVD and "
        f"{nmr_simpls_selection['selected_ncomp']} SIMPLS-family components "
        "within the evaluated grid. The smaller SIMPLS-family model therefore "
        "required fewer component updates; Figure S12 reports the complete "
        "held-out error, runtime and memory paths.",
    )
    replace_started(
        document,
        "The accelerated implementations improved",
        "The accelerated implementations improved both computation and "
        "prediction relative to the deposited 165-component PLS-SVD workflow. "
        "The main comparison retained the predefined reporting points of "
        f"{nmr_plssvd_cuda['ncomp']} PLS-SVD and "
        f"{nmr_simpls_cuda['ncomp']} SIMPLS-family components; the former is "
        "not the PLS-SVD value selected by the one-standard-error rule. "
        f"CUDA PLS-SVD required {plssvd_cuda_time:.3f} s at "
        f"{nmr_plssvd_cuda['ncomp']} components and CUDA SIMPLS-family execution "
        f"required {simpls_cuda_time:.3f} s at {nmr_simpls_cuda['ncomp']} "
        f"components, compared with {deposited_time:.3f} s for the deposited "
        f"calculation. The corresponding differences were {plssvd_speed:.0f}- "
        f"and {simpls_speed:.0f}-fold. Test RMSD was "
        f"{finite(nmr_plssvd_cuda['RMSD']):.7f} and "
        f"{finite(nmr_simpls_cuda['RMSD']):.7f}, compared with "
        f"{finite(deposited['RMSD']):.7f} for the deposited workflow (Figure 3).",
    )
    replace_started(
        document,
        "Memory use decreased substantially",
        "Memory use also decreased substantially. The deposited calculation "
        f"increased host RSS by {finite(deposited['incremental_rss_mib']):.0f} "
        f"MiB, compared with {finite(nmr_plssvd_cuda['incremental_rss_mib']):.0f} "
        f"MiB for CUDA PLS-SVD and "
        f"{finite(nmr_simpls_cuda['incremental_rss_mib']):.0f} MiB for CUDA "
        f"SIMPLS-family execution; their device peaks were "
        f"{finite(nmr_plssvd_cuda['gpu_peak_mib']):.0f} and "
        f"{finite(nmr_simpls_cuda['gpu_peak_mib']):.0f} MiB. CPU host-RSS "
        f"increments were {finite(nmr_plssvd_cpu['incremental_rss_mib']):.0f} "
        f"and {finite(nmr_simpls_cpu['incremental_rss_mib']):.0f} MiB. Because "
        "the deposited and accelerated workflows differ in family, solver, "
        "precision and component count, these are workflow-level rather than "
        "single-factor improvements.",
    )

    image_rows = facts["figure4_imagenet"]
    image_50_argmax = row_for(
        image_rows, method="simpls", classifier="argmax", ncomp_requested="50"
    )
    image_50_lda = row_for(
        image_rows, method="simpls", classifier="lda", ncomp_requested="50"
    )
    image_1000_argmax = row_for(
        image_rows, method="simpls", classifier="argmax", ncomp_requested="1000"
    )
    image_1000_lda = row_for(
        image_rows, method="simpls", classifier="lda", ncomp_requested="1000"
    )
    replace_started(
        document,
        "Foundation models produce dense embedding matrices",
        "Foundation models produce dense reusable embeddings that require a "
        "task-specific downstream model. PLS can construct a supervised compact "
        "representation of these correlated features, including representations "
        "from biomedical imaging models. We therefore used ImageNet/DINOv2 both "
        "as a million-sample feasibility test and to examine classification "
        "across SIMPLS-family component prefixes (Figure 4). At 50 components, "
        f"argmax attained {percent(image_50_argmax['top1_accuracy'])} top-1 and "
        f"{percent(image_50_argmax['top5_accuracy'])} top-5 accuracy, whereas "
        f"LDA attained {percent(image_50_lda['top1_accuracy'])} and "
        f"{percent(image_50_lda['top5_accuracy'])}. At 1,000 components, the "
        f"corresponding argmax values were "
        f"{percent(image_1000_argmax['top1_accuracy'])} and "
        f"{percent(image_1000_argmax['top5_accuracy'])}, and the LDA values were "
        f"{percent(image_1000_lda['top1_accuracy'])} and "
        f"{percent(image_1000_lda['top5_accuracy'])}. LDA therefore separated "
        "classes effectively with a short score representation, whereas argmax "
        "benefited more strongly from additional components. Nonlinear kernel "
        "PLS was not evaluated because a one-million-sample Gram matrix contains "
        "one trillion entries, requiring about 3.6 TiB in float32 before other "
        "workspaces. These exploratory results demonstrate feasibility rather "
        "than clinical performance; the split is nonstandard, extraction metadata "
        "are incomplete and 1,000 components is the evaluated boundary, not an "
        "optimum.",
    )

    replace_started(
        document,
        "The results show that fastPLS gains",
        "The results show that fastPLS gains its largest practical advantage "
        "when compact representation and output construction matter as much as "
        "the low-rank calculation. Small tasks were often separated from IKPLS "
        "by only milliseconds, but fastPLS completed the million-sample ImageNet "
        f"calculation {image_ratio:.2f} times faster and was the only compared "
        "implementation to complete the large multivariate NMR task. These gains "
        "are consistent with retaining compact predictor and response factors and "
        "constructing only requested outputs rather than dense coefficient paths.",
    )
    replace_started(
        document,
        "The numerical results also clarify",
        "The numerical results also clarify the scope of the accelerated "
        "estimators. rSVD is approximate and depends on the singular spectrum, "
        "seed, oversampling and power iterations. Likewise, bounded candidate "
        "blocks preserve sequential orthogonalization and deflation but do not "
        "recompute the leading direction after every accepted component; the "
        "method is therefore described as SIMPLS-family rather than classical "
        "de Jong SIMPLS. The current dense-reference, multi-seed, float32/float64 "
        "and backend checks met the stated numerical criteria, but do not "
        "establish equality for every matrix.",
    )
    replace_started(
        document,
        "CUDA was beneficial only after",
        f"CUDA was beneficial only after workloads became large enough to offset "
        f"initialization, transfer and synchronization. It reduced runtime in "
        f"{len(faster)} of {len(paired)} paired calculations, while increasing "
        f"the host-RSS increment in "
        f"{figure2['cuda_host_memory_increase_count']} pairs. The largest matrices "
        "instead moved substantial storage to device memory. The compiled CV "
        f"workflow required a median of "
        f"{figure2['cv_over_fit_prediction']['cpu']['median']:.2f} times one fit "
        f"and prediction on CPU and "
        f"{figure2['cv_over_fit_prediction']['cuda']['median']:.2f} times on "
        "CUDA. Ratios above ten on individual routes reflect route setup and a "
        "particularly fast direct-fit denominator rather than ten independent "
        "refits. Leakage-free reuse of additive statistics and component prefixes "
        "is especially relevant to methods such as KODAMA that repeatedly fit "
        "cross-validated PLS models [19,20].",
    )
    replace_started(
        document,
        "The NMR experiment provides",
        "The NMR experiment provides the clearest biomedical demonstration. "
        f"Training-only selection retained "
        f"{nmr_plssvd_selection['selected_ncomp']} PLS-SVD and "
        f"{nmr_simpls_selection['selected_ncomp']} SIMPLS-family components "
        "within the evaluated grid. The accelerated CUDA workflows were "
        f"{plssvd_speed:.0f}- and {simpls_speed:.0f}-fold faster than the "
        "deposited workflow, with lower global RMSD and substantially lower host "
        "memory. This changes the practical scope of spectral reconstruction by "
        "making component paths and repeated validation feasible. The comparison "
        "remains workflow-level because family, solver, precision and component "
        "count differ; localized and low-intensity errors also require scientific "
        "interpretation beyond the global metric.",
    )

    cuda_rows = [row for row in facts["cuda_software"] if row.get("status") == "success"]
    cuda_groups = {}
    for row in cuda_rows:
        cuda_groups.setdefault(row["dataset"], {})[row["implementation"]] = row
    cold_ratios = []
    warm_ratios = []
    host_memory_ratios = []
    for implementations in cuda_groups.values():
        fast = implementations.get("fastPLS_cuda")
        ikpls = implementations.get("IKPLS_jax_cuda_alg2")
        if not fast or not ikpls:
            continue
        cold_ratios.append(
            finite(ikpls["median_cold_total_sec"])
            / finite(fast["median_cold_total_sec"])
        )
        warm_ratios.append(
            finite(ikpls["median_warm_total_sec"])
            / finite(fast["median_warm_total_sec"])
        )
        host_memory_ratios.append(
            finite(ikpls["median_peak_rss_mib"])
            / finite(fast["median_peak_rss_mib"])
        )
    if cold_ratios and warm_ratios:
        replace_started(
            document,
            "Across the 11 datasets completed",
            f"Across the {len(cold_ratios)} datasets completed by both CUDA "
            "implementations, the median IKPLS/fastPLS first-use runtime ratio "
            f"was {median(cold_ratios):.2f}, with a range of "
            f"{min(cold_ratios):.2f}-{max(cold_ratios):.2f} (Figure S13). "
            "First-use timing includes process, framework and CUDA-context "
            "initialization. After those resources were initialized, the median "
            f"ratio was {median(warm_ratios):.2f}. IKPLS used a median "
            f"{median(host_memory_ratios):.2f}-fold more absolute host-process "
            "memory. These are software-workflow comparisons rather than "
            "estimator-kernel comparisons.",
        )
        replace_started(
            document,
            "These findings suggest that fastPLS benefits",
            "The CUDA software comparison also separates cold and warm timing. "
            "Cold timing includes process and accelerator initialization, whereas "
            "the repeated timing begins after the CUDA context has been created. "
            "Across matched successful tasks, the median IKPLS/fastPLS "
            f"ratio was {median(cold_ratios):.2f} for cold execution and "
            f"{median(warm_ratios):.2f} after initialization. This pattern "
            "indicates that compact factors, "
            "persistent workspaces and reduced framework overhead contribute in "
            "addition to GPU matrix multiplication.",
        )


def replace_supplement_figures(document, figures: Path) -> None:
    width = usable_width(document)
    for number in range(1, 14):
        image = figures / f"FigureS{number}.png"
        if not image.is_file():
            raise FileNotFoundError(image)
        caption = paragraph_starting(document, f"Figure S{number}.")
        replace_picture_before_caption(caption, image, width)
        if number == 12:
            replace_caption(
                caption,
                "Figure S12. NMR prediction and computational cost across "
                "component counts on CPU and CUDA. Each model was fitted on "
                "all 1,200 training spectra and evaluated on the fixed 321 "
                "test spectra. Rows show held-out RMSD, fitting and prediction "
                "time, and baseline-corrected peak host RSS; columns show the "
                "four PLS families. Dotted lines mark the training-only "
                "one-standard-error selections of 75 components for PLS-SVD "
                "and 50 for the SIMPLS-family estimator. OPLS and kernel PLS "
                "paths are shown "
                "without implying an independently selected component count."
            )

    image = figures / "FigureS14.png"
    if not image.is_file():
        raise FileNotFoundError(image)
    existing = [p for p in document.paragraphs if p.text.startswith("Figure S14.")]
    if existing:
        replace_picture_before_caption(existing[0], image, width)
        replace_caption(existing[0], FIGURE_S14_CAPTION)
        return
    last_caption = paragraph_starting(document, "Figure S13.")
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(image), width=width)
    caption = document.add_paragraph(FIGURE_S14_CAPTION)
    insert_after(last_caption._p, picture_paragraph._p)
    insert_after(picture_paragraph._p, caption._p)


def replace_method_grid_sentence(document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Classification differences were expressed"):
            text = paragraph.text
            text = re.sub(
                r"Dataset- and family-specific component counts were selected "
                r"within the evaluated candidate grids; the complete component "
                r"paths and retained counts were carried forward to the Results\.",
                "Dataset- and family-specific component counts were selected "
                "within explicit training-only candidate grids. The grids, "
                "selection rules, eligible NMR values and retained counts are "
                "reported in the Supplementary material; no selected count is "
                "described as a global optimum.",
                text,
            )
            replace_caption(paragraph, text)
            return
    raise RuntimeError("Could not locate the component-selection methods paragraph")


def correct_availability_cross_reference(document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Data and software availability:"):
            text = paragraph.text.replace(
                "Dataset retrieval and redistribution conditions are described "
                "in Supplementary Section S1.",
                "Dataset retrieval, preparation and redistribution conditions "
                "are described in Section 2.4.",
            )
            replace_caption(paragraph, text)
            return
    raise RuntimeError("Could not locate the data and software availability statement")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("main_input", type=Path)
    parser.add_argument("supplement_input", type=Path)
    parser.add_argument("assets", type=Path)
    parser.add_argument("output_directory", type=Path)
    args = parser.parse_args()
    figures = args.assets / "figures"
    tables = args.assets / "tables"
    output = args.output_directory.resolve()
    output.mkdir(parents=True, exist_ok=True)
    narrative_path = args.assets / "cmpb_narrative_summary.json"
    if not narrative_path.is_file():
        raise FileNotFoundError(narrative_path)
    narrative = json.loads(narrative_path.read_text())

    main = Document(args.main_input)
    normalize_shared_notation(main)
    replace_main_table(main, tables / "Table1_prepared_benchmark_dimensions.csv")
    replace_main_figures(main, figures)
    replace_method_grid_sentence(main)
    correct_availability_cross_reference(main)
    normalize_main_prose_and_references(main)
    rewrite_quantitative_narrative(main, narrative)
    main_output = output / "fastPLS_CMPB.docx"
    main.save(main_output)

    supplement = Document(args.supplement_input)
    normalize_shared_notation(supplement)
    normalize_supplement_algorithms(supplement)
    replace_supplement_tables(supplement, tables)
    add_component_grid_text(supplement, tables)
    replace_supplement_figures(supplement, figures)
    supplement_output = output / "fastPLS_CMPB_supplement.docx"
    supplement.save(supplement_output)

    print(main_output)
    print(supplement_output)


if __name__ == "__main__":
    main()
