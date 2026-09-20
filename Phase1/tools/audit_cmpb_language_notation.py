#!/usr/bin/env python3

import argparse
from pathlib import Path
import re
import sys

from docx import Document


FORBIDDEN_LITERAL = (
    "PLSSVD",
    "RSVD",
    "five-fold validation",
    "five-fold cross-validation",
    "ten-fold validation",
    "ten-fold cross-validation",
    "training-validation",
    "component k",
    "component count k",
    "Selected k",
    "Reference k",
    "k shown",
    "Predictive k",
    "n/p/q/k",
    "SIMPLS-rSVD",
    "SIMPLS family",
    "PLS-SVD/rSVD",
    "kernel-PLS",
    "candidate component set A",
    "Steps 1-11",
    "define T as all row indices",
    "n - C",
    "IRLBA",
    "warm-start",
    "Current-release",
    "current-release",
)


def text_blocks(document):
    for index, paragraph in enumerate(document.paragraphs):
        yield f"P{index}", paragraph.text
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                yield f"T{table_index}R{row_index}C{cell_index}", cell.text


def audit(path):
    document = Document(path)
    findings = []
    all_text = "\n".join(text for _, text in text_blocks(document))

    for location, text in text_blocks(document):
        for term in FORBIDDEN_LITERAL:
            if term in text:
                findings.append(f"{location}: forbidden term {term!r}: {text}")
        if re.search(r"\bk-\s*1\b", text) and "k-nearest" not in text:
            findings.append(f"{location}: component expression uses k-1: {text}")

    is_supplement = "supplement" in path.name.lower()
    is_cmpb_document = "cmpb" in path.name.lower()
    old_versions = sorted(set(re.findall(r"\b0\.99\.\d+\b", all_text)))
    if old_versions:
        findings.append(
            "obsolete package version labels remain: " + ", ".join(old_versions)
        )
    release_mentions = re.findall(
        r"fastPLS(?:\s+version)?\s+0\.3\b", all_text, flags=re.IGNORECASE
    )
    if is_cmpb_document and not is_supplement and release_mentions:
        findings.append("the package version must not appear in the main manuscript")
    if is_cmpb_document and is_supplement and len(release_mentions) != 1:
        findings.append(
            "the supplement must identify fastPLS 0.3 exactly once in its "
            "reproducibility section"
        )
    if is_cmpb_document:
        if "PLS-SVD" not in all_text:
            findings.append("document does not contain PLS-SVD")
        if "rSVD" not in all_text:
            findings.append("document does not contain rSVD")
    if is_cmpb_document:
        required_phrases = {
            "requested component-count set C": "requested component set C",
            "A = max(C)": "maximum retained component count A",
        }
        if not is_supplement:
            required_phrases.update({
                "G classes": "class count G",
                "K-fold map": "fold count K",
            })
        for phrase, description in required_phrases.items():
            if phrase not in all_text:
                findings.append(f"document does not define {description}")

    return findings


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("documents", nargs="+", type=Path)
    args = parser.parse_args()

    failed = False
    for path in args.documents:
        findings = audit(path)
        if findings:
            failed = True
            print(f"FAIL {path}")
            for finding in findings:
                print(f"  {finding}")
        else:
            print(f"PASS {path}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
