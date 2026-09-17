#!/usr/bin/env python3
"""Insert the matched CPU/GPU cross-validation figure into CMPB documents."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def insert_before(document: Document, target, text: str = "", style=None):
    paragraph = document.add_paragraph(text, style=style)
    target._p.addprevious(paragraph._p)
    return paragraph


def style_paragraph(paragraph, size: float = 9.0, italic: bool = False) -> None:
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.italic = italic


def cv_summary(summary_path: Path) -> dict[str, float | int]:
    frame = pd.read_csv(summary_path)
    frame = frame[
        (frame["status"] == "success")
        & frame["cpu_over_accelerator"].notna()
    ]
    output: dict[str, float | int] = {}
    for platform, prefix in (
        ("linux_nvidia", "cuda"),
        ("mac_apple_silicon", "metal"),
    ):
        part = frame[frame["platform"] == platform]
        output[f"{prefix}_completed"] = len(part)
        output[f"{prefix}_faster"] = int(
            (part["cpu_over_accelerator"] > 1.0).sum()
        )
        output[f"{prefix}_maximum"] = float(
            part["cpu_over_accelerator"].max()
        )
    return output


def update_manuscript(
    source: Path, destination: Path, stats: dict[str, float | int]
) -> None:
    document = Document(source)
    target = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Verified four-thread CPU execution")
    )
    text = (
        "Complete ten-fold cross-validation showed the same dependence on "
        "workload size. CUDA was faster than the same-host CPU route in "
        f"{stats['cuda_faster']} of {stats['cuda_completed']} completed "
        "family-dataset comparisons, with a maximum CPU/CUDA runtime ratio "
        f"of {stats['cuda_maximum']:.1f}. Metal was faster in "
        f"{stats['metal_faster']} of {stats['metal_completed']} matched "
        "comparisons on the Apple workstation; most Metal ratios remained "
        "near one because the hybrid route retains CPU-dependent stages. "
        "The complete paired matrix is shown in Supplementary Figure S20."
    )
    paragraph = insert_before(document, target, text)
    style_paragraph(paragraph, 10.0)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def update_supplement(
    source: Path,
    destination: Path,
    figure_path: Path,
    stats: dict[str, float | int],
) -> None:
    document = Document(source)
    headings = {paragraph.text: paragraph for paragraph in document.paragraphs}
    target = headings["S10. ImageNet stress test"]

    heading = insert_before(
        document,
        target,
        "S10. Cross-validation runtime across CPU and GPU backends",
        style="Heading 1",
    )
    style_paragraph(heading, 13.0)
    heading.paragraph_format.keep_with_next = True

    explanation = insert_before(
        document,
        target,
        (
            "Complete ten-fold cross-validation was measured with float32 "
            "inputs, fixed folds, seed 123 and the training-selected component "
            "counts in Table S11. Each cell compares CPU and accelerator "
            "execution on the same host. The timed public call includes fold "
            "preparation, fitting, out-of-fold prediction, metric reduction "
            "and construction of the documented return object. CUDA timing "
            "also includes context initialization, data transfer and "
            "synchronization; Metal timing includes command submission and "
            "synchronization. Values above one favour the GPU route. CUDA was "
            f"faster in {stats['cuda_faster']} of {stats['cuda_completed']} "
            "completed comparisons, whereas Metal was faster in "
            f"{stats['metal_faster']} of {stats['metal_completed']}. The "
            "ImageNet public CV workload was not evaluated because retaining "
            "the 1,000-component out-of-fold score object exceeded the memory "
            "budget of the benchmark systems."
        ),
    )
    style_paragraph(explanation, 9.5)
    explanation.paragraph_format.keep_with_next = True

    picture = insert_before(document, target)
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(figure_path), width=Inches(6.45))

    caption = insert_before(
        document,
        target,
        (
            "Figure S20. Same-host CPU-to-GPU runtime ratio for complete "
            "ten-fold cross-validation. Panel A compares Linux CPU with "
            "NVIDIA CUDA; Panel B compares Mac CPU with the hybrid Metal "
            "route. PLS-SVD, SIMPLS, OPLS and linear kernel PLS used matched "
            "float32 inputs, fixed folds, seed 123 and fixed component counts. "
            "Classification used argmax and regression used continuous "
            "prediction. Cell values are CPU CV time divided by GPU CV time: "
            "values above one indicate faster GPU execution and values below "
            "one indicate faster CPU execution. Points summarize five fresh "
            "processes per evaluated cell; accelerator initialization, "
            "transfer and synchronization are included. NE denotes not "
            "evaluated."
        ),
    )
    style_paragraph(caption, 9.5)

    headings["S10. ImageNet stress test"].text = "S11. ImageNet stress test"
    headings[
        "S11. Reproducibility and software checks"
    ].text = "S12. Reproducibility and software checks"

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manuscript", type=Path, required=True)
    parser.add_argument("--supplement", type=Path, required=True)
    parser.add_argument("--figure", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--manuscript-out", type=Path, required=True)
    parser.add_argument("--supplement-out", type=Path, required=True)
    args = parser.parse_args()

    stats = cv_summary(args.summary)
    update_manuscript(args.manuscript, args.manuscript_out, stats)
    update_supplement(
        args.supplement,
        args.supplement_out,
        args.figure,
        stats,
    )


if __name__ == "__main__":
    main()
