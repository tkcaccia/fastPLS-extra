#!/usr/bin/env python3
"""Apply the Leonardo revision to the CMPB manuscript and supplement.

The script accepts tracked changes, removes review comments, replaces the five
pseudocode tables with implementation-faithful algorithms using native Word
equations, and applies the requested editorial corrections.
"""

from __future__ import annotations

import argparse
import copy
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W_NS}
MATH_XSL = Path(
    "/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl"
)


def accept_revisions(source: Path, destination: Path) -> None:
    """Accept tracked content and formatting changes in a DOCX."""
    work = Path(tempfile.mkdtemp(prefix="fastpls-docx-accept-"))
    try:
        with zipfile.ZipFile(source) as archive:
            archive.extractall(work)
        for xml_path in (work / "word").rglob("*.xml"):
            try:
                tree = etree.parse(str(xml_path))
            except etree.XMLSyntaxError:
                continue
            changed = False
            for tag in ("del", "moveFrom"):
                for node in tree.xpath(f"//w:{tag}", namespaces=NS):
                    node.getparent().remove(node)
                    changed = True
            for tag in ("ins", "moveTo"):
                for node in tree.xpath(f"//w:{tag}", namespaces=NS):
                    parent = node.getparent()
                    position = parent.index(node)
                    for child in list(node):
                        parent.insert(position, child)
                        position += 1
                    parent.remove(node)
                    changed = True
            # Property-change elements retain the previous formatting after a
            # tracked change is accepted.  Removing the history node preserves
            # the current paragraph, run, table, row, cell and section style.
            for tag in (
                "rPrChange",
                "pPrChange",
                "tblPrChange",
                "tblPrExChange",
                "tblGridChange",
                "trPrChange",
                "tcPrChange",
                "sectPrChange",
                "numberingChange",
            ):
                for node in tree.xpath(f"//w:{tag}", namespaces=NS):
                    node.getparent().remove(node)
                    changed = True
            if changed:
                tree.write(
                    str(xml_path),
                    encoding="UTF-8",
                    xml_declaration=True,
                    standalone=True,
                )
        archive_base = destination.with_suffix("")
        shutil.make_archive(str(archive_base), "zip", work)
        archive_path = archive_base.with_suffix(".zip")
        destination.parent.mkdir(parents=True, exist_ok=True)
        os.replace(archive_path, destination)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def remove_review_markup(path: Path) -> None:
    """Remove comment anchors, comment parts and comment relationships."""
    work = Path(tempfile.mkdtemp(prefix="fastpls-docx-clean-"))
    try:
        with zipfile.ZipFile(path) as archive:
            archive.extractall(work)

        word_dir = work / "word"
        for xml_path in word_dir.rglob("*.xml"):
            try:
                tree = etree.parse(str(xml_path))
            except etree.XMLSyntaxError:
                continue
            changed = False
            for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                for node in tree.xpath(f"//w:{tag}", namespaces=NS):
                    node.getparent().remove(node)
                    changed = True
            if changed:
                tree.write(
                    str(xml_path),
                    encoding="UTF-8",
                    xml_declaration=True,
                    standalone=True,
                )

        rels = word_dir / "_rels" / "document.xml.rels"
        if rels.exists():
            tree = etree.parse(str(rels))
            root = tree.getroot()
            for node in list(root):
                relation_type = node.get("Type", "")
                target = node.get("Target", "")
                if "comment" in relation_type.lower() or "people" in target.lower():
                    root.remove(node)
            tree.write(
                str(rels),
                encoding="UTF-8",
                xml_declaration=True,
                standalone=True,
            )

        content_types = work / "[Content_Types].xml"
        tree = etree.parse(str(content_types))
        root = tree.getroot()
        for node in list(root):
            part_name = node.get("PartName", "").lower()
            if "comment" in part_name or part_name.endswith("/people.xml"):
                root.remove(node)
        tree.write(
            str(content_types),
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )

        for pattern in ("comments*.xml", "people.xml"):
            for stale in word_dir.glob(pattern):
                stale.unlink()

        archive_base = path.with_suffix("")
        shutil.make_archive(str(archive_base), "zip", work)
        archive_path = archive_base.with_suffix(".zip")
        os.replace(archive_path, path)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def clear_paragraph(paragraph) -> None:
    element = paragraph._element
    properties = element.pPr
    for child in list(element):
        if child is not properties:
            element.remove(child)


def set_paragraph_text(paragraph, text: str, *, bold_prefix: str | None = None) -> None:
    clear_paragraph(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(f"paragraph not found: {prefix}")


def set_by_prefix(document: Document, prefix: str, text: str) -> None:
    paragraph = find_paragraph(document, prefix)
    set_paragraph_text(paragraph, text)


def replace_literal_in_paragraph(
    document: Document,
    prefix: str,
    source: str,
    target: str,
) -> None:
    """Replace visible run text without rebuilding EndNote field structures."""
    paragraph = find_paragraph(document, prefix)
    replaced = False
    for run in paragraph.runs:
        if source in run.text:
            run.text = run.text.replace(source, target)
            replaced = True
    if not replaced:
        raise KeyError(f"text {source!r} not found in paragraph {prefix!r}")


def replace_in_runs(
    document: Document,
    replacements: dict[str, str],
    *,
    preserve_references: bool = False,
) -> None:
    """Apply spelling corrections without destroying existing Word equations."""
    body_paragraphs = document.paragraphs
    if preserve_references:
        reference_index = next(
            (
                index
                for index, paragraph in enumerate(body_paragraphs)
                if paragraph.text.strip() == "References"
            ),
            len(body_paragraphs),
        )
        body_paragraphs = body_paragraphs[:reference_index]
    containers = [body_paragraphs]
    for table in document.tables:
        containers.extend(cell.paragraphs for row in table.rows for cell in row.cells)
    for paragraphs in containers:
        for paragraph in paragraphs:
            for run in paragraph.runs:
                original = run.text
                text = original
                for source, target in replacements.items():
                    text = re.sub(rf"\b{re.escape(source)}\b", target, text)
                    text = re.sub(
                        rf"\b{re.escape(source.capitalize())}\b",
                        target.capitalize(),
                        text,
                    )
                if text != original:
                    run.text = text


def omml_from_latex(latex: str):
    mathml = etree.fromstring(latex_to_mathml(latex).encode("utf-8"))
    transform = etree.XSLT(etree.parse(str(MATH_XSL)))
    result = transform(mathml)
    root = result.getroot()
    return copy.deepcopy(root)


def add_text_line(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.2)


def add_equation_line(cell, latex: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph._element.append(omml_from_latex(latex))


def add_algorithm_items_inline(cell, items: list[tuple[str, str]]) -> None:
    """Write one compact operation paragraph with inline native equations."""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for index, (kind, value) in enumerate(items):
        if index:
            spacer = paragraph.add_run("  ")
            spacer.font.size = Pt(8.8)
        if kind == "text":
            run = paragraph.add_run(value)
            run.font.size = Pt(8.8)
        elif kind == "math":
            paragraph._element.append(omml_from_latex(value))
        else:
            raise ValueError(kind)


def clear_cell(cell) -> None:
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)
    cell.add_paragraph()


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def rebuild_algorithm_table(table, rows: list[tuple[str, list[tuple[str, str]]]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    table.autofit = False
    grid_columns = table._tbl.tblGrid.gridCol_lst
    if len(grid_columns) >= 2:
        grid_columns[0].set(qn("w:w"), "1050")
        grid_columns[1].set(qn("w:w"), "8600")
    header = table.rows[0]
    header.cells[0].width = Inches(0.72)
    header.cells[1].width = Inches(5.98)
    for cell, label in zip(header.cells, ("Step", "Operation")):
        clear_cell(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(9.2)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=45, start=70, bottom=45, end=70)
    set_repeat_table_header(header)
    prevent_row_split(header)

    for label, items in rows:
        row = table.add_row()
        row.cells[0].width = Inches(0.72)
        row.cells[1].width = Inches(5.98)
        clear_cell(row.cells[0])
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        run.bold = label in {"Input", "Output"}
        run.font.size = Pt(8.8)
        clear_cell(row.cells[1])
        add_algorithm_items_inline(row.cells[1], items)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=40, start=65, bottom=40, end=65)
        prevent_row_split(row)


SIMPLS_ROWS = [
    (
        "Input",
        [
            ("text", "Training predictors and responses or compact class labels; requested component set and its maximum; bounded candidate-block size."),
            ("math", r"X\in\mathbb{R}^{n\times p},\quad Y_c\in\mathbb{R}^{n\times q},\quad C,\quad A=\max(C),\quad b"),
        ],
    ),
    (
        "1",
        [
            ("text", "Estimate centring and scaling from the training data. Form the cross-covariance explicitly or expose only its matrix products; compact labels replace a dense one-hot response."),
            ("math", r"S_0=X_c^{\mathsf T}Y_c"),
        ],
    ),
    (
        "2",
        [
            ("text", "Initialise the compact weight, loading and deflation bases. Cache a predictor moment matrix or the smaller right Gram matrix only when the route heuristic selects it."),
            ("math", r"R=Q=V=\varnothing,\qquad M=X_c^{\mathsf T}X_c,\qquad G=S^{\mathsf T}S"),
        ],
    ),
    (
        "3",
        [
            ("text", "At each refresh, obtain at most b leading left candidates from the current deflated state by rSVD or a small eigendecomposition. Candidates in one block deliberately share that state."),
        ],
    ),
    (
        "4",
        [
            ("text", "For each candidate direction, compute and normalise the score, then calculate predictor and response loadings."),
            ("math", r"t=X_cr,\quad t^{\mathsf T}t=1,\quad p=X_c^{\mathsf T}t\;(\mathrm{or}\;Mr),\quad q=S_0^{\mathsf T}r=Y_c^{\mathsf T}t"),
        ],
    ),
    (
        "5",
        [
            ("text", "Orthogonalise the predictor loading against the retained deflation basis and apply one rank-one update to the cross-covariance and, when present, its cached Gram matrix."),
            ("math", r"v=\frac{(I-VV^{\mathsf T})p}{\lVert(I-VV^{\mathsf T})p\rVert_2},\quad h=v^{\mathsf T}S,\quad S\leftarrow S-vh,\quad G\leftarrow G-h^{\mathsf T}h"),
        ],
    ),
    (
        "6",
        [
            ("text", "Append the effective direction and continue through the current block. Resketch only after that block has been consumed, stopping at A or at numerical rank e."),
        ],
    ),
    (
        "7",
        [
            ("text", "For each requested prefix, use the available directions. Prefixes above rank repeat the last estimable path; a zero-direction regression path returns the training-response mean and classification uses the training prior."),
            ("math", r"c^{\ast}=\min(c,e),\qquad \widehat{Y}_c=(X_{\mathrm{new},c}R_{c^{\ast}})Q_{c^{\ast}}^{\mathsf T}+\mathbf{1}\bar y^{\mathsf T}"),
        ],
    ),
    (
        "Output",
        [
            ("text", "Compact predictor weights R, response loadings Q, deflation basis V, optional scores T, effective rank e, diagnostics and the complete requested prediction path."),
        ],
    ),
]


PLSSVD_ROWS = [
    (
        "Input",
        [
            ("text", "Centred or scaled training predictors and centred responses or compact labels, with requested component set C and A = max(C)."),
            ("math", r"X_c\in\mathbb{R}^{n\times p},\qquad Y_c\in\mathbb{R}^{n\times q}"),
        ],
    ),
    (
        "1",
        [
            ("text", "Compute one rank-e randomised decomposition of the cross-covariance, either explicitly or through matrix-free products."),
            ("math", r"S=X_c^{\mathsf T}Y_c\approx U_eD_eV_e^{\mathsf T},\qquad e\leq A"),
        ],
    ),
    (
        "2",
        [
            ("text", "Construct the retained-score Gram matrix directly from scores or from a cached predictor moment matrix."),
            ("math", r"T=X_cU_e,\qquad H=T^{\mathsf T}T=U_e^{\mathsf T}(X_c^{\mathsf T}X_c)U_e"),
        ],
    ),
    (
        "3",
        [
            ("text", "For every requested prefix, solve its own leading latent system by Cholesky factorisation, with a general solve as fallback; do not invert H."),
            ("math", r"c^{\ast}=\min(c,e),\qquad H_{c^{\ast}}L_c=D_{c^{\ast}}"),
        ],
    ),
    (
        "4",
        [
            ("text", "Store the compact prefix-specific response map. Requests above rank repeat the last map; if e = 0, retain a zero map and an intercept-only path."),
            ("math", r"W_c=L_cV_{c^{\ast}}^{\mathsf T}"),
        ],
    ),
    (
        "5",
        [
            ("text", "Predict in latent space and restore the training-response mean. Form dense predictor-by-response coefficients only when explicitly requested."),
            ("math", r"\widehat{Y}_c=(X_{\mathrm{new},c}U_{c^{\ast}})W_c+\mathbf{1}\bar y^{\mathsf T}"),
        ],
    ),
    (
        "Output",
        [
            ("text", "Truncated bases Ue and Ve, singular values De, prefix-specific latent maps Wc, effective rank and complete requested prediction path."),
        ],
    ),
]


CV_ROWS = [
    (
        "Input",
        [
            ("text", "Predictors, responses or compact labels, a group-preserving fold map, requested components, PLS family, preprocessing rule, prediction head, backend and seed."),
        ],
    ),
    (
        "1",
        [
            ("text", "Before the fold loop, cache only eligible uncentred additive statistics: sample counts, column sums, squared sums and cross-products. Classification uses class counts and class-wise predictor sums instead of dense one-hot responses."),
        ],
    ),
    (
        "2",
        [
            ("text", "For fold h, recover exact training statistics by subtracting held-out raw contributions from the full-data statistics."),
            ("math", r"M_{\mathrm{train},h}=M_{\mathrm{full}}-M_{\mathrm{hold},h}"),
        ],
    ),
    (
        "3",
        [
            ("text", "Derive training-only means and scales, then centre the recovered products. The same training quantities are applied unchanged to the held-out observations."),
            ("math", r"S_{xy,h}=M_{xy,h}-n_h\bar x_h\bar y_h^{\mathsf T},\qquad S_{xx,h}=M_{xx,h}-n_h\bar x_h\bar x_h^{\mathsf T}"),
        ],
    ),
    (
        "4",
        [
            ("text", "When response-sample storage is advantageous, extract the training principal submatrix of the raw response Gram matrix and double-centre it using training-only row sums. OPLS filtering and nonlinear kernels remain fold-specific."),
        ],
    ),
    (
        "5",
        [
            ("text", "Fit one maximal component path in the training fold. For requested prefixes above the effective rank, repeat the last estimable prediction. A zero-direction regression fold predicts its training mean; degenerate classification uses its training prior or sole class."),
        ],
    ),
    (
        "6",
        [
            ("text", "For classification, apply argmax to reconstructed responses or fit LDA using only training-fold scores or algebraically equivalent training-only moments. Write each held-out prediction to its original row."),
        ],
    ),
    (
        "7",
        [
            ("text", "Aggregate complete out-of-fold predictions, compute the declared metric and select the lowest component count among ties within C."),
        ],
    ),
    (
        "8",
        [
            ("text", "For nested validation, repeat Steps 1-7 solely within each outer-training partition, refit the selected model on that complete partition, and predict the outer holdout."),
        ],
    ),
    (
        "Output",
        [
            ("text", "Out-of-fold prediction paths, validation metrics, the component count selected within the evaluated grid, effective-component diagnostics, fallback metadata and optional split indices."),
        ],
    ),
]


OPLS_ROWS = [
    (
        "Input",
        [
            ("text", "Training predictors, responses or compact labels, O orthogonal components and requested predictive component set C."),
            ("math", r"X_c\in\mathbb{R}^{n\times p},\qquad Y_c\in\mathbb{R}^{n\times q}"),
        ],
    ),
    ("1", [("text", "Initialise the residual predictors and their current predictor-response cross-covariance."), ("math", r"X^{(1)}=X_c,\qquad S^{(1)}=X^{(1)\mathsf T}Y_c")]),
    ("2", [("text", "For orthogonal component o, obtain the leading predictive direction from the current cross-covariance."), ("math", r"w^{(o)}=\operatorname{leadleft}(S^{(o)})")]),
    ("3", [("text", "Compute its score and predictor loading."), ("math", r"t^{(o)}=X^{(o)}w^{(o)},\qquad p^{(o)}=\frac{X^{(o)\mathsf T}t^{(o)}}{t^{(o)\mathsf T}t^{(o)}}")]),
    ("4", [("text", "Remove the part of the loading parallel to the predictive direction and normalise the remaining orthogonal weight."), ("math", r"w_{\perp}^{(o)}=p^{(o)}-w^{(o)}\frac{w^{(o)\mathsf T}p^{(o)}}{w^{(o)\mathsf T}w^{(o)}},\qquad w_{\perp}^{(o)}\leftarrow\frac{w_{\perp}^{(o)}}{\lVert w_{\perp}^{(o)}\rVert_2}")]),
    ("5", [("text", "Form the orthogonal score and loading, then apply the rank-one predictor deflation."), ("math", r"t_{\perp}^{(o)}=X^{(o)}w_{\perp}^{(o)},\quad p_{\perp}^{(o)}=\frac{X^{(o)\mathsf T}t_{\perp}^{(o)}}{t_{\perp}^{(o)\mathsf T}t_{\perp}^{(o)}},\quad X^{(o+1)}=X^{(o)}-t_{\perp}^{(o)}p_{\perp}^{(o)\mathsf T}")]),
    ("6", [("text", "Stop at the completed prefix if a required norm is zero or non-finite; otherwise refresh the cross-covariance and continue to O."), ("math", r"S^{(o+1)}=X^{(o+1)\mathsf T}Y_c")]),
    ("7", [("text", "Fit one maximal SIMPLS-family predictive path to the filtered predictors. At prediction time, apply the stored orthogonal deflations in extraction order before the predictive path."),
    ]),
    ("Output", [("text", "Orthogonal weights, loadings and scores; filtered predictors; the predictive SIMPLS-family model; preprocessing and effective-component diagnostics."),
    ]),
]


KERNEL_ROWS = [
    (
        "Input",
        [
            ("text", "Training predictors, responses or compact labels, requested components C, and a linear, radial-basis or polynomial kernel."),
        ],
    ),
    ("1", [("text", "Estimate training centring and scaling. For the linear kernel, use the standardised predictor matrix directly as the design."), ("math", r"Z=X_c")]),
    ("2", [("text", "For a nonlinear kernel, form the training Gram matrix from standardised rows."), ("math", r"K_{ij}=\exp(-\gamma\lVert x_i-x_j\rVert_2^2)\quad\mathrm{or}\quad K_{ij}=(\gamma x_i^{\mathsf T}x_j+c)^d")]),
    ("3", [("text", "Double-centre the nonlinear Gram matrix and store the training column means and grand mean."), ("math", r"H=I_n-\frac{1}{n}\mathbf{1}\mathbf{1}^{\mathsf T},\qquad Z=K_c=HKH")]),
    ("4", [("text", "Fit one maximal SIMPLS-family path to Z and the centred response or compact labels; retain every requested prefix and the effective-rank diagnostics."), ("math", r"S_Z=Z^{\mathsf T}Y_c")]),
    ("5", [("text", "For new observations, use standardised predictors directly for the linear route. For a nonlinear route, form the cross-kernel to the stored training rows and centre it with its row means and the stored training means."), ("math", r"K_{\mathrm{new},c}=K_{\mathrm{new}}-r_{\mathrm{new}}\mathbf{1}^{\mathsf T}-\mathbf{1}\mu_K^{\mathsf T}+\bar K")]),
    ("6", [("text", "Predict with the corresponding SIMPLS-family prefix. During cross-validation, recompute all nonlinear kernel centring inside each training fold; only the linear route uses additive sufficient statistics."),
    ]),
    ("Output", [("text", "The shared SIMPLS-family latent model, kernel parameters, training reference data and centring statistics, complete requested prediction path and diagnostics."),
    ]),
]


def replace_figure_before_caption(document: Document, caption_prefix: str, image: Path) -> None:
    caption = find_paragraph(document, caption_prefix)
    caption_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._element is caption._element
    )
    image_paragraph = None
    for paragraph in reversed(document.paragraphs[:caption_index]):
        if paragraph._element.xpath('.//*[local-name()="drawing"]') or paragraph._element.xpath(
            './/*[local-name()="pict"]'
        ):
            image_paragraph = paragraph
            break
    if image_paragraph is None:
        raise RuntimeError(f"image before {caption_prefix!r} not found")
    clear_paragraph(image_paragraph)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image), width=Inches(6.82))


def add_reference(document: Document, text: str) -> None:
    reference_heading = find_paragraph(document, "References")
    insertion_point = document.paragraphs[-1]._element
    paragraph = document.add_paragraph(style="EndNote Bibliography")
    paragraph.add_run(text)
    insertion_point.addnext(paragraph._element)


def clear_review_formatting(document: Document) -> None:
    """Remove residual reviewer highlighting and restore black manuscript text."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(
            paragraph
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.highlight_color = None
            run_properties = run._r.get_or_add_rPr()
            for tag in ("highlight", "shd"):
                for node in run_properties.findall(qn(f"w:{tag}")):
                    run_properties.remove(node)


def edit_main(source: Path, destination: Path, figure1: Path) -> None:
    accepted = Path(tempfile.mkstemp(suffix=".docx")[1])
    try:
        accept_revisions(source, accepted)
        document = Document(accepted)
        replace_in_runs(
            document,
            {
                "randomized": "randomised",
                "optimized": "optimised",
                "specialized": "specialised",
                "centered": "centred",
                "centering": "centring",
                "factorization": "factorisation",
                "normalization": "normalisation",
                "modeling": "modelling",
                "analyzed": "analysed",
                "favor": "favour",
                "behavior": "behaviour",
                "standardized": "standardised",
                "utilized": "utilised",
                "Bidiagonalization": "Bidiagonalisation",
            },
            preserve_references=True,
        )

        # Keep every existing EndNote field intact.  The two new sources are
        # deliberately left as PMID placeholders for insertion through EndNote.
        replace_literal_in_paragraph(
            document,
            "Partial least squares (PLS) regression",
            "REF",
            "PMID: 26002472",
        )
        replace_literal_in_paragraph(
            document,
            "As biological applications have grown",
            "REF",
            "PMID: 24706821",
        )
        replace_literal_in_paragraph(
            document,
            "Single value decomposition",
            "Single value decomposition",
            "Singular value decomposition",
        )
        replace_literal_in_paragraph(
            document,
            "Singular value decomposition",
            "idiagonalization",
            "idiagonalisation",
        )
        set_by_prefix(
            document,
            "Here we present fastPLS",
            "Here we present fastPLS, which provides optimised implementations of PLS-SVD, a SIMPLS-family estimator, OPLS, kernel PLS and related discriminant models. The implementation reduces repeated computation through reusable matrix products, truncated decompositions, compact prediction structures and specialised operations for multivariate and multiclass problems. It also provides a compiled cross-validation framework that limits repeated model construction and storage. We evaluated predictive performance, runtime and memory across different data dimensions and on CPU and GPU hardware.",
        )
        set_by_prefix(
            document,
            "fastPLS handles classification",
            "For classification, fastPLS either assigned the class with the largest reconstructed dummy-response score (argmax) or fitted linear discriminant analysis (LDA) to the PLS scores. LDA used a pooled within-class covariance matrix and automatically increased the ridge term if Cholesky factorisation failed. For regression, the package returned continuous response predictions and evaluated them using RMSD, R²Y or Q²Y.",
        )
        set_by_prefix(document, "2.3 Leakage-Free Cross-Validation", "2.3 Leakage-free cross-validation")
        set_by_prefix(
            document,
            "Cross-validation (CV) is essential",
            "Cross-validation (CV) estimates out-of-sample performance and can select a component count within a prespecified grid. Recomputing preprocessing and matrix products independently in every training fold is expensive, but reusing quantities calculated with validation rows would leak information. fastPLS therefore caches only uncentred additive statistics before the fold loop and reconstructs each training-fold statistic by subtracting the held-out contribution.",
        )
        set_by_prefix(
            document,
            "fastPLS implements an optimised",
            "For each fold, training means, scales and centred cross-products are derived from the recovered training statistics. The same training parameters are then applied to the held-out observations. User-supplied grouping identifiers keep related observations, such as repeated samples from one donor, in the same fold. OPLS filtering is re-estimated within the fold, and nonlinear kernels are constructed and centred entirely from the fold's training observations.",
        )
        set_by_prefix(
            document,
            "Centring, scaling, model fitting",
            "The fitted path retains every requested component position. When a fold yields fewer effective directions than requested, higher prefixes repeat the last estimable prediction. A zero-direction regression fold predicts its training-response mean; a degenerate classification fold uses the training prior or its sole observed class. For classification, users can choose argmax or LDA, and LDA is fitted only from training-fold scores or equivalent training-only moments.",
        )
        set_by_prefix(
            document,
            "To ensure numerical stability during cross-validation",
            "Single cross-validation pools the out-of-fold predictions before calculating the selected accuracy, balanced accuracy, RMSD, R²Y, Q²Y or other documented task metric. Ties are resolved in favour of the lower component count. Nested cross-validation repeats component selection solely within each outer-training partition, refits the selected model on that partition and predicts the outer holdout. Algorithm 3 gives the complete leakage-free procedure.",
        )
        set_by_prefix(
            document,
            "To evaluate execution performance",
            "For the timing comparison, one complete 10-fold validation was compared with one fit on the full training partition followed by prediction of the fixed test partition. Each pair used the same dataset, family, component count, precision, rSVD controls, seed and backend. Classification used LDA in both workflows, and regression used continuous predictions. Timings included fold construction, fitting, prediction and out-of-fold assembly but excluded initial data loading and type conversion; CUDA timings included initialisation, transfers and synchronisation.",
        )

        set_by_prefix(document, "2.3 Dataset retrieval and preparation", "2.4 Dataset retrieval and preparation")
        set_by_prefix(document, "2.4 Benchmark design", "2.5 Benchmark design")
        set_by_prefix(
            document,
            "2.5 fastPLS performance benchmark and independent comparators",
            "2.6 fastPLS performance benchmark and comparator workflows",
        )
        set_by_prefix(document, "3.4 Multivariate NMR prediction", "3.3 Multivariate NMR prediction")
        set_by_prefix(
            document,
            "3.5 Foundation-model embedding analysis",
            "3.4 Foundation-model embedding analysis",
        )
        set_by_prefix(
            document,
            "fastPLS was evaluated against eight established R packages",
            "fastPLS was evaluated against eight R package workflows (pls, plsgenomics, mdatools, plsdepot, pcv, chemometrics, mixOmics and spls) and two Python workflows (IKPLS and scikit-learn). These were package-level comparisons rather than eight independent numerical kernels: several R packages share PLS ancestry, and pcv exposes the internal simpls routine used by its validation workflow. Supplementary Table S3 records every function, estimator, component count, prediction contract and software version.",
        )

        set_by_prefix(
            document,
            "At these reporting points, CUDA PLS-SVD",
            "At these reporting points, CUDA PLS-SVD required 0.492 s and CUDA SIMPLS-family required 0.468 s, compared with 438.862 s for the deposited 165-component workflow. The deposited median comprised 380.439 s for fitting and 58.423 s for held-out prediction; it therefore represents the complete workflow rather than a partial fitting stage. The corresponding runtime reductions were 892-fold and 938-fold. Test RMSD was also lower for fastPLS: 0.0007194 for PLS-SVD and 0.0007230 for the SIMPLS-family estimator, compared with 0.0007858 for the deposited workflow.",
        )
        set_by_prefix(
            document,
            "Figure 3. NMR prediction",
            "Figure 3. NMR prediction using float32 fastPLS at 100 PLS-SVD and 50 SIMPLS-family components, versus the deposited float64 165-component workflow. A: median fitting-plus-prediction time and interquartile range from three processes; the deposited total comprises both stages. B: test RMSD over 28,355 responses. C: host-RSS increment and CUDA allocation. The same fixed 1,200-spectrum training and 321-spectrum test partitions were used.",
        )

        rebuild_algorithm_table(document.tables[0], SIMPLS_ROWS)
        rebuild_algorithm_table(document.tables[1], PLSSVD_ROWS)
        rebuild_algorithm_table(document.tables[2], CV_ROWS)
        for paragraph in document.paragraphs:
            if not paragraph.text.strip() and paragraph.style.name.startswith("Heading"):
                paragraph.style = document.styles["Normal"]
        replace_figure_before_caption(document, "Figure 1.", figure1)

        clear_review_formatting(document)

        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        remove_review_markup(destination)
    finally:
        accepted.unlink(missing_ok=True)


def edit_supplement(source: Path, destination: Path) -> None:
    accepted = Path(tempfile.mkstemp(suffix=".docx")[1])
    try:
        accept_revisions(source, accepted)
        document = Document(accepted)
        replace_in_runs(
            document,
            {
                "randomized": "randomised",
                "optimized": "optimised",
                "specialized": "specialised",
                "centered": "centred",
                "centering": "centring",
                "factorization": "factorisation",
                "normalization": "normalisation",
                "modeling": "modelling",
                "analyzed": "analysed",
                "favor": "favour",
                "behavior": "behaviour",
                "standardized": "standardised",
                "utilized": "utilised",
            },
        )
        rebuild_algorithm_table(document.tables[0], OPLS_ROWS)
        rebuild_algorithm_table(document.tables[1], KERNEL_ROWS)
        clear_review_formatting(document)
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        remove_review_markup(destination)
    finally:
        accepted.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--main-source", type=Path, required=True)
    parser.add_argument("--supp-source", type=Path, required=True)
    parser.add_argument("--figure1", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()

    main_out = args.out_dir / "fastPLS_CMPB_manuscript_compact_algorithms.docx"
    supp_out = args.out_dir / "fastPLS_CMPB_supplement_compact_algorithms.docx"
    edit_main(args.main_source, main_out, args.figure1)
    edit_supplement(args.supp_source, supp_out)
    print(main_out)
    print(supp_out)


if __name__ == "__main__":
    main()
