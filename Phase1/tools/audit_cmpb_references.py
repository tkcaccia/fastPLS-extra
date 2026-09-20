#!/usr/bin/env python3

import argparse
import re
from pathlib import Path

from docx import Document


REFERENCE_RE = re.compile(r"^(?:\[(\d+)\]|(\d+)\.)\s+")
CITATION_RE = re.compile(r"\[((?:\d+\s*(?:[-,]\s*)?)+)\]")


def document_text(document, include_tables=False):
    for paragraph in document.paragraphs:
        yield paragraph.text
    if include_tables:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield cell.text


def expand_citation(text):
    numbers = []
    for part in text.replace(" ", "").split(","):
        if "-" in part:
            start, end = (int(x) for x in part.split("-", 1))
            numbers.extend(range(start, end + 1))
        elif part:
            numbers.append(int(part))
    return numbers


def audit(main_path, supplement_path):
    main = Document(main_path)
    supplement = Document(supplement_path)

    references = []
    in_references = False
    for paragraph in main.paragraphs:
        text = paragraph.text.strip()
        if text.lower() == "references":
            in_references = True
            continue
        if not in_references:
            continue
        match = REFERENCE_RE.match(text)
        if match:
            references.append(int(match.group(1) or match.group(2)))

    errors = []
    expected = list(range(1, len(references) + 1))
    if references != expected:
        errors.append(
            f"reference sequence is {references}; expected {expected}"
        )
    if len(references) != len(set(references)):
        errors.append("duplicate reference numbers detected")

    valid = set(references)
    citations = []
    main_narrative = []
    for text in document_text(main):
        if text.strip().lower() == "references":
            break
        main_narrative.append(text)
    for label, texts in (
        ("main", main_narrative),
        ("supplement", list(document_text(supplement))),
    ):
        for text in texts:
            if REFERENCE_RE.match(text.strip()):
                continue
            for match in CITATION_RE.finditer(text):
                for number in expand_citation(match.group(1)):
                    citations.append((label, number, text.strip()))
                    if number not in valid:
                        errors.append(
                            f"{label} cites [{number}], which is absent from references"
                        )

    joined_main = "\n".join(main_narrative)
    joined_supplement = "\n".join(document_text(supplement))
    required = {
        "main CIFAR-100": (joined_main, r"CIFAR-100[^\n]*\[\d+\]"),
        "main ImageNet/DINOv2": (
            joined_main,
            r"ImageNet/DINOv2[^\n]*\[\d+(?:\s*,\s*\d+)*\]",
        ),
        "main pathology": (
            joined_main,
            r"(?:UNI-?2?|Prov-GigaPath)[\s\S]{0,500}\[\d+(?:\s*,\s*\d+)*\]",
        ),
    }
    for label, (text, pattern) in required.items():
        if not re.search(pattern, text, flags=re.DOTALL):
            errors.append(f"expected citation mapping not found: {label}")

    if errors:
        raise SystemExit("Reference audit failed:\n- " + "\n- ".join(errors))

    counts = {number: 0 for number in references}
    for _, number, _ in citations:
        counts[number] += 1
    print(
        f"Reference audit passed: {len(references)} sequential references, "
        f"{len(citations)} resolved citation links."
    )
    print("Citation counts:", counts)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("main_docx", type=Path)
    parser.add_argument("supplement_docx", type=Path)
    args = parser.parse_args()
    audit(args.main_docx, args.supplement_docx)


if __name__ == "__main__":
    main()
