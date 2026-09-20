#!/usr/bin/env python3
"""Check that CMPB pseudocode retains the active C++ algorithm contracts."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys

from docx import Document


def table_text(document: Document, rows: int, marker: str) -> str:
    for table in document.tables:
        if len(table.rows) != rows or len(table.columns) != 2:
            continue
        text = "\n".join(cell.text for row in table.rows for cell in row.cells)
        if marker in text:
            return text
    raise RuntimeError(f"Could not locate {marker!r} pseudocode table")


def require_terms(text: str, terms: tuple[str, ...], label: str,
                  failures: list[str]) -> None:
    for term in terms:
        if term not in text:
            failures.append(f"{label} is missing {term!r}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("main_document", type=Path)
    parser.add_argument("supplement_document", type=Path)
    parser.add_argument("package_root", type=Path)
    args = parser.parse_args()

    main_document = Document(args.main_document)
    supplement_document = Document(args.supplement_document)
    failures: list[str] = []

    simpls = table_text(main_document, 9, "initial cross-covariance")
    require_terms(
        simpls,
        (
            "requested component-count set C", "A = max(C)",
            "bounded block", "deflated state", "rank-one",
            "a in C", "min(a,e)", "e = 0",
        ),
        "Algorithm 1",
        failures,
    )
    plssvd = table_text(main_document, 7, "implicit products")
    require_terms(
        plssvd,
        (
            "requested component-count set C", "A = max(C)",
            "rank-e randomized decomposition", "score Gram matrix H",
            "Cholesky", "do not invert H", "repeat the map",
        ),
        "Algorithm 2",
        failures,
    )
    cross_validation = table_text(main_document, 12, "fixed K-fold map")
    require_terms(
        cross_validation,
        (
            "additive full-data statistics", "Do not centre or scale with full-data means",
            "I_train,h", "I_hold,h", "training-fold statistics",
            "nonlinear kernel PLS", "min(a,eh)", "single cross-validation",
            "nested cross-validation", "outer-training partition",
        ),
        "Algorithm 3",
        failures,
    )
    opls = table_text(supplement_document, 11, "orthogonal components")
    require_terms(
        opls,
        (
            "component-count set C", "A = max(C)", "w⊥(o)",
            "X(o+1) = X(o) - t⊥(o)p⊥(o)ᵀ", "maximal A-component",
            "min(a,e)", "training scores",
        ),
        "Algorithm S1",
        failures,
    )
    kernel = table_text(supplement_document, 10, "radial-basis")
    require_terms(
        kernel,
        (
            "component-count set C", "A = max(C)",
            "without forming an n by n Gram matrix", "double-centre",
            "training column means", "maximal A-component",
            "min(a,e)", "within each training fold",
        ),
        "Algorithm S2",
        failures,
    )

    source_contracts = {
        "inst/include/fastpls/core/simpls.hpp": (
            "simpls_candidate_block_size", "refresh_directions",
            "backend_rank1_subtract", "model.completed_components",
        ),
        "inst/include/fastpls/core/plssvd.hpp": (
            "backend.cholesky_solve", "backend.general_solve",
            "model.prediction_weights",
        ),
        "inst/include/fastpls/core/opls.hpp": (
            "orthogonal_weight", "orthogonal_loading",
            "fit_opls_filter_from_moments",
        ),
        "inst/include/fastpls/core/kernelpls.hpp": (
            "KernelType::linear", "KernelType::radial_basis",
            "KernelType::polynomial", "center_kernel_train",
            "center_kernel_test", "kernel_grand_mean",
        ),
        "inst/include/fastpls/core/cross_validation.hpp": (
            "dense_sufficient_statistics", "label_sufficient_statistics",
            "make_kernel_fold", "fold_components", "effective_components",
            "model.completed_components == 0",
        ),
    }
    for relative, terms in source_contracts.items():
        path = args.package_root / relative
        if not path.is_file():
            failures.append(f"missing source file {path}")
            continue
        require_terms(path.read_text(), terms, relative, failures)

    if failures:
        print("FAIL pseudocode/source contract")
        for failure in failures:
            print(f"  {failure}")
        sys.exit(1)
    print("PASS pseudocode/source contract")


if __name__ == "__main__":
    main()
