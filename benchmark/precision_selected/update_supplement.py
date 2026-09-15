#!/usr/bin/env python3

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def new_paragraph_after(paragraph, text="", style=None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    if style is not None:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def picture_after(paragraph, image_path, page_break=True):
    inserted = new_paragraph_after(paragraph)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = inserted.add_run()
    if page_break:
        run.add_break(WD_BREAK.PAGE)
    run.add_picture(str(image_path), width=Inches(6.45))
    return inserted


def median(values):
    ordered = sorted(values)
    size = len(ordered)
    middle = size // 2
    if size % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / 2


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--ratios", required=True)
    parser.add_argument("--runtime-figure", required=True)
    parser.add_argument("--memory-figure", required=True)
    args = parser.parse_args()

    document = Document(args.input)
    ratios = read_csv(args.ratios)
    completed = [
        row for row in ratios
        if row.get("runtime_ratio") and row.get("absolute_peak_rss_ratio")
    ]

    panel_order = [
        ("Linux", "cpu", "Linux CPU"),
        ("Linux", "cuda", "Linux CUDA"),
        ("macOS", "cpu", "macOS CPU"),
    ]
    panel_summaries = []
    for platform, backend, label in panel_order:
        rows = [
            row for row in completed
            if row["platform"] == platform and row["backend"] == backend
        ]
        runtime = [float(row["runtime_ratio"]) for row in rows]
        memory = [float(row["absolute_peak_rss_ratio"]) for row in rows]
        panel_summaries.append(
            f"{label}: {median(runtime):.2f} runtime and "
            f"{median(memory):.2f} absolute-RSS ratio"
        )

    accuracy_rows = [
        row for row in completed if row["metric_name"] == "accuracy"
    ]
    rmsd_rows = [row for row in completed if row["metric_name"] == "RMSD"]
    maximum_accuracy_difference = max(
        abs(float(row["metric_difference"])) for row in accuracy_rows
    )
    maximum_relative_rmsd_difference = max(
        abs(float(row["metric_difference"])) /
        abs(float(row["metric_median_float64"]))
        for row in rmsd_rows
    )

    heading = find_paragraph(document, "S7.")
    heading.text = "S7. CPU core and numerical-precision comparisons"
    first_methods = new_paragraph_after(
        heading, "S7.1 One-core versus four-core CPU execution",
        document.styles["Heading 2"],
    )

    anchor = find_paragraph(document, "Figure S16.")
    subsection = new_paragraph_after(
        anchor, "S7.2 Float32 versus float64 execution",
        document.styles["Heading 2"],
    )
    methods = new_paragraph_after(
        subsection,
        (
            "The precision comparison used the same 13 family-dataset workloads, "
            "fixed train/test partitions, component counts, model settings, seed "
            "and output contract as Figure 2. PLS-SVD, SIMPLS, OPLS and linear "
            "kernel PLS were evaluated with argmax for classification and "
            "continuous prediction for regression. Each cell is the median of "
            "three fresh processes with one CPU core requested. Input conversion "
            "was completed before the monitored fitting-and-prediction interval. "
            "Linux CPU and CUDA were measured on the same Intel/NVIDIA computer; "
            "macOS CPU used Apple Accelerate. Metal is not included because the "
            "public Metal route accepts float32 but not float64."
        ),
        document.styles["Normal"],
    )
    results = new_paragraph_after(
        methods,
        (
            "Across the paired cells, the median float64/float32 runtime and "
            "float32/float64 absolute peak process-RSS ratios were "
            + "; ".join(panel_summaries)
            + ". The largest paired classification-accuracy difference was "
            f"{maximum_accuracy_difference:.4f}, and the largest relative "
            "RMSD difference was "
            f"{100 * maximum_relative_rmsd_difference:.2f}%. Ratios are "
            "interpreted within a platform and backend, not between computers."
        ),
        document.styles["Normal"],
    )

    runtime_picture = picture_after(
        results, Path(args.runtime_figure), page_break=True
    )
    runtime_caption = new_paragraph_after(
        runtime_picture,
        (
            "Figure S17. Runtime effect of float32 relative to float64 for the "
            "selected PLS workloads. Cell values are median float64 total time "
            "divided by median float32 total time; values above one indicate "
            "faster float32 execution. Total time includes fitting and held-out "
            "prediction. Every comparison uses matched data, component count, "
            "family, backend, seed and returned prediction, with three fresh "
            "processes per precision. NE denotes a missing matched result."
        ),
        anchor.style,
    )
    memory_picture = picture_after(
        runtime_caption, Path(args.memory_figure), page_break=True
    )
    new_paragraph_after(
        memory_picture,
        (
            "Figure S18. Absolute peak process-memory effect of float32 relative "
            "to float64 for the selected PLS workloads. Cell values are median "
            "float32 peak RSS divided by median float64 peak RSS; values below "
            "one indicate lower complete-process peak RSS with float32. Absolute "
            "RSS includes R, loaded inputs, numerical libraries, fitting and "
            "held-out prediction. Every comparison uses three fresh processes "
            "per precision. NE denotes a missing matched result."
        ),
        anchor.style,
    )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    main()
