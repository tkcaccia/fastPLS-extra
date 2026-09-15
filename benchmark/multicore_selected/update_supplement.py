#!/usr/bin/env python3

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


FAMILY_ORDER = ["plssvd", "simpls", "opls", "kernelpls"]
FAMILY_LABELS = {
    "plssvd": "PLS-SVD",
    "simpls": "SIMPLS",
    "opls": "OPLS",
    "kernelpls": "linear kernel PLS",
}


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def paragraph_before(paragraph):
    element = paragraph._p.getprevious()
    while element is not None and element.tag != qn("w:p"):
        element = element.getprevious()
    if element is None:
        raise ValueError("No paragraph precedes the caption.")
    return Paragraph(element, paragraph._parent)


def new_paragraph_after(paragraph, text="", style=None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    if style is not None:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def replace_picture(paragraph, image_path, page_break=False):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    if page_break:
        run.add_break(WD_BREAK.PAGE)
    run.add_picture(str(image_path), width=Inches(6.45))


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def format_table(table):
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(
                cell,
                "17365D" if row_index == 0 else (
                    "EDF3F8" if row_index % 2 == 0 else "FFFFFF"
                ),
            )
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index < 2
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7.1)
                    run.font.bold = row_index == 0
                    if row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(repeat)


def replace_table(document, index, headers, rows):
    old = document.tables[index]
    new = document.add_table(rows=1, cols=len(headers))
    if old.style is not None:
        new.style = old.style
    for column, label in enumerate(headers):
        new.rows[0].cells[column].text = label
    for values in rows:
        cells = new.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = str(value)
    old._tbl.addprevious(new._tbl)
    old._element.getparent().remove(old._element)
    format_table(new)


def median(values):
    ordered = sorted(values)
    size = len(ordered)
    middle = size // 2
    if size % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / 2


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--ratios", required=True)
    parser.add_argument("--runtime-figure", required=True)
    parser.add_argument("--memory-figure", required=True)
    args = parser.parse_args()

    document = Document(args.input)
    ratios = read_csv(args.ratios)
    summary_rows = []
    for platform in ("Linux", "Mac"):
        for family in FAMILY_ORDER:
            rows = [
                row for row in ratios
                if row["platform"] == platform and row["family"] == family
            ]
            runtime = [float(row["runtime_ratio"]) for row in rows]
            memory = [float(row["absolute_peak_rss_ratio"]) for row in rows]
            metric = []
            for row in rows:
                difference = abs(float(row["metric_difference"]))
                if row["metric_name"].lower() == "accuracy":
                    metric.append(100 * difference)
                else:
                    reference = abs(float(row["metric_median_one"]))
                    metric.append(100 * difference / max(reference, 1e-15))
            summary_rows.append([
                platform,
                "OpenBLAS" if platform == "Linux" else "Apple Accelerate",
                FAMILY_LABELS[family],
                len(rows),
                f"{median(runtime):.2f}",
                f"{min(runtime):.2f}-{max(runtime):.2f}",
                f"{median(memory):.2f}",
                f"{max(metric):.3g}",
            ])

    heading = find_paragraph(document, "S7.")
    heading.text = "S7. CPU core scaling on selected benchmark workloads"
    methods = new_paragraph_after(
        heading,
        (
            "The one-core and four-core comparison used the same 13 float32 "
            "family-dataset workloads, fixed train/test partitions, component "
            "counts, model settings and output contract as Figure 2. PLS-SVD, "
            "SIMPLS, OPLS and linear kernel PLS were evaluated with argmax for "
            "classification and continuous prediction for regression. Each cell "
            "is the median of three fresh processes. Linux used the OpenBLAS-linked "
            "fastPLS 0.99.65 build; macOS used Apple Accelerate. The one-core/four-"
            "core time ratio is greater than one when four cores are faster. The "
            "four-core/one-core absolute peak RSS ratio is below one when the "
            "four-core process uses less complete-process peak memory. Absolute "
            "RSS includes R, loaded inputs, numerical libraries, fitting and "
            "held-out prediction."
        ),
        document.styles["Normal"],
    )

    table_caption = find_paragraph(document, "Table S15.")
    table_caption.text = (
        "Table S15. Summary of current-release one-core versus four-core CPU "
        "measurements across 13 selected workloads per PLS family."
    )
    replace_table(
        document,
        16,
        [
            "Platform", "Numerical library", "PLS family", "Paired workloads",
            "Median time ratio", "Time-ratio range", "Median RSS ratio",
            "Maximum metric change (%)",
        ],
        summary_rows,
    )

    runtime_caption = find_paragraph(document, "Figure S15.")
    replace_picture(paragraph_before(runtime_caption), Path(args.runtime_figure))
    runtime_caption.text = (
        "Figure S15. Runtime effect of requesting four rather than one CPU core "
        "for the selected float32 PLS workloads. Cell values are median one-core "
        "total time divided by median four-core total time; values above one "
        "indicate faster four-core execution. Total time includes fitting and "
        "held-out prediction. Linux and Mac ratios are paired within their "
        "respective computers. Each median uses three fresh processes."
    )

    memory_caption = new_paragraph_after(
        runtime_caption,
        (
            "Figure S16. Absolute peak process-memory effect of requesting four "
            "rather than one CPU core for the selected float32 PLS workloads. "
            "Cell values are median four-core peak RSS divided by median one-core "
            "peak RSS; values below one indicate lower complete-process peak RSS "
            "with four cores. Linux and Mac ratios are paired within their "
            "respective computers. Each median uses three fresh processes."
        ),
        runtime_caption.style,
    )
    memory_picture = new_paragraph_after(runtime_caption)
    replace_picture(memory_picture, Path(args.memory_figure), page_break=True)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    main()
