#!/usr/bin/env python3
"""Expand the CMPB independent-method benchmark description."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/stefano/Documents/GPUPLS/document_work/independent_methods_20260915")
MANUSCRIPT_IN = ROOT / "manuscript_input.docx"
SUPPLEMENT_IN = ROOT / "supplement_input.docx"
MANUSCRIPT_OUT = ROOT / "fastPLS_manuscript_0.99.67.docx"
SUPPLEMENT_OUT = ROOT / "fastPLS_supplementary_material_0.99.67.docx"

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
GRID = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), GRID)


def keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and node is not None:
        p_pr.remove(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_table_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def style_table(table, widths, font_size=7.2):
    table.style = "Table Grid"
    table.autofit = False
    set_table_widths(table, widths)
    set_table_borders(table)
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_fill(cell, BLUE if row_index == 0 else
                          (PALE_BLUE if row_index % 2 == 0 else WHITE))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if col_index == 0 and row_index > 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size if row_index else font_size + 0.2)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)


def add_table(doc, headers, rows, widths, font_size=7.2):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table, widths, font_size=font_size)
    return table


def append_before(reference_paragraph, block):
    element = block._p if hasattr(block, "_p") else block._tbl
    reference_paragraph._p.addprevious(element)


def add_paragraph_before(doc, reference, text, style=None, bold_lead=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(6)
    append_before(reference, paragraph)
    return paragraph


def replace_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def update_manuscript():
    doc = Document(MANUSCRIPT_IN)
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if text.startswith("Several PLS formulations address different computational"):
            replace_text(
                paragraph,
                "Several PLS formulations address different computational and "
                "statistical settings. PLS-SVD derives directions from a singular "
                "value decomposition of the predictor-response cross-covariance, "
                "whereas SIMPLS extracts sequential components while preserving a "
                "deflation geometry in the original variable space [8]. Orthogonal "
                "PLS separates response-orthogonal variation [9], and kernel PLS "
                "permits nonlinear relations through a kernel representation [10]. "
                "The R package pls provides established implementations [11], while "
                "IKPLS supplies NumPy- and JAX-based improved-kernel PLS algorithms "
                "for CPU and GPU computation [12]. Scikit-learn provides another "
                "widely used Python PLSRegression implementation [13]."
            )
        elif text.startswith("For large matrices, the dominant directions"):
            replace_text(paragraph, text.replace("[15]", "[14]"))
        elif text.startswith("Large embedding matrices create a related"):
            replace_text(
                paragraph,
                text.replace("[16,17]", "[15,16]")
                .replace("[18]", "[17]")
                .replace("[19]", "[18]")
            )
        elif text.startswith("The same computational pressure arises"):
            replace_text(paragraph, text.replace("[20,21]", "[19,20]"))
        elif text.startswith("The shared CPU core is compiled C++"):
            replace_text(paragraph, text.replace("[22]", "[21]").replace("[23]", "[22]"))
        elif text.startswith("The classification comparison used CCLE"):
            replace_text(paragraph, text.replace("[24]", "[23]"))
        elif text.startswith("The independent-implementation comparison used an Ubuntu"):
            replace_text(
                paragraph,
                "The independent-implementation comparison used an Ubuntu 22.04 "
                "workstation with an Intel Core i7-13700 processor and 32 GiB RAM. "
                "Every implementation used one effective CPU thread. The fastPLS "
                "0.99.65 build was linked to OpenBLAS 0.3.29 and fitted SIMPLS with "
                "rSVD to centred float32 predictors, using oversampling 32, five "
                "power iterations and seed 123. Dataset-specific component counts "
                "were selected from training data and then held fixed across "
                "implementations. Fitted responses, variance summaries and loading "
                "matrices were not requested from fastPLS. The independent R and "
                "Python calls retained their native model objects and used the "
                "precision supported by each workflow. Classification scores were "
                "decoded by argmax unless a package supplied its own classifier. "
                "Exact functions, package versions, parameter settings, timing and "
                "memory instrumentation are reported in the Supplementary Material. "
                "The comparison therefore "
                "measures complete software workflows rather than an identical "
                "numerical kernel. Complete-process RSS, failures and numerical "
                "warnings were retained."
            )
        elif text.startswith("Among completed R workflows"):
            replace_text(
                paragraph,
                text + " The dominant operations and storage of each implementation "
                "are summarized in Supplementary Table S3, and the supported CPU, "
                "threading, accelerator and precision capabilities are listed in "
                "Supplementary Table S6."
            )
        elif text.startswith("These end-to-end comparisons used the same prepared split"):
            replace_text(
                paragraph,
                "These end-to-end comparisons used the same prepared split, "
                "response encoding, component count and one-thread execution "
                "contract. The ImageNet/DINOv2 column in Figure 1 is a separate "
                "single-CPU, single-run, 1,000-component boundary stress test: "
                "fastPLS SIMPLS-LDA achieved accuracy 0.8094 in 63.3 s, whereas "
                "IKPLS achieved accuracy 0.7999 in 197.1 s with argmax. It "
                "demonstrates completion of the million-sample workflow but is "
                "neither classifier matched nor a biomedical validation. Detailed "
                "Python results and the matched CUDA software comparison are "
                "reported in Supplementary Table S17 and Figure S17. Exact "
                "independent-method calls and benchmark instrumentation are given "
                "in Supplementary Tables S18 and S19; per-dataset component counts "
                "are listed in Supplementary Table S11. The timings include "
                "language-interface validation and output-object construction. For "
                "sub-second workflows, allocation and assembly of R model objects "
                "can represent a larger fraction of elapsed time than NumPy-array "
                "assembly in Python; the reported values therefore compare complete "
                "software workflows rather than isolated numerical kernels."
            )
        elif text.startswith("A separate 50-component, native-float64 Python feasibility comparison"):
            replace_text(
                paragraph,
                "A separate 50-component, native-float64 Python feasibility "
                "comparison with scikit-learn PLSRegression completed in 107.5 s "
                "with RMSD 0.000745 and peak RSS 9180 MiB (Supplementary Table "
                "S17). This row is not pooled with the family-specific NMR "
                "implementation comparison."
            )
        elif text.startswith("The ImageNet/DINOv2 experiment used fastPLS 0.99.66"):
            replace_text(paragraph, text.replace("Supplementary Table S18", "Supplementary Table S20"))
        elif text.startswith("The study has limitations."):
            replace_text(paragraph, text.replace("Supplementary Table S19", "Supplementary Table S21"))

    references = []
    for paragraph in list(doc.paragraphs):
        text = paragraph.text
        if text.startswith("[14] Beurier G."):
            paragraph._element.getparent().remove(paragraph._element)
            continue
        if text.startswith("[") and "]" in text:
            try:
                number = int(text[1:text.index("]")])
            except ValueError:
                continue
            if number >= 15:
                replace_text(paragraph, f"[{number - 1}]" + text[text.index("]") + 1:])
            references.append(paragraph)

    for paragraph in doc.paragraphs:
        if "nirs4all" in paragraph.text:
            raise RuntimeError("Unremoved nirs4all reference in manuscript")
    doc.save(MANUSCRIPT_OUT)


def update_supplement():
    doc = Document(SUPPLEMENT_IN)
    paragraphs = doc.paragraphs
    section_heading = next(p for p in paragraphs if p.text == "S9. Independent implementations")
    cuda_heading = next(p for p in paragraphs if p.text == "CUDA comparison with IKPLS")
    for run in list(cuda_heading._p.findall(qn("w:r"))):
        page_breaks = [
            br for br in run.findall(qn("w:br"))
            if br.get(qn("w:type")) == "page"
        ]
        for page_break in page_breaks:
            run.remove(page_break)
        if len(run) == 0:
            cuda_heading._p.remove(run)

    for paragraph in paragraphs:
        text = paragraph.text
        if text.startswith("In Figure 1, all implementations were measured"):
            replace_text(
                paragraph,
                "Figure 1 compares complete fitting-and-prediction workflows on "
                "the same Intel Core i7-13700 workstation. All ordinary tasks used "
                "one effective CPU thread, identical prepared train/test splits and "
                "the component counts in Table S11. The fastPLS 0.99.65 R workflow "
                "used centred float32 predictors, SIMPLS with rSVD, CPU execution, "
                "oversampling 32, five power iterations and seed 123; classification "
                "used LDA and regression returned continuous predictions. The "
                "independent implementations retained their native estimators, "
                "precision and model objects, as detailed below. Python and MATLAB "
                "fastPLS bindings were not included in this comparison."
            )
        elif text == "Table S17. Independent Python PLS results across the benchmark datasets.":
            replace_text(
                paragraph,
                "Table S17. IKPLS and scikit-learn PLS results across the benchmark datasets."
            )
        elif text.startswith("nirs4all-methods was tested"):
            replace_text(
                paragraph,
                "IKPLS and scikit-learn were evaluated with the component counts "
                "in Table S11. ImageNet and NMR were treated as single-run "
                "feasibility tests. Failures and convergence warnings were retained "
                "rather than excluded. IKPLS could not fit the 50-component NMR "
                "model because its coefficient path required a 68.7-GiB float32 "
                "tensor on the 32-GiB workstation; scikit-learn completed the "
                "separate native-float64 feasibility run."
            )
        elif text == "Table S18. Current-release ImageNet/DINOv2 float32 SIMPLS stress test.":
            replace_text(paragraph, "Table S20. Current-release ImageNet/DINOv2 float32 SIMPLS stress test.")
        elif text == "Table S19. Reproducibility, platform and CPU numerical-library summary.":
            replace_text(paragraph, "Table S21. Reproducibility, platform and CPU numerical-library summary.")

    python_results = doc.tables[18]
    for row in list(python_results.rows)[1:]:
        if row.cells[0].text.startswith("nirs4all"):
            python_results._tbl.remove(row._tr)
    capability_table = doc.tables[7]
    for row in list(capability_table.rows)[1:]:
        if row.cells[0].text.startswith("nirs4all"):
            capability_table._tbl.remove(row._tr)

    heading = add_paragraph_before(doc, cuda_heading, "Software and benchmark settings", style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    intro = add_paragraph_before(
        doc,
        cuda_heading,
        "The comparison adapters used only documented package functions except "
        "for pcv:::simpls, which is an internal routine and is identified as such. "
        "A denotes the fixed dataset-specific component count in Table S11. All "
        "classification packages received the same factor labels and dummy-coded "
        "training response where required; all regression packages received the "
        "same multivariate response matrix.",
    )
    caption_18 = add_paragraph_before(
        doc,
        cuda_heading,
        "Table S18. Independent PLS software, functions and parameter settings used in Figure 1.",
    )
    keep_with_next(caption_18)

    method_rows = [
        (
            "pls 2.9.0",
            "R; pls::simpls.fit",
            "SIMPLS; classification and regression",
            "simpls.fit(X, Y, ncomp = A); remaining arguments at package defaults",
            "Final coefficient slice combined with stored X and Y means; argmax for classification",
        ),
        (
            "plsgenomics 1.5.3",
            "R; pls.lda / pls.regression",
            "PLS-LDA for classification; PLS regression",
            "Xtest supplied; ncomp = A; nruncv = 0 for pls.lda",
            "Package test predictions; predclass for classification and final response prediction for regression",
        ),
        (
            "mdatools 0.15.0",
            "R; plsda / pls",
            "PLS-DA for classification; SIMPLS regression",
            "ncomp = A, center = TRUE, scale = FALSE, cv = NULL where accepted; method = 'simpls' for regression",
            "stats::predict on the held-out matrix; final class-score or response matrix",
        ),
        (
            "plsdepot 0.3.1",
            "R; plsdepot::simpls",
            "SIMPLS; classification and regression",
            "Training-derived removal of non-finite or constant predictors; external centering; comps = A",
            "Prediction reconstructed from returned scores and weights; argmax for classification",
        ),
        (
            "pcv 1.1.0",
            "R; pcv:::simpls",
            "Internal SIMPLS routine; classification and regression",
            "Externally centred X and Y; ncomp = A",
            "Prediction reconstructed from scores and weights; argmax for classification",
        ),
        (
            "chemometrics 1.4.4",
            "R; chemometrics::pls_eigen",
            "Eigenvalue PLS; classification and regression",
            "Externally centred X and Y; a = min(A, q, n - 1, p)",
            "Latent-score regression reconstructed on held-out X; argmax for classification",
        ),
        (
            "mixOmics 6.36.0",
            "R; mixOmics::plsda / mixOmics::pls",
            "PLS-DA for classification; PLS regression",
            "ncomp = A, scale = FALSE; mode = 'regression' for continuous responses",
            "stats::predict; max.dist class at component A or final multivariate response prediction",
        ),
        (
            "spls 2.3.2",
            "R; spls::splsda / spls::spls",
            "Sparse PLS-DA and sparse PLS regression",
            "K = A, eta = 0.9; classifier = 'lda' for splsda; scale.x = FALSE, scale.y = FALSE and fit = 'simpls' for regression",
            "stats::predict; package class prediction or final response matrix",
        ),
        (
            "IKPLS 6.1.2",
            "Python; ikpls.numpy.PLS",
            "Improved-kernel PLS implementation selected by algorithm = 2",
            "External float32 centering; center_X/Y = FALSE, scale_X/Y = FALSE, copy = FALSE, dtype = float32; fit(..., A)",
            "predict(..., n_components = A); argmax for classification; coefficient path retained internally",
        ),
        (
            "scikit-learn 1.7.2",
            "Python; sklearn.cross_decomposition.PLSRegression",
            "PLSRegression; classification and regression",
            "n_components = A, scale = FALSE, max_iter = 500, tol = 10^-6, copy = FALSE; float32 interchange, float64 fitted model",
            "Native model plus final held-out prediction; argmax for classification",
        ),
    ]
    table_18 = add_table(
        doc,
        ("Package", "Language and function", "Estimator and task", "Parameters used", "Prediction/output contract"),
        method_rows,
        (0.83, 1.10, 1.17, 1.95, 1.45),
        font_size=6.9,
    )
    tol_paragraph = table_18.rows[-1].cells[3].paragraphs[0]
    tol_text = tol_paragraph.text
    before, after = tol_text.split("10^-6", 1)
    tol_paragraph.clear()
    tol_paragraph.add_run(before + "10")
    exponent = tol_paragraph.add_run("-6")
    exponent.font.superscript = True
    tol_paragraph.add_run(after)
    for run in tol_paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(6.9)
    append_before(cuda_heading, table_18)

    note_18 = add_paragraph_before(
        doc,
        cuda_heading,
        "The R adapters converted float-package matrices to ordinary double "
        "matrices before calling independent R packages. IKPLS retained float32 "
        "through fitting and prediction. scikit-learn received float32 NumPy "
        "arrays, but its fitted coefficients and predictions were float64. These "
        "precision differences are part of the reported complete-workflow "
        "comparison and are not interpreted as estimator equivalence.",
    )
    caption_19 = add_paragraph_before(
        doc,
        cuda_heading,
        "Table S19. Runtime and complete-process memory measurement software used for the R and Python comparisons.",
    )
    keep_with_next(caption_19)
    instrumentation_rows = [
        (
            "fastPLS R worker",
            "Base R system.time()[['elapsed']] around fitting and held-out prediction separately; the two elapsed values were summed",
            "ps 1.9.2 recorded pre-fit RSS; an external Python monitor read /proc/<pid>/status every 5 ms during the timed boundary",
            "Median and IQR over ten fresh processes for ordinary datasets; one process for ImageNet; absolute and baseline-corrected peak RSS in MiB",
        ),
        (
            "Independent R-package workers",
            "Base R proc.time()[3] around each complete package adapter, including fitting and held-out prediction",
            "ps 1.9.2 recorded RSS at the fit boundary; the same external 5-ms /proc/<pid>/status monitor recorded peak process RSS",
            "Three fresh processes when completed within the repetition limit; slow, failed and resource-limited runs retained explicitly",
        ),
        (
            "IKPLS and scikit-learn Python workers",
            "Python time.perf_counter() around fitting and prediction separately; elapsed values were summed",
            "psutil 7.1.3 recorded pre-fit RSS and sampled the worker plus child processes every 5 ms; GNU /usr/bin/time -v supplied the Linux maximum RSS when available",
            "Median and IQR over ten fresh processes for ordinary datasets; NMR and ImageNet feasibility rows used one process; RSS reported in MiB",
        ),
    ]
    table_19 = add_table(
        doc,
        ("Workflow", "Elapsed-time measurement", "Memory measurement", "Replication and summary"),
        instrumentation_rows,
        (1.15, 1.85, 2.25, 1.25),
        font_size=7.1,
    )
    append_before(cuda_heading, table_19)
    add_paragraph_before(
        doc,
        cuda_heading,
        "Peak RSS is the largest complete-process resident set size observed "
        "during the timed interval. It includes the language runtime, loaded "
        "packages, prepared data, model object, predictions and temporary "
        "allocations; it is not an isolated measurement of algorithmic workspace. "
        "The 5-ms sampler was identical for the R workflows, while GNU time "
        "provided an operating-system maximum-RSS check for the Python workflows.",
    )

    for paragraph in doc.paragraphs:
        if "nirs4all" in paragraph.text:
            raise RuntimeError("Unremoved nirs4all reference in supplement")
    doc.save(SUPPLEMENT_OUT)


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    update_manuscript()
    update_supplement()
    print(MANUSCRIPT_OUT)
    print(SUPPLEMENT_OUT)
