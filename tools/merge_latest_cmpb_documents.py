#!/usr/bin/env python3
"""Merge the independent-software documentation into the latest CMPB files."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/stefano/Documents/GPUPLS")
SOURCE = ROOT / "manuscript_style_revision_20260914" / "documents"
OUTPUT = ROOT / "document_work" / "cmpb_latest_20260915"
MANUSCRIPT_IN = SOURCE / "fastPLS_CMPB_manuscript.docx"
SUPPLEMENT_IN = SOURCE / "fastPLS_CMPB_supplement.docx"
MANUSCRIPT_OUT = OUTPUT / "fastPLS_CMPB_manuscript_updated.docx"
SUPPLEMENT_OUT = OUTPUT / "fastPLS_CMPB_supplement_updated.docx"

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
GRID = "D9D9D9"


def replace_paragraph(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def append_before(reference, block):
    element = block._p if hasattr(block, "_p") else block._tbl
    reference._p.addprevious(element)


def add_paragraph_before(doc, reference, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(6)
    append_before(reference, paragraph)
    return paragraph


def keep_with_next(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    if properties.find(qn("w:keepNext")) is None:
        properties.append(OxmlElement("w:keepNext"))


def set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, value=70):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), GRID)


def style_table(table, widths, font_size=7.0):
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            fill = BLUE if row_index == 0 else (
                PALE_BLUE if row_index % 2 == 0 else WHITE
            )
            set_cell_fill(cell, fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index > 0 and column_index == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size + (0.2 if row_index == 0 else 0))
                    run.font.bold = row_index == 0
                    run.font.color.rgb = (
                        RGBColor(255, 255, 255)
                        if row_index == 0
                        else RGBColor(0, 0, 0)
                    )


def add_table(doc, headers, rows, widths, font_size=7.0):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table, widths, font_size)
    return table


def update_manuscript():
    document = Document(MANUSCRIPT_IN)
    for paragraph in document.paragraphs:
        text = paragraph.text
        if text.startswith("The scikit-learn comparison completed all 11"):
            replace_paragraph(
                paragraph,
                "The scikit-learn comparison completed all 11 standard "
                "classification and regression tasks. Separate large-data "
                "feasibility tests are reported in Table S14, together with the "
                "other Python results. The exact R and Python functions, package "
                "versions, parameters and prediction contracts are documented in "
                "Table S15; Table S16 describes the timing, memory measurement and "
                "replication procedures. The CUDA comparison with IKPLS is shown "
                "in Figure S17."
            )
        elif text.startswith("Figure 1. Prediction, runtime and memory"):
            replace_paragraph(
                paragraph,
                text.replace(
                    "other precisions and repetition counts are given in Table S14.",
                    "independent-software settings and repetition counts are given "
                    "in Tables S14-S16.",
                )
            )
        elif text.startswith("The ImageNet experiment fitted 1,000,000"):
            replace_paragraph(
                paragraph,
                text.replace("Table S15", "Table S17").replace(
                    "Table S16", "Table S18"
                )
            )
    document.save(MANUSCRIPT_OUT)


def update_supplement():
    document = Document(SUPPLEMENT_IN)
    paragraphs = document.paragraphs
    cuda_heading = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text == "CUDA comparison with IKPLS"
    )

    for paragraph in paragraphs:
        text = paragraph.text
        if text.startswith("Figure 1 compares the complete public workflows"):
            replace_paragraph(
                paragraph,
                "Figure 1 compares complete fitting-and-prediction workflows on "
                "the same Intel workstation. All ordinary tasks used one effective "
                "CPU thread, identical prepared train/test splits and the fixed "
                "component counts reported with each result. fastPLS 0.99.65 used "
                "centred float32 predictors, SIMPLS with rSVD, CPU execution, "
                "oversampling 32, five power iterations and seed 123; classification "
                "used LDA and regression returned continuous predictions. Fitted "
                "responses, projected training outputs and variance summaries were "
                "omitted. Independent implementations retained their native "
                "estimators, numerical precision and model objects. Table S14 "
                "reports the Python results, Table S15 gives the exact R and Python "
                "calls, and Table S16 documents measurement and replication."
            )
        elif text == (
            "Table S15. ImageNet/DINOv2 predictions with the top-5 API used in "
            "Figure 4. Each classifier uses one maximal 1,000-component fit; rows "
            "report its component prefixes. Reported time and memory refer to the "
            "complete multi-prefix calculation, not a separate fit at each row."
        ):
            replace_paragraph(paragraph, text.replace("Table S15", "Table S17"))
        elif text == (
            "Table S16. Software interfaces, hardware and CPU numerical libraries "
            "associated with the experiments."
        ):
            replace_paragraph(paragraph, text.replace("Table S16", "Table S18"))

    heading = add_paragraph_before(
        document,
        cuda_heading,
        "Software and benchmark settings",
        style="Heading 2",
    )
    keep_with_next(heading)
    add_paragraph_before(
        document,
        cuda_heading,
        "The comparison adapters used documented package functions except for "
        "pcv:::simpls, which is an internal routine and is identified as such. "
        "A denotes the fixed number of components used for a dataset. All "
        "classification implementations received the same factor labels and, "
        "where required, the same dummy-coded training response. Regression "
        "implementations received the same multivariate response matrix.",
    )

    caption_15 = add_paragraph_before(
        document,
        cuda_heading,
        "Table S15. Independent PLS software, functions and parameter settings "
        "used in Figure 1.",
    )
    keep_with_next(caption_15)
    method_rows = [
        (
            "pls 2.9.0", "R; pls::simpls.fit",
            "SIMPLS; classification and regression",
            "simpls.fit(X, Y, ncomp = A); remaining arguments at defaults",
            "Final coefficient slice combined with stored X and Y means; argmax "
            "for classification",
        ),
        (
            "plsgenomics 1.5.3", "R; pls.lda / pls.regression",
            "PLS-LDA; PLS regression",
            "Xtest supplied; ncomp = A; nruncv = 0 for pls.lda",
            "Package test predictions; predclass for classification and the final "
            "response prediction for regression",
        ),
        (
            "mdatools 0.15.0", "R; plsda / pls",
            "PLS-DA; SIMPLS regression",
            "ncomp = A, center = TRUE, scale = FALSE, cv = NULL where accepted; "
            "method = 'simpls' for regression",
            "stats::predict on held-out X; final class-score or response matrix",
        ),
        (
            "plsdepot 0.3.1", "R; plsdepot::simpls",
            "SIMPLS; classification and regression",
            "Training-derived removal of non-finite or constant predictors; "
            "external centering; comps = A",
            "Prediction reconstructed from returned scores and weights; argmax "
            "for classification",
        ),
        (
            "pcv 1.1.0", "R; pcv:::simpls",
            "Internal SIMPLS routine; classification and regression",
            "Externally centred X and Y; ncomp = A",
            "Prediction reconstructed from scores and weights; argmax for "
            "classification",
        ),
        (
            "chemometrics 1.4.4", "R; chemometrics::pls_eigen",
            "Eigenvalue PLS; classification and regression",
            "Externally centred X and Y; a = min(A, q, n - 1, p)",
            "Latent-score regression on held-out X; argmax for classification",
        ),
        (
            "mixOmics 6.36.0", "R; mixOmics::plsda / mixOmics::pls",
            "PLS-DA; PLS regression",
            "ncomp = A, scale = FALSE; mode = 'regression' for continuous responses",
            "stats::predict; max.dist class at component A or final multivariate "
            "response prediction",
        ),
        (
            "spls 2.3.2", "R; spls::splsda / spls::spls",
            "Sparse PLS-DA; sparse PLS regression",
            "K = A, eta = 0.9; classifier = 'lda' for splsda; scale.x = FALSE, "
            "scale.y = FALSE and fit = 'simpls' for regression",
            "stats::predict; package class prediction or final response matrix",
        ),
        (
            "IKPLS 6.1.2", "Python; ikpls.numpy.PLS",
            "Improved-kernel PLS, algorithm 2",
            "External float32 centering; center_X/Y = FALSE, scale_X/Y = FALSE, "
            "copy = FALSE, dtype = float32; fit(..., A)",
            "predict(..., n_components = A); argmax for classification; coefficient "
            "path retained internally",
        ),
        (
            "scikit-learn 1.7.2",
            "Python; sklearn.cross_decomposition.PLSRegression",
            "PLSRegression; classification and regression",
            "n_components = A, scale = FALSE, max_iter = 500, tol = 10^-6, "
            "copy = FALSE; float32 interchange, float64 fitted model",
            "Native model and held-out prediction; argmax for classification",
        ),
    ]
    table_15 = add_table(
        document,
        (
            "Package", "Language and function", "Estimator and task",
            "Parameters used", "Prediction and output contract",
        ),
        method_rows,
        (0.83, 1.10, 1.17, 1.95, 1.45),
        font_size=6.9,
    )
    append_before(cuda_heading, table_15)

    add_paragraph_before(
        document,
        cuda_heading,
        "The R adapters converted float-package matrices to ordinary double "
        "matrices before calling independent R packages. IKPLS retained float32 "
        "through fitting and prediction. scikit-learn received float32 NumPy "
        "arrays, but its fitted coefficients and predictions were float64. These "
        "precision differences are part of the complete-workflow comparison and "
        "are not interpreted as estimator equivalence.",
    )

    caption_16 = add_paragraph_before(
        document,
        cuda_heading,
        "Table S16. Runtime and complete-process memory measurement software used "
        "for the R and Python comparisons.",
    )
    keep_with_next(caption_16)
    instrumentation_rows = [
        (
            "fastPLS R worker",
            "Base R system.time()[['elapsed']] around fitting and held-out "
            "prediction separately; elapsed values were summed",
            "ps 1.9.2 recorded pre-fit RSS; an external Python monitor read "
            "/proc/<pid>/status every 5 ms during the timed interval",
            "Median and IQR over ten fresh processes for ordinary datasets; one "
            "process for ImageNet; absolute and baseline-corrected RSS in MiB",
        ),
        (
            "Independent R-package workers",
            "Base R proc.time()[3] around the complete package adapter, including "
            "fitting and held-out prediction",
            "ps 1.9.2 recorded RSS at the fit boundary; the same external 5-ms "
            "/proc/<pid>/status monitor recorded peak process RSS",
            "Three fresh processes when completed within the repetition limit; "
            "slow, failed and resource-limited runs retained",
        ),
        (
            "IKPLS and scikit-learn Python workers",
            "Python time.perf_counter() around fitting and prediction separately; "
            "elapsed values were summed",
            "psutil 7.1.3 recorded pre-fit RSS and sampled the process tree every "
            "5 ms; GNU /usr/bin/time -v supplied Linux maximum RSS when available",
            "Median and IQR over ten fresh processes for ordinary datasets; NMR "
            "and ImageNet used one process; RSS in MiB",
        ),
    ]
    table_16 = add_table(
        document,
        (
            "Workflow", "Elapsed-time measurement", "Memory measurement",
            "Replication and summary",
        ),
        instrumentation_rows,
        (1.15, 1.85, 2.25, 1.25),
        font_size=7.1,
    )
    append_before(cuda_heading, table_16)
    add_paragraph_before(
        document,
        cuda_heading,
        "Peak RSS is the largest complete-process resident set size observed "
        "during the timed interval. It includes the language runtime, loaded "
        "packages, prepared data, model object, predictions and temporary "
        "allocations; it is not an isolated measure of algorithmic workspace. "
        "The 5-ms sampler was identical for R workflows, while GNU time supplied "
        "an operating-system maximum-RSS check for Python workflows.",
    )
    document.save(SUPPLEMENT_OUT)


def audit():
    manuscript = Document(MANUSCRIPT_OUT)
    supplement = Document(SUPPLEMENT_OUT)
    manuscript_text = "\n".join(p.text for p in manuscript.paragraphs)
    supplement_text = "\n".join(p.text for p in supplement.paragraphs)
    table_citations = (
        "Table S1", "Table S2", "Table S3", "Tables S4 and S5", "Table S6",
        "Table S7", "Tables S8-S9", "Table S10", "Table S11", "Table S12",
        "Table S13", "Table S14", "Table S15", "Table S16", "Table S17",
        "Table S18",
    )
    figure_citations = (
        "Figures S1-S11", "Figure S12", "Figure S13", "Figure S14",
        "Figures S15-S16", "Figure S17", "Figure S18", "Figure S19",
        "Figure S20",
    )
    for citation in table_citations:
        if citation not in manuscript_text:
            raise RuntimeError(f"Missing manuscript citation: {citation}")
    for citation in figure_citations:
        if citation not in manuscript_text:
            raise RuntimeError(f"Missing manuscript citation: {citation}")
    for number in range(1, 19):
        if f"Table S{number}." not in supplement_text:
            raise RuntimeError(f"Table S{number} caption is missing")
    for number in range(1, 21):
        if f"Figure S{number}." not in supplement_text:
            raise RuntimeError(f"Figure S{number} caption is missing")
    for forbidden in ("nirs4all", "Table S19", "Table S20", "Table S21"):
        if forbidden in supplement_text:
            raise RuntimeError(f"Unexpected text remains: {forbidden}")


if __name__ == "__main__":
    OUTPUT.mkdir(parents=True, exist_ok=True)
    update_manuscript()
    update_supplement()
    audit()
    print(MANUSCRIPT_OUT)
    print(SUPPLEMENT_OUT)
