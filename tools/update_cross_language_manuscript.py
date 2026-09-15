#!/usr/bin/env python3
"""Reframe fastPLS as cross-language software while preserving R benchmarks."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/documents")
MAIN_INPUT = ROOT / "fastPLS_manuscript_ordered_crossreferences_0.99.66.docx"
SUPP_INPUT = ROOT / "fastPLS_supplement_ordered_crossreferences_0.99.66.docx"
MAIN_OUTPUT = ROOT / "fastPLS_manuscript_cross_language_0.99.66.docx"
SUPP_OUTPUT = ROOT / "fastPLS_supplement_cross_language_0.99.66.docx"


def replace_paragraph(document, old, new):
    matches = [p for p in document.paragraphs if p.text == old]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph match, found {len(matches)}: {old[:80]}")
    paragraph = matches[0]
    if not paragraph.runs:
        paragraph.add_run(new)
        return
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_after(paragraph, text):
    new_xml = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_xml.append(deepcopy(paragraph._p.pPr))
    paragraph._p.addnext(new_xml)
    inserted = Paragraph(new_xml, paragraph._parent)
    inserted.add_run(text)
    return inserted


def replace_cell(table, row_index, col_index, expected, replacement):
    cell = table.rows[row_index].cells[col_index]
    if cell.text != expected:
        raise RuntimeError(
            f"Unexpected table cell at ({row_index}, {col_index}): {cell.text!r}"
        )
    paragraph = cell.paragraphs[0]
    paragraph.runs[0].text = replacement
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_table_row_after(table, row_index, values):
    source_row = table.rows[row_index]
    new_row = deepcopy(source_row._tr)
    source_row._tr.addnext(new_row)
    cells = table.rows[row_index + 1].cells
    if len(cells) != len(values):
        raise RuntimeError("Inserted table row has an unexpected column count")
    for cell, value in zip(cells, values):
        cell.text = value


def update_main():
    document = Document(MAIN_INPUT)
    replacements = {
        "Background and objective: Partial least-squares (PLS) models are valuable for biomedical data with correlated predictors and multivariate responses, but repeated component extraction and dense prediction paths can become computationally limiting. We developed fastPLS around an accelerated execution of the sequential SIMPLS algorithm.":
            "Background and objective: Partial least-squares (PLS) models are valuable for biomedical data with correlated predictors and multivariate responses, but repeated component extraction and dense prediction paths can become computationally limiting. We developed fastPLS to accelerate sequential SIMPLS execution.",
        "Methods: fastPLS combines compiled sequential component updates, reusable deflation products, sample-space Gram updates for response-wide problems, compact latent prediction, label-aware class products and randomized singular value decomposition (rSVD). The package also implements PLS-SVD, orthogonal PLS (OPLS) and kernel PLS. Numerical routes include compiled CPU execution, fully device-resident NVIDIA CUDA float32/float64 execution and a fixed CPU/Metal operation split for Apple silicon in float32. Standard R matrices use float64, whereas float-package matrices select float32. Benchmarks measured prediction, fitting-plus-prediction time and process memory on biomedical and controlled matrix regimes. The independent-implementation comparison used fastPLS 0.99.65 on Linux.":
            "Methods: fastPLS combines compiled sequential updates, reusable deflation products, sample-space response Gram updates, compact latent prediction, label-aware class products and randomized singular value decomposition (rSVD). A shared C++17 core implements SIMPLS, PLS-SVD, orthogonal PLS and kernel PLS through R, Python and MATLAB interfaces. The R interface supplies the CUDA and Apple Metal routes evaluated here, and all reported analyses and benchmarks used that interface. Python and MATLAB parity was validated separately. Benchmarks measured predictive performance, fitting-plus-prediction time and process memory.",
        "Conclusions: fastPLS extends sequential PLS computation to matrix regimes that are difficult to process through conventional R workflows. Its gains arise from the combined execution and storage strategy rather than from a universally faster individual optimization. rSVD remains approximate, and precision, CPU threading and accelerator use should be selected according to numerical agreement and matrix shape.":
            "Conclusions: fastPLS extends sequential PLS computation to matrix regimes that are difficult for conventional PLS software. Its gains arise from the combined execution and storage strategy rather than from a universally faster individual optimization. rSVD remains approximate, and precision, CPU threading and accelerator use should be selected according to numerical agreement and matrix shape.",
        "Several PLS formulations address different computational and statistical settings. PLS-SVD derives directions from a singular value decomposition of the predictor-response cross-covariance, whereas SIMPLS extracts sequential components while preserving a deflation geometry in the original variable space [8]. Orthogonal PLS separates response-orthogonal variation [9], and kernel PLS permits nonlinear relations through a kernel representation [10]. The R package pls provides established implementations [11], while IKPLS supplies NumPy- and JAX-based improved-kernel PLS algorithms for CPU and GPU computation [12]. Python alternatives include the widely used scikit-learn PLSRegression estimator [13] and the compiled multi-solver nirs4all-methods engine [14].":
            "Several PLS formulations address different computational and statistical settings. PLS-SVD derives directions from a singular value decomposition of the predictor-response cross-covariance, whereas SIMPLS extracts sequential components while preserving a deflation geometry in the original variable space [8]. Orthogonal PLS separates response-orthogonal variation [9], and kernel PLS permits nonlinear relations through a kernel representation [10]. The R package pls provides established implementations [11], while IKPLS supplies NumPy- and JAX-based improved-kernel PLS algorithms for CPU and GPU computation [12]. Other Python alternatives include the widely used scikit-learn PLSRegression estimator [13] and the compiled multi-solver nirs4all-methods engine [14].",
        "We present fastPLS, a Bioconductor R package centred on accelerated SIMPLS execution. Its contribution is not the component-path concept already available in conventional SIMPLS software, but the combination of compiled sequential updates, reusable products and persistent workspaces, sample-space Gram updates for response-wide problems, compact prediction, label-aware class products, matrix-free operators and corresponding CPU, CUDA and Metal implementations. PLS-SVD, OPLS, kernel PLS, argmax decoding, latent-space linear discriminant analysis (LDA) and compiled cross-validation share the same public interface.":
            "We present fastPLS, a cross-language implementation centred on accelerated SIMPLS execution and a shared MIT-licensed C++17 numerical core. Public interfaces are available for R, Python and MATLAB. Its contribution is not the component-path concept already available in conventional SIMPLS software, but the combination of compiled sequential updates, reusable products and persistent workspaces, sample-space Gram updates for response-wide problems, compact prediction, label-aware class products and matrix-free operators. The R interface additionally exposes the CPU, CUDA and Metal routes, prediction heads and compiled cross-validation evaluated in this study.",
        "CPU routes are compiled C++ and call the linked BLAS/LAPACK; multicore execution is available only when the linked numerical library and operation use multiple threads. CUDA routes use NVIDIA cuBLAS, cuSOLVER and cuRAND for device matrix products, factorizations and randomized sketches [22]. After input transfer, CUDA retains preprocessing, decomposition, SIMPLS updates, OPLS filtering, nonlinear-kernel construction and centering, prediction and LDA on the device in float32 and float64. The Apple route instead uses a fixed operation split: persistent Metal Performance Shaders and custom kernels evaluate fitting products involving the training sample matrix, whereas the CPU performs centering and scaling, reduced QR/SVD calculations, sequential orthogonalization and deflation, coefficient assembly, LDA and compact prediction [23]. This assignment is invariant to dataset shape. Metal accepts float32 because the tested Apple GPU does not provide the float64 arithmetic required by the kernels. Unsupported accelerator or precision combinations return an error rather than silently falling back to CPU.":
            "The shared CPU core is compiled C++ and calls the linked BLAS/LAPACK; multicore execution is available only when the linked numerical library and operation use multiple threads. The R interface additionally exposes CUDA routes that use NVIDIA cuBLAS, cuSOLVER and cuRAND for device matrix products, factorizations and randomized sketches [22]. After input transfer, CUDA retains preprocessing, decomposition, SIMPLS updates, OPLS filtering, nonlinear-kernel construction and centering, prediction and LDA on the device in float32 and float64. The R interface also exposes an Apple route with a fixed operation split: persistent Metal Performance Shaders and custom kernels evaluate fitting products involving the training sample matrix, whereas the CPU performs centering and scaling, reduced QR/SVD calculations, sequential orthogonalization and deflation, coefficient assembly, LDA and compact prediction [23]. This assignment is invariant to dataset shape. Metal accepts float32 because the tested Apple GPU does not provide the float64 arithmetic required by the kernels. Unsupported accelerator or precision combinations return an error rather than silently falling back to CPU.",
        "R stores ordinary numeric matrices in float64. Inputs created with the float package occupy float32 storage and automatically select a supported float32 numerical route. This can reduce the size of input and intermediate arrays, but complete-process memory and runtime also include the R session, loaded data, libraries, allocator state and device context. Precision comparisons consequently report measured memory and prediction differences rather than inferring a twofold process-memory reduction.":
            "For the R analyses reported here, ordinary numeric matrices use float64. Inputs created with the float package occupy float32 storage and automatically select a supported float32 numerical route. This can reduce the size of input and intermediate arrays, but complete-process memory and runtime also include the R session, loaded data, libraries, allocator state and device context. Precision comparisons consequently report measured memory and prediction differences rather than inferring a twofold process-memory reduction.",
        "The compiled single and double cross-validation procedures preserve groups supplied through constrain, preventing observations from one subject from entering both training and validation folds. Component number and prediction-relevant model arguments can be selected using accuracy, balanced accuracy, RMSD or Q² as appropriate. Training R², independent-test Q² relative to the training mean and fold-based cross-validated Q² are returned as distinct quantities. Permutation p-values use the finite-sample correction (b + 1)/(B + 1); when groups are supplied, labels are permuted at the exchangeability-block level.":
            "The compiled single and double cross-validation procedures exposed by the R interface preserve groups supplied through constrain, preventing observations from one subject from entering both training and validation folds. Component number and prediction-relevant model arguments can be selected using accuracy, balanced accuracy, RMSD or Q² as appropriate. Training R², independent-test Q² relative to the training mean and fold-based cross-validated Q² are returned as distinct quantities. Permutation p-values use the finite-sample correction (b + 1)/(B + 1); when groups are supplied, labels are permuted at the exchangeability-block level.",
        "The effect of the low-rank solver was isolated in the companion benchmark by holding the simulated matrices, 20-component workload, preprocessing, family and prediction output constant. PLS-SVD and SIMPLS were compared with native rSVD and the IRLBA companion on balanced, predictor-wide and response-wide shapes in five isolated processes. This comparison evaluates solver cost within the same fastPLS execution design; IRLBA is not part of the public package (Supplementary Table S10 and Figure S14). Training-only component settings and their associations with predictive and computational metrics are reported in Supplementary Tables S11 and S12.":
            "The effect of the low-rank solver was isolated in the companion benchmark by holding the simulated matrices, 20-component workload, preprocessing, family and prediction output constant. PLS-SVD and SIMPLS were compared with native rSVD and the IRLBA companion on balanced, predictor-wide and response-wide shapes in five isolated processes. This comparison evaluates solver cost within the same fastPLS execution design; IRLBA is not part of the public fastPLS interfaces (Supplementary Table S10 and Figure S14). Training-only component settings and their associations with predictive and computational metrics are reported in Supplementary Tables S11 and S12.",
        "fastPLS accelerates SIMPLS through execution and storage choices rather than redefining the underlying sequential PLS objective. Reusing deflation products and workspaces reduces repeated allocation, while compact latent prediction avoids dense p by q coefficient arrays for every requested prefix. Label-aware products are particularly helpful in multiclass problems because the response cross-product can be formed from class sums without materializing a dense n by J indicator matrix. These gains explain why the current single-CPU workflow was fastest or tied on most completed classification datasets, while the remaining small task showed that package overhead can dominate when all methods finish in hundredths of a second.":
            "fastPLS accelerates SIMPLS through execution and storage choices rather than redefining the underlying sequential PLS objective. Reusing deflation products and workspaces reduces repeated allocation, while compact latent prediction avoids dense p by q coefficient arrays for every requested prefix. Label-aware products are particularly helpful in multiclass problems because the response cross-product can be formed from class sums without materializing a dense n by J indicator matrix. These gains explain why the current single-CPU workflow was fastest or tied on most completed classification datasets, while the remaining small task showed that interface and object-construction overhead can dominate when all methods finish in hundredths of a second.",
        "fastPLS combines a compiled sequential SIMPLS path with compact prediction, rSVD and optional CUDA and Metal execution. It extends practical PLS analysis to response and sample dimensions that were previously difficult to process in R, while preserving explicit reporting of precision, backend residency and numerical controls. The NMR application demonstrates the value of the implementation for high-response biomedical prediction; appropriate solver and backend choices remain workload dependent.":
            "fastPLS combines a shared compiled SIMPLS core with compact prediction, rSVD and optional CUDA and Metal execution through its R interface. Public R, Python and MATLAB interfaces extend practical PLS analysis to response and sample dimensions that were previously difficult for conventional PLS software, while preserving explicit reporting of precision and numerical controls. The NMR application demonstrates the value of the implementation for high-response biomedical prediction; appropriate solver and backend choices remain workload dependent.",
        "Data and software availability: fastPLS is available through Bioconductor and at https://github.com/tkcaccia/fastPLS. Executable benchmark and figure-generation scripts are maintained separately at https://github.com/tkcaccia/fastPLS-extra; the frozen publication archive will include the result tables and restricted-data acquisition instructions. Public datasets remain subject to their source licences and access conditions.":
            "Data and software availability: fastPLS uses a shared C++17 numerical core and is available for R through Bioconductor and GitHub (https://github.com/tkcaccia/fastPLS), for Python through PyPI (https://pypi.org/project/fastPLS/) and GitHub (https://github.com/tkcaccia/fastPLS-py), and for MATLAB through the versioned release at https://github.com/tkcaccia/fastPLS-matlab/releases/tag/0.1.0. Every fastPLS analysis and benchmark reported here used the R interface. Executable R benchmark and figure-generation scripts are maintained separately at https://github.com/tkcaccia/fastPLS-extra; the frozen publication archive will include the result tables and restricted-data acquisition instructions. Public datasets remain subject to their source licences and access conditions.",
    }
    for old, new in replacements.items():
        replace_paragraph(document, old, new)

    anchor = next(
        p for p in document.paragraphs
        if p.text.startswith("The Metal installation step compiles and embeds")
    )
    insert_after(
        anchor,
        "fastPLS is organized around one portable C++17 CPU core exposed through R, "
        "Python and MATLAB. The Python and MATLAB 0.1.0 releases vendor the same core "
        "as fastPLS 0.99.66. Interface-parity tests use identical matrices, precision, "
        "controls and seeds. All model fitting, cross-validation, timing, memory "
        "measurement and figure or table generation reported in this article used the "
        "R interface. The Python and MATLAB interfaces were validated separately and "
        "were not mixed into benchmark timings. The CUDA and Metal routes evaluated "
        "here are currently exposed through R; the Python and MATLAB releases expose "
        "the portable CPU core and reject unsupported accelerator requests explicitly.",
    )
    document.save(MAIN_OUTPUT)


def update_supplement():
    document = Document(SUPP_INPUT)
    replacements = {
        "CPU/accelerator runtime ratios above one favour the accelerator. Every row was evaluated with fastPLS 0.99.66 in three fresh processes and includes device and pipeline initialization. The Metal route uses one invariant operation partition: Metal evaluates fitting products involving the training sample matrix, while the CPU performs reduced factorizations, sequential PLS updates and compact prediction. Host RSS is the baseline-corrected complete-process increment. CUDA device allocation can include context, library and allocator state. Values are paired within workstation, and no completed timing is suppressed on the basis of metric disagreement.":
            "CPU/accelerator runtime ratios above one favour the accelerator. Every row was evaluated through the fastPLS 0.99.66 R interface in three fresh processes and includes device and pipeline initialization. The Metal route uses one invariant operation partition: Metal evaluates fitting products involving the training sample matrix, while the CPU performs reduced factorizations, sequential PLS updates and compact prediction. Host RSS is the baseline-corrected complete-process increment. CUDA device allocation can include context, library and allocator state. Values are paired within workstation, and no completed timing is suppressed on the basis of metric disagreement.",
        "The matrices, preprocessing, 20-component workload, model family and held-out prediction output were identical within each pair. IRLBA is retained only in the GPL companion benchmark package; it is not a public fastPLS solver. The comparison therefore isolates low-rank solver choice within the same surrounding PLS implementation.":
            "The matrices, preprocessing, 20-component workload, model family and held-out prediction output were identical within each pair. IRLBA is retained only for the separate GPL comparison workflow; it is not exposed by the public fastPLS interfaces. The comparison therefore isolates low-rank solver choice within the same surrounding PLS implementation.",
        "For the eleven general benchmark datasets, component counts were selected using only the training partition. Ten fixed folds (seed 123) were used with float32 inputs and the CPU rSVD route in fastPLS 0.99.65. At every admissible count, classification evaluated both argmax and LDA and retained the component-head pair with the greatest pooled out-of-fold accuracy. Regression retained the component count with the smallest pooled out-of-fold RMSD. Ties were resolved by the ordered grid in favour of the smaller count. OPLS removed one orthogonal component and kernel PLS used the linear kernel. PLS-SVD classification was restricted to q - 1 components. Selections at the largest evaluated count or an intrinsic rank limit are labelled explicitly and are not interpreted as unconstrained optima. The package selected its public automatic rSVD controls; the effective oversampling and power values were recorded with every result.":
            "For the eleven general benchmark datasets, component counts were selected using only the training partition. Ten fixed folds (seed 123) were used with float32 inputs and the CPU rSVD route through the fastPLS 0.99.65 R interface. At every admissible count, classification evaluated both argmax and LDA and retained the component-head pair with the greatest pooled out-of-fold accuracy. Regression retained the component count with the smallest pooled out-of-fold RMSD. Ties were resolved by the ordered grid in favour of the smaller count. OPLS removed one orthogonal component and kernel PLS used the linear kernel. PLS-SVD classification was restricted to q - 1 components. Selections at the largest evaluated count or an intrinsic rank limit are labelled explicitly and are not interpreted as unconstrained optima. The R interface selected its public automatic rSVD controls; the effective oversampling and power values were recorded with every result.",
        "float32 support means that the principal supported numerical buffers use single precision. It does not imply that every R bookkeeping object or complete-process allocation is half the corresponding float64 measurement. Metal is an operation-level CPU/GPU implementation, not a whole-model shape dispatcher: the same operations are assigned to Metal and CPU for every dataset. An unsupported backend or precision combination raises an informative error before fitting; no CPU fallback is used.":
            "For the R benchmark routes, float32 support means that the principal supported numerical buffers use single precision. It does not imply that every R bookkeeping object or complete-process allocation is half the corresponding float64 measurement. Metal is an operation-level CPU/GPU implementation, not a whole-model shape dispatcher: the same operations are assigned to Metal and CPU for every dataset. An unsupported backend or precision combination raises an informative error before fitting; no CPU fallback is used.",
        "In Figure 1, all implementations were measured on the same Intel Core i7-13700 workstation. fastPLS 0.99.65 used method = 'simpls', svd.method = 'rsvd', backend = 'cpu', scaling = 'centering', fit = FALSE, proj = FALSE, return_variance = FALSE, float32 predictors, oversample = 32, power = 5 and seed = 123. The classifier was 'argmax' or 'lda'; the dataset-specific ncomp values were selected from training data and are listed in Table S11. IKPLS used its improved-kernel formulation in float32. nirs4all-methods and scikit-learn used their native float64 Python routes. All ordinary tasks used one effective CPU thread, identical inputs, train/test splits and component counts.":
            "In Figure 1, all implementations were measured on the same Intel Core i7-13700 workstation. The fastPLS 0.99.65 R interface used method = 'simpls', svd.method = 'rsvd', backend = 'cpu', scaling = 'centering', fit = FALSE, proj = FALSE, return_variance = FALSE, float32 predictors, oversample = 32, power = 5 and seed = 123. The classifier was 'argmax' or 'lda'; the dataset-specific ncomp values were selected from training data and are listed in Table S11. IKPLS used its improved-kernel formulation in float32. nirs4all-methods and scikit-learn used their native float64 Python routes. All ordinary tasks used one effective CPU thread, identical inputs, train/test splits and component counts. Python and MATLAB fastPLS bindings were not included in this performance comparison.",
        "The Basic Linear Algebra Subprograms (BLAS) define standard dense matrix operations but do not identify a specific vendor library. fastPLS is linked at compilation to one CPU BLAS/LAPACK implementation. Apple Accelerate is the default optimized implementation on macOS, whereas OpenBLAS is used for the authors' Linux and Windows performance builds. If OpenBLAS is unavailable, installation remains possible using the BLAS/LAPACK libraries supplied with R. These implementations can differ substantially in latency and thread scaling, so the linked library and active thread count must be reported with CPU benchmarks. The fastPLS_blas() function returns the library selected at compilation. Apple Accelerate is a CPU numerical framework and is distinct from the package's Metal CPU/GPU backend.":
            "The Basic Linear Algebra Subprograms (BLAS) define standard dense matrix operations but do not identify a specific vendor library. The fastPLS R interface used for the reported benchmarks is linked at compilation to one CPU BLAS/LAPACK implementation. Apple Accelerate is the default optimized implementation on macOS, whereas OpenBLAS is used for the authors' Linux and Windows performance builds. If OpenBLAS is unavailable, installation remains possible using the BLAS/LAPACK libraries supplied with R. These implementations can differ substantially in latency and thread scaling, so the linked library and active thread count must be reported with CPU benchmarks. The R function fastPLS_blas() returns the library selected at compilation. Apple Accelerate is a CPU numerical framework and is distinct from the R interface's Metal CPU/GPU backend.",
        "Benchmark scripts record requested method, component count, precision, seed, rSVD controls, executed route, timing, metrics and completion status. CUDA timing includes transfer and synchronization in the public workflow. Metal float64 is tested as an informative error, not as a CPU fallback.":
            "All fastPLS benchmark and figure-generation scripts use the R interface and record requested method, component count, precision, seed, rSVD controls, executed route, timing, metrics and completion status. CUDA timing includes transfer and synchronization in the public R workflow. Metal float64 is tested as an informative error, not as a CPU fallback. Python and MATLAB parity tests are maintained separately and do not contribute timings to the manuscript benchmarks.",
    }
    for old, new in replacements.items():
        replace_paragraph(document, old, new)

    anchor = next(
        p for p in document.paragraphs
        if p.text == "Here A is the maximum retained component count and l is the randomized sketch width. Explicit means that the p by q cross-covariance is stored. Implicit means that only its action on a narrow matrix is evaluated, so S is not stored."
    )
    insert_after(
        anchor,
        "The portable C++17 CPU core is shared by the R, Python and MATLAB "
        "interfaces. Python and MATLAB release 0.1.0 vendor the same core as "
        "fastPLS 0.99.66 and are checked with identical deterministic inputs. All "
        "reported fastPLS analyses and benchmarks use the R interface. Compiled "
        "cross-validation and the CUDA and Metal routes evaluated here are currently "
        "exposed through R; the public Python and MATLAB releases expose the CPU core "
        "and reject unsupported accelerator requests explicitly. Cross-language parity "
        "is therefore software validation, not an additional performance panel.",
    )

    capabilities = document.tables[7]
    replace_cell(
        capabilities,
        1,
        1,
        "Compiled C++/linked BLAS",
        "Shared C++17 core; R, Python and MATLAB interfaces",
    )
    replace_cell(capabilities, 1, 3, "CUDA", "CUDA through R")
    replace_cell(capabilities, 1, 4, "Metal", "Metal through R")

    reproducibility = document.tables[20]
    insert_table_row_after(
        reproducibility,
        1,
        [
            "Public interfaces",
            "R 0.99.66; Python 0.1.0; MATLAB 0.1.0",
            "Shared portable C++17 CPU core; manuscript benchmarks use R only",
        ],
    )
    document.save(SUPP_OUTPUT)


if __name__ == "__main__":
    update_main()
    update_supplement()
    print(MAIN_OUTPUT)
    print(SUPP_OUTPUT)
