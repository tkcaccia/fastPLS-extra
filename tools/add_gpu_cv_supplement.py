#!/usr/bin/env python3
"""Append the accelerator cross-validation audit to a fastPLS supplement."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def margins(cell, value: int = 55) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.first_child_found_in("w:tcMar")
    if element is None:
        element = OxmlElement("w:tcMar")
        properties.append(element)
    for name in ("top", "start", "bottom", "end"):
        side = element.find(qn(f"w:{name}"))
        if side is None:
            side = OxmlElement(f"w:{name}")
            element.append(side)
        side.set(qn("w:w"), str(value))
        side.set(qn("w:type"), "dxa")


def add_table(document: Document, caption: str, headers: list[str],
              rows: list[list[str]], font_size: float = 6.4) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(caption).bold = True
    paragraph.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    for row_index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        properties.append(no_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            properties.append(repeat)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if row_index == 0:
                shade(cell, "D9E2F3")
            elif row_index % 2 == 0:
                shade(cell, "F6F8FA")
            for cell_paragraph in cell.paragraphs:
                cell_paragraph.paragraph_format.space_after = Pt(0)
                cell_paragraph.paragraph_format.line_spacing = 1.0
                for run in cell_paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    if row_index == 0:
                        run.bold = True


def fmt(value: str, digits: int = 3) -> str:
    return f"{float(value):.{digits}f}"


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit(
            "Usage: add_gpu_cv_supplement.py INPUT_DOCX SUMMARY_DIR OUTPUT_DOCX"
        )
    input_docx = Path(sys.argv[1])
    summary_dir = Path(sys.argv[2])
    output_docx = Path(sys.argv[3])
    comparison = read_rows(summary_dir / "cv_comparison_summary.csv")
    ablation = read_rows(summary_dir / "cv_path_ablation_summary.csv")

    document = Document(input_docx)
    document.add_page_break()
    document.add_heading(
        "S11. Accelerator cross-validation residency and transfer audit", 1
    )
    document.add_paragraph(
        "The public cross-validation arguments and fold-construction rules were "
        "kept unchanged. In particular, constrain was converted once to a "
        "deterministic group-aware fold vector, and every row sharing a "
        "constraint value remained in one fold. The audit separated host "
        "orchestration from accelerator-resident numerical work and compared "
        "complete ten-fold validation with one fit and prediction on the first "
        "matched held-out fold."
    )

    residency_rows = [
        ["Fold assignment", "Host", "Host", "Host once; upload fold vector once"],
        ["Fold row extraction", "Host matrices", "Host matrices", "Device gather or fold masks"],
        ["Centering and scaling", "Accelerator after fold upload", "Accelerator after fold upload", "Device fold moments from full and held-out sums"],
        ["PLS fitting", "Resident model created per fold", "Resident model created per fold", "One reusable context and workspace across folds"],
        ["Component-path prediction", "Resident; shared test projection", "Resident; shared test projection", "Resident; retain fold output on device"],
        ["Argmax or LDA", "Resident", "Resident", "Resident"],
        ["Metric reduction", "Host", "Host", "Device reduction"],
        ["Documented score path", "Transferred after each fold", "Transferred after each fold", "One final transfer after all folds"],
    ]
    add_table(
        document,
        "Table S17. Current and target cross-validation residency.",
        ["Stage", "CUDA 0.99.40", "Metal 0.99.40", "Native-CV target"],
        residency_rows,
        6.2,
    )

    numerical = {
        ("metal", "argmax"): "Identical predictions and metrics",
        ("metal", "lda"): "Identical predictions and metrics",
        ("cuda", "argmax"): "Class agreement 0.99972; score error 4.83 × 10⁻⁵",
        ("cuda", "lda"): "Class agreement 1.00000; score error 4.83 × 10⁻⁵",
    }
    ablation_rows = [
        [
            row["backend"].upper() if row["backend"] == "cuda" else "Metal",
            row["classifier"],
            fmt(row["baseline_median_sec"]),
            fmt(row["combined_path_median_sec"]),
            fmt(row["speedup"], 2),
            numerical[(row["backend"], row["classifier"])],
        ]
        for row in ablation
    ]
    add_table(
        document,
        "Table S18. Shared component-path ablation on grouped synthetic classification data.",
        ["Backend", "Head", "Separate path s", "Shared path s", "Speed-up", "Numerical comparison"],
        ablation_rows,
        6.2,
    )
    document.add_paragraph(
        "Values are medians of seven alternating fresh-process runs. The CUDA "
        "relative score error is the Frobenius norm of the score difference "
        "divided by the Frobenius norm of the baseline score array; score "
        "correlation was 1.000000 for both heads."
    )

    dataset_names = {
        "metref": "MetRef",
        "retina": "Retina",
        "cifar": "CIFAR-100",
        "nmr": "NMR",
    }
    comparison_rows = []
    for row in comparison:
        cpu = row["same_host_cpu_cv_median_sec"]
        ratio = row["cpu_to_accelerator_ratio"]
        comparison_rows.append([
            dataset_names[row["dataset"]],
            "CUDA" if row["backend"] == "cuda" else "Metal",
            row["repetitions"],
            fmt(row["cv_median_sec"]),
            f"{fmt(row['cv_q1_sec'])}-{fmt(row['cv_q3_sec'])}",
            fmt(row["one_fold_median_sec"]),
            fmt(row["cv_to_one_fold_ratio"], 2),
            "NA" if cpu == "NA" else fmt(cpu),
            "NA" if ratio == "NA" else fmt(ratio, 2),
            row["best_ncomp"],
            row["metric_name"].upper(),
            fmt(row["metric_min"], 5),
        ])
    add_table(
        document,
        "Table S19. Complete ten-fold SIMPLS-rSVD validation and matched one-fold controls.",
        ["Dataset", "Backend", "Runs", "CV median s", "CV IQR s", "One-fold s", "CV/one-fold", "CPU CV s", "CPU/backend", "Selected A", "Metric", "Value"],
        comparison_rows,
        5.7,
    )
    document.add_paragraph(
        "Ratios greater than one in the CPU/backend column favour the "
        "accelerator. CPU comparisons are made only within the same computer. "
        "The Apple M3 CPU was faster than Metal for Retina and CIFAR-100, "
        "whereas CUDA was 1.07-fold faster for Retina, 5.59-fold faster for "
        "CIFAR-100, and 1.94-fold faster for NMR on the Intel Core i7-13700 and "
        "NVIDIA GeForce RTX 5060 Ti workstation. MetRef remained CPU-favourable "
        "on the CUDA workstation. "
        "All runs preserved constraint groups, selected the same component "
        "within repetitions, and produced deterministic fold assignments."
    )
    document.add_paragraph(
        "The NMR rows are single guarded feasibility probes and are not used "
        "to estimate timing dispersion. Classification rows used the LDA head; "
        "NMR is multivariate regression and therefore has no classification head."
    )
    document.add_paragraph(
        "The results identify two different optimization regimes. Reusing one "
        "test-fold projection for every requested component reduced Metal "
        "cross-validation time by 23-24%, but changed CUDA time by at most 3% "
        "because CUDA was dominated by repeated fold setup. Fully accelerator-"
        "native numerical cross-validation is feasible without changing the "
        "public interface or constrain semantics, but it requires a dedicated "
        "persistent CV workspace rather than a dispatch-only modification. The "
        "full predictor matrix, response or compact labels, and fold vector "
        "would be uploaded once; fold-specific moments would be formed on the "
        "device; model, decomposition, classifier, and metric workspaces would "
        "be reused; and the documented score path would be transferred once at "
        "the end. This design should first be qualified for PLS-SVD and SIMPLS. "
        "OPLS requires fold-specific orthogonal-filter statistics, whereas "
        "nonlinear kernel PLS requires separate treatment because each fold "
        "constructs a quadratic-size Gram matrix."
    )

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    print(output_docx)


if __name__ == "__main__":
    main()
