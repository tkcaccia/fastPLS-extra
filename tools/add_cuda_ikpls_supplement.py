#!/usr/bin/env python3
"""Insert the matched CUDA fastPLS/IKPLS comparison into the supplement."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches


def insert_before(anchor, paragraph) -> None:
    anchor._p.addprevious(paragraph._p)


def add_text_before(document, anchor, text, style=None, page_break=False):
    paragraph = document.add_paragraph(style=style)
    if page_break:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run(text)
    insert_before(anchor, paragraph)
    return paragraph


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("figure_png", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()

    document = Document(args.input_docx)
    anchors = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "S10. ImageNet stress test"
    ]
    if len(anchors) != 1:
        raise RuntimeError("Expected one S10. ImageNet stress test anchor")
    anchor = anchors[0]

    add_text_before(
        document,
        anchor,
        "CUDA comparison with IKPLS",
        style="Heading 2",
        page_break=True,
    )
    add_text_before(
        document,
        anchor,
        (
            "A matched CUDA software comparison used the same float32 training and "
            "test matrices and SIMPLS component counts as Figure 1. fastPLS 0.99.65 "
            "used SIMPLS-LDA for classification and SIMPLS for multivariate "
            "regression; IKPLS 6.1.2 used improved-kernel PLS algorithm 2 through "
            "JAX 0.6.2 and converted classification response scores to labels by "
            "argmax. Because the estimators and classification heads differ, this "
            "is an end-to-end software comparison rather than an estimator-matched "
            "comparison. Data loading and file conversion occurred before the timed "
            "boundary."
        ),
    )
    add_text_before(
        document,
        anchor,
        (
            "For the eleven standard tasks, values are medians of ten isolated "
            "processes. NMR and ImageNet are single-run feasibility measurements. "
            "Cold totals include CUDA initialization or JAX compilation where "
            "incurred, host-to-device transfer, synchronization, prediction and "
            "result transfer. Warm totals repeat the same public fitting and "
            "prediction workflow in the initialized process. Host memory is the "
            "absolute peak process resident set size; device memory is the peak "
            "memory attributed by nvidia-smi to the benchmark process."
        ),
    )
    add_text_before(
        document,
        anchor,
        (
            "fastPLS had the lower cold and warm total time in every one of the "
            "eleven IKPLS-completed tasks and lower absolute peak host memory in "
            "each of those tasks. Among the nine completed classification pairs, "
            "fastPLS accuracy was higher in seven, equal in one and lower in one; "
            "these differences also reflect LDA versus argmax. IKPLS could not fit "
            "the 50-component NMR response because its coefficient path requested "
            "68.66 GiB of device memory. At 1,000 ImageNet components, IKPLS could "
            "not allocate an additional 3.74 GiB on the 16-GiB GPU, whereas "
            "fastPLS completed the compact workflow."
        ),
    )

    figure_paragraph = document.add_paragraph()
    figure_paragraph.alignment = 1
    figure_paragraph.add_run().add_picture(str(args.figure_png), width=Inches(6.45))
    insert_before(anchor, figure_paragraph)
    add_text_before(
        document,
        anchor,
        (
            "Figure S17. CUDA software comparison between fastPLS and IKPLS on "
            "matched float32 inputs. Classification uses fastPLS SIMPLS-LDA and "
            "IKPLS algorithm 2 with argmax; regression uses each implementation's "
            "multivariate response workflow. Requested component counts follow the "
            "SIMPLS contract in Table S11. Time labels show cold total with warm "
            "total in parentheses. Standard-task cells are medians of ten isolated "
            "processes; NMR and ImageNet are single-run feasibility measurements. "
            "The NVIDIA GeForce RTX 5060 Ti had 16 GiB of device memory. Gray cells "
            "retain resource failures rather than treating them as missing data."
        ),
    )

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)


if __name__ == "__main__":
    main()
