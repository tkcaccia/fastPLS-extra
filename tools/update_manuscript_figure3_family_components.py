#!/usr/bin/env python3

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SOURCE = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "documents/fastPLS_manuscript_figure4_0.99.66.docx"
)
FIGURE = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "figures/figure3_nmr_family_components.png"
)
OUTPUT = Path(
    "/Users/stefano/Documents/GPUPLS/manuscript_revision_0.99.66/"
    "documents/fastPLS_manuscript_figure3_figure4_0.99.66.docx"
)


REPLACEMENTS = {
    "At a common 165-component NMR workload, matched Linux SIMPLS required 16.759 s on CPU and 0.496 s on CUDA. The Apple CPU/Metal route completed the same float32 SIMPLS calculation in 2.520 s; cross-platform timings were not interpreted as hardware speed ratios.":
        "At the family-specific NMR workloads shown in the main analysis, matched Linux CPU/CUDA execution required 6.128/0.500 s for 100-component PLS-SVD and 1.619/0.463 s for 50-component SIMPLS. The Apple Metal route completed the corresponding calculations in 0.877 and 1.141 s; cross-platform timings were not interpreted as hardware speed ratios.",
    "Predictive selection and a separate fixed 165-component implementation comparison were not pooled.":
        "The displayed implementation comparison used 100 PLS-SVD components, an eligible value under the one-standard-error rule, and 50 SIMPLS components, the smallest eligible value. Because the families use different component counts, family contrasts are descriptive; backend comparisons remain matched within each family.",
    "At the common 165-component workload, float32 PLS-SVD required 68.985 s on the Linux CPU, 0.424 s on CUDA and 1.543 s through the Apple CPU/Metal route. The corresponding SIMPLS times were 16.759, 0.496 and 2.520 s. CPU and CUDA were measured on the same Intel/NVIDIA workstation, giving a 33.8-fold SIMPLS runtime ratio. Metal was measured on the Apple M3 and is therefore reported as a separate implementation result rather than a direct ratio to the Linux CPU. All six routes used the same prepared split, float32 inputs, rSVD controls, seed and component count.":
        "At 100 PLS-SVD components, float32 fitting plus prediction required a median of 6.128 s on the Linux CPU and 0.500 s on CUDA, a 12.3-fold same-workstation runtime ratio. The Apple CPU/Metal route completed the same 100-component calculation in 0.877 s. At 50 SIMPLS components, the corresponding medians were 1.619 s on the Linux CPU, 0.463 s on CUDA and 1.141 s through the Apple CPU/Metal route, giving a 3.5-fold same-workstation CPU/CUDA ratio. Held-out RMSD was 0.0007194 across the three PLS-SVD routes and ranged from 0.0007230 to 0.0007376 across the SIMPLS routes. Metal results are reported separately rather than as ratios to the Linux CPU. All routes used the same prepared split, float32 inputs, rSVD controls and seed; component count was fixed within each family.",
    "These rows are not pooled with the 165-component implementation comparison.":
        "These rows are not pooled with the family-specific NMR implementation comparison.",
    "The choice between PLS-SVD and SIMPLS depends on the scientific and matrix regime. PLS-SVD obtains a retained subspace in one decomposition and was faster for the extreme-response NMR workload, but its component count is response-rank limited. SIMPLS pays for sequential orthogonalization and deflation, yet can continue beyond a low response rank and supports an interpretable component path.":
        "The choice between PLS-SVD and SIMPLS depends on the scientific and matrix regime. PLS-SVD obtains a retained subspace in one decomposition, but its component count is response-rank limited. SIMPLS pays for sequential orthogonalization and deflation, yet can continue beyond a low response rank and supports an interpretable component path. In the family-specific NMR comparison, the relative runtime also depended on the retained component count: PLS-SVD used 100 components and SIMPLS used 50, so the between-family timing difference is not interpreted as an estimator-matched speed comparison.",
}


CAPTION = (
    "Figure 3. NMR prediction at family-specific component counts. PLS-SVD was "
    "evaluated at 100 components and SIMPLS at 50 components; CPU, CUDA and "
    "Metal routes are shown for each family. Panels compare median fitting-plus-"
    "prediction time with interquartile range over three isolated processes, "
    "held-out RMSD, baseline-corrected peak process RSS, CUDA device peak memory, "
    "per-spectrum RMSD and observed/predicted spectra for the held-out sample "
    "nearest the median Linux CPU SIMPLS error. The runtime axis begins at zero. "
    "CPU and CUDA were measured on the Intel/NVIDIA workstation, whereas Metal "
    "was measured on the Apple M3; cross-platform timings are not interpreted as "
    "hardware speed ratios. Input conversion is excluded, accelerator "
    "initialization and transfer are included, and all 28,355 response variables "
    "contribute to the metrics."
)


def replace_once(paragraph, old, new):
    if old not in paragraph.text:
        return False
    full = paragraph.text.replace(old, new)
    formatting = None
    if paragraph.runs:
        run = paragraph.runs[0]
        formatting = (run.bold, run.italic, run.underline)
    paragraph.clear()
    run = paragraph.add_run(full)
    if formatting is not None:
        run.bold, run.italic, run.underline = formatting
    return True


def main():
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError("The source manuscript or revised Figure 3 is missing.")
    document = Document(SOURCE)

    matched = {old: 0 for old in REPLACEMENTS}
    for paragraph in document.paragraphs:
        for old, new in REPLACEMENTS.items():
            if replace_once(paragraph, old, new):
                matched[old] += 1

    missing = [old for old, count in matched.items() if count != 1]
    if missing:
        raise RuntimeError(
            "Expected each prose replacement exactly once; mismatches: "
            + repr([(text, matched[text]) for text in missing])
        )

    caption_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.startswith("Figure 3.")
    )
    if caption_index == 0:
        raise RuntimeError("Figure 3 has no preceding image paragraph.")
    image_paragraph = document.paragraphs[caption_index - 1]
    image_paragraph.clear()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(FIGURE), width=Inches(6.65))

    caption_paragraph = document.paragraphs[caption_index]
    caption_paragraph.clear()
    caption_run = caption_paragraph.add_run(CAPTION)
    caption_run.italic = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
