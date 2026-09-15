#!/usr/bin/env python3
"""Clarify why Supplementary Tables S4 and S5 are retained."""

import argparse

from docx import Document


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def update_main(document):
    prefixes = (
        "Separate numerical checks supported interpretation",
        "Table S4 summarizes numerical agreement",
    )
    replacement = (
        "Table S4 summarizes numerical agreement with dense and family-specific "
        "reference calculations and the qualification results for the approximate "
        "rSVD routes. Table S5 addresses a different question by isolating the "
        "runtime, memory and numerical effect of each SIMPLS execution optimization "
        "within the same implementation. Together, these analyses distinguish "
        "numerical reliability from the source of computational gains."
    )
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefixes):
            replace_text(paragraph, replacement)
            return
    raise ValueError("Could not find the main-text validation paragraph")


def update_supplement(document):
    s4_caption = "Table S4. Numerical agreement and rSVD qualification summary."
    s5_caption = (
        "Table S5. Contribution of individual SIMPLS execution optimizations "
        "to runtime, host memory and numerical output."
    )
    seen = set()
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Table S4."):
            replace_text(paragraph, s4_caption)
            seen.add("S4")
        elif text.startswith("Table S5."):
            replace_text(paragraph, s5_caption)
            paragraph.paragraph_format.page_break_before = True
            seen.add("S5")
        elif text.startswith("The component-wise SIMPLS validation compares"):
            replace_text(
                paragraph,
                "Table S4 evaluates numerical reliability rather than speed. The "
                "component-wise SIMPLS comparison covers well-conditioned, "
                "collinear, rank-deficient, high-response, p<n, p>n and near-tied "
                "cases. rSVD qualification uses route-specific controls and stated "
                "numerical tolerances; meeting them is not described as exact "
                "equality. Prefix-level errors and replicate records are retained "
                "with the benchmark evidence.",
            )
        elif text.startswith("A speed-up above one favours"):
            replace_text(
                paragraph,
                "Table S5 isolates implementation effects within the same SIMPLS "
                "code base. A speed-up above one favours the optimized state, while "
                "negative RSS reduction denotes a larger measured process-memory "
                "increment. The results show that individual optimizations are "
                "matrix-regime dependent and explain why the public implementation "
                "combines them selectively rather than claiming universal gains.",
            )
    missing = {"S4", "S5"}.difference(seen)
    if missing:
        raise ValueError(f"Could not find captions: {sorted(missing)}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--main-input", required=True)
    parser.add_argument("--supplement-input", required=True)
    parser.add_argument("--main-output", required=True)
    parser.add_argument("--supplement-output", required=True)
    args = parser.parse_args()

    main_document = Document(args.main_input)
    update_main(main_document)
    main_document.save(args.main_output)

    supplement = Document(args.supplement_input)
    update_supplement(supplement)
    supplement.save(args.supplement_output)


if __name__ == "__main__":
    main()
