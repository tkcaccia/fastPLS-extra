#!/usr/bin/env python3

"""Separate CMPB biomedical evidence from the future JSS software study."""

from copy import deepcopy
from pathlib import Path
import csv
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/stefano/Documents/GPUPLS")
INPUT_MAIN = ROOT / "document_work/cmpb_latest_20260915/fastPLS_CMPB_manuscript_updated.docx"
INPUT_SUPP = ROOT / "document_work/cmpb_latest_20260915/fastPLS_CMPB_supplement_updated.docx"
INPUT_JSS = ROOT / "manuscript_style_revision_20260914/documents/fastPLS_JSS_future_draft.docx"
OUTPUT = ROOT / "document_work/cmpb_jss_restructure_20260915/documents"
FIGURES = ROOT / "document_work/cmpb_jss_restructure_20260915/figures"
PATH_FIGURES = ROOT / "document_work/cmpb_jss_restructure_20260915/component_paths_current/component_path_plots"
CONTRACT = ROOT / "fastPLS-extra/benchmark/gpu_cross_validation/selected_component_contract.csv"


def paragraph_text(element, document):
    return " ".join(Paragraph(element, document).text.split())


def body_elements(document):
    return list(document.element.body.iterchildren())


def find_paragraph_element(document, prefix):
    for element in body_elements(document):
        if element.tag == qn("w:p") and paragraph_text(element, document).startswith(prefix):
            return element
    raise ValueError(f"Paragraph not found: {prefix}")


def find_paragraph(document, prefix):
    return Paragraph(find_paragraph_element(document, prefix), document)


def set_paragraph(document, prefix, text, style=None):
    paragraph = find_paragraph(document, prefix)
    paragraph.clear()
    paragraph.add_run(text)
    if style:
        paragraph.style = style
    return paragraph


def delete_paragraph(document, prefix):
    element = find_paragraph_element(document, prefix)
    element.getparent().remove(element)


def delete_section(document, start_prefix, end_prefix):
    start = find_paragraph_element(document, start_prefix)
    end = find_paragraph_element(document, end_prefix)
    elements = body_elements(document)
    first = elements.index(start)
    last = elements.index(end)
    for element in elements[first:last]:
        element.getparent().remove(element)


def move_section_before(document, start_prefix, end_prefix, target_prefix):
    start = find_paragraph_element(document, start_prefix)
    end = find_paragraph_element(document, end_prefix)
    target = find_paragraph_element(document, target_prefix)
    elements = body_elements(document)
    block = elements[elements.index(start):elements.index(end)]
    for element in block:
        target.addprevious(element)


def nearest_previous_drawing(element):
    previous = element.getprevious()
    while previous is not None:
        if previous.tag == qn("w:p") and previous.xpath(".//w:drawing"):
            return previous
        if previous.tag == qn("w:p") and paragraph_text(previous, element.getparent()):
            return None
        previous = previous.getprevious()
    return None


def delete_caption_object(document, prefix, remove_following_table=True):
    caption = find_paragraph_element(document, prefix)
    drawing = caption.getprevious()
    while drawing is not None and drawing.tag == qn("w:p"):
        if drawing.xpath(".//w:drawing"):
            drawing.getparent().remove(drawing)
            break
        if paragraph_text(drawing, document):
            break
        previous = drawing.getprevious()
        drawing.getparent().remove(drawing)
        drawing = previous
    following = caption.getnext()
    caption.getparent().remove(caption)
    if remove_following_table:
        while following is not None and following.tag == qn("w:p") and not paragraph_text(following, document):
            next_element = following.getnext()
            following.getparent().remove(following)
            following = next_element
        if following is not None and following.tag == qn("w:tbl"):
            following.getparent().remove(following)


def table_after(document, caption_prefix):
    element = find_paragraph_element(document, caption_prefix).getnext()
    while element is not None and element.tag != qn("w:tbl"):
        element = element.getnext()
    if element is None:
        raise ValueError(f"Table not found after {caption_prefix}")
    return Table(element, document)


def filter_table_rows(table, column, keep):
    header = [cell.text.strip() for cell in table.rows[0].cells]
    index = header.index(column)
    for row in list(table.rows[1:]):
        if row.cells[index].text.strip() not in keep:
            table._tbl.remove(row._tr)


def format_table(table, size=7.0):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)
                    if row_index == 0:
                        run.font.bold = True


def replace_table(document, caption_prefix, headers, rows):
    old = table_after(document, caption_prefix)
    new = document.add_table(rows=1, cols=len(headers))
    if old.style:
        new.style = old.style
    for index, value in enumerate(headers):
        new.rows[0].cells[index].text = str(value)
    for values in rows:
        cells = new.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    old._tbl.addprevious(new._tbl)
    old._tbl.getparent().remove(old._tbl)
    format_table(new)


def replace_picture_before_caption(document, caption_prefix, image_path, width=6.5):
    caption = find_paragraph_element(document, caption_prefix)
    element = caption.getprevious()
    while element is not None:
        if element.tag == qn("w:p") and element.xpath(".//w:drawing"):
            paragraph = Paragraph(element, document)
            paragraph.clear()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(width))
            return
        if element.tag == qn("w:p") and paragraph_text(element, document):
            break
        element = element.getprevious()
    raise ValueError(f"Image not found before {caption_prefix}")


def replace_references(document, mapping):
    patterns = sorted(mapping, key=len, reverse=True)
    placeholders = {key: f"@@REF{index}@@" for index, key in enumerate(patterns)}
    for paragraph in document.paragraphs:
        text = paragraph.text
        changed = False
        for key in patterns:
            pattern = re.escape(key) + r"(?!\d)"
            if re.search(pattern, text):
                text = re.sub(pattern, placeholders[key], text)
                changed = True
        if changed:
            for key in patterns:
                text = text.replace(placeholders[key], mapping[key])
            paragraph.clear()
            paragraph.add_run(text)


def insert_before(anchor, element):
    anchor.addprevious(element)


def add_paragraph_before(document, anchor, text, style=None):
    paragraph = document.add_paragraph(text)
    if style:
        paragraph.style = style
    insert_before(anchor, paragraph._p)
    return paragraph


def add_table_copy_before(document, anchor, source_table, size=7.0):
    element = deepcopy(source_table._tbl)
    insert_before(anchor, element)
    table = Table(element, document)
    format_table(table, size=size)
    return table


def add_picture_before(document, anchor, path, width=6.4):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    insert_before(anchor, paragraph._p)


def fixed_component_rows():
    labels = {
        "ccle": "CCLE", "cifar100": "CIFAR-100", "gtex_v8": "GTEx v8",
        "metref": "MetRef", "retina": "Retina", "tabula": "Tabula Muris",
        "tcga_brca": "TCGA-BRCA",
        "tcga_hnsc_methylation": "TCGA-HNSC methylation",
        "tcga_pan_cancer": "TCGA Pan-Cancer", "cbmc_citeseq": "CBMC CITE-seq",
        "prism": "PRISM", "nmr": "NMR", "imagenet": "ImageNet/DINOv2"
    }
    order = list(labels)
    values = {}
    with CONTRACT.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            values.setdefault(row["dataset"], {})[row["family"]] = row["selected_ncomp"]
    return [
        [labels[key], values[key]["plssvd"], values[key]["simpls"],
         values[key]["opls"], values[key]["kernelpls"]]
        for key in order
    ]


def build_cmpb_supplement():
    document = Document(INPUT_SUPP)
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        if section.page_width < section.page_height:
            continue
        section.page_width, section.page_height = section.page_height, section.page_width

    delete_section(document, "S2. Algorithms", "S4. CPU and GPU")
    set_paragraph(document, "S4. CPU and GPU", "S3. CPU and CUDA comparisons", "Heading 1")
    filter_table_rows(table_after(document, "Table S4."), "Platform", {"CUDA"})
    filter_table_rows(table_after(document, "Table S5."), "Platform", {"CUDA"})
    timing_table = table_after(document, "Table S4.")
    timing_headers = {
        cell.text.strip(): index
        for index, cell in enumerate(timing_table.rows[0].cells)
    }
    for row in timing_table.rows[1:]:
        cells = row.cells
        if (cells[timing_headers["Dataset"]].text.strip() == "NMR" and
                cells[timing_headers["Family"]].text.strip() == "PLS-SVD"):
            replacements = {
                "A": "100", "CPU s": "6.128", "Accel. s": "0.500",
                "Time ratio": "12.26", "CPU metric": "7.195e-04",
                "Accel. metric": "7.194e-04"
            }
            for header, value in replacements.items():
                cells[timing_headers[header]].text = value
    memory_table = table_after(document, "Table S5.")
    memory_headers = {
        cell.text.strip(): index
        for index, cell in enumerate(memory_table.rows[0].cells)
    }
    for row in memory_table.rows[1:]:
        cells = row.cells
        if (cells[memory_headers["Dataset"]].text.strip() == "NMR" and
                cells[memory_headers["Family"]].text.strip() == "PLS-SVD"):
            replacements = {
                "A": "100", "CPU RSS MiB": "593.1",
                "Accel. RSS MiB": "679.1", "RSS ratio": "1.15",
                "GPU peak MiB": "564.0"
            }
            for header, value in replacements.items():
                cells[memory_headers[header]].text = value
    set_paragraph(
        document, "Table S4.",
        "Table S4. Fitting and prediction time and predictive performance for matched float32 CPU/CUDA pairs. A is the number of components; the predictive measure is accuracy for classification and RMSD for regression."
    )
    set_paragraph(
        document, "Table S5.",
        "Table S5. Host-memory increase and CUDA device allocation for the CPU/CUDA comparisons in Table S4. RSS is reported in MiB."
    )
    delete_caption_object(document, "Table S6.")
    delete_caption_object(document, "Table S7.")

    set_paragraph(document, "S5. Component", "S4. Component counts and prediction paths", "Heading 1")
    delete_caption_object(document, "Table S9.")
    if any(p.text.strip().startswith("Table S9 reports") for p in document.paragraphs):
        delete_paragraph(document, "Table S9 reports")
    set_paragraph(
        document, "Table S8.",
        "Table S8. Component counts retained for the PLS-SVD, SIMPLS, OPLS and linear kernel-PLS benchmark workflows. PLS-SVD counts for GTEx v8, MetRef and TCGA Pan-Cancer are constrained by response rank."
    )
    replace_table(
        document, "Table S8.",
        ["Dataset", "PLS-SVD", "SIMPLS", "OPLS", "Linear kernel PLS"],
        fixed_component_rows()
    )
    for number, dataset in enumerate([
        "ccle", "cifar100", "gtex_v8", "metref", "retina", "tabula",
        "tcga_brca", "tcga_hnsc_methylation", "tcga_pan_cancer",
        "cbmc_citeseq", "prism"
    ], start=1):
        replace_picture_before_caption(
            document, f"Figure S{number}.", PATH_FIGURES / f"component_path_{dataset}.png"
        )
        caption = find_paragraph(document, f"Figure S{number}.")
        text = caption.text
        text = text.replace(
            "colours and point shapes identify the hardware in the legend.",
            "colours and point shapes distinguish Linux CPU and CUDA."
        ).replace(
            "Dotted lines mark training-selected counts.",
            "Dotted lines mark the retained benchmark component counts."
        )
        caption.clear()
        caption.add_run(text)

    delete_paragraph(document, "S6. NMR")
    delete_caption_object(document, "Table S10.")
    replace_picture_before_caption(document, "Figure S12.", FIGURES / "figureS12_nmr_test_component_path.png")
    set_paragraph(
        document, "Figure S12.",
        "Figure S12. NMR prediction and computational cost across component counts on Linux CPU and CUDA. Each model was fitted on all 1,200 training spectra and evaluated on the fixed 321 test spectra. Rows show held-out RMSD, fitting and prediction time, and host-RSS increase; columns show the four PLS families. Dotted lines mark the retained benchmark counts: 100 for PLS-SVD and 50 for SIMPLS, OPLS and kernel PLS."
    )
    delete_caption_object(document, "Table S11.")
    if any(p.text.strip().startswith("Table S11 and Figure 3") for p in document.paragraphs):
        delete_paragraph(document, "Table S11 and Figure 3")
    delete_caption_object(document, "Figure S13.", remove_following_table=False)
    delete_caption_object(document, "Figure S14.", remove_following_table=False)

    set_paragraph(
        document, "Each setting used three fresh",
        "Each setting used three fresh processes and the fastPLS 0.99.66 R interface. Time includes CUDA initialization, fitting and test prediction. CPU/CUDA results are paired on the Intel/NVIDIA workstation. A CPU/CUDA time ratio above one indicates faster CUDA execution. Memory values are increases in process RSS above the recorded baseline, including runtime allocations; device allocation may also include CUDA context, library and allocator storage."
    )
    delete_paragraph(document, "The comparison recorded 360")
    if any(p.text.strip().startswith("Each solver pair used") for p in document.paragraphs):
        delete_paragraph(document, "Each solver pair used")
    set_paragraph(
        document, "Component selection for the eleven",
        "The retained benchmark component counts are listed in Table S6. Figures S2-S12 show the corresponding test-set paths for the general datasets on Linux CPU and CUDA. Classification paths report argmax and LDA where both were evaluated; regression paths report RMSD. The dotted line in each family panel identifies the retained count."
    )
    set_paragraph(
        document, "The NMR selection grid was",
        "The NMR path used the supplied training/test split and component counts from 1 to 300. Figure S12 reports held-out RMSD, fitting and prediction time, and host-memory increase. The retained comparison uses 100 PLS-SVD components and 50 components for SIMPLS, OPLS and kernel PLS."
    )

    delete_section(document, "S7. CPU thread", "S9. Independent")
    set_paragraph(document, "S9. Independent", "S2. Independent implementations", "Heading 1")
    set_paragraph(document, "Figure S18 contains", "Figure S14 contains the argmax classification results for matched Linux CPU/CUDA workflows. Inputs, component counts and three-process repetition were retained across the four PLS families.")
    if any(p.text.strip().startswith("Figure S18 uses") for p in document.paragraphs):
        delete_paragraph(document, "Figure S18 uses")
    replace_picture_before_caption(document, "Figure S18.", FIGURES / "figure_s14_argmax_cuda_cmpb.png")
    set_paragraph(
        document, "Figure S18.",
        "Figure S18. Argmax classification with Linux CPU and CUDA execution. A: CPU/CUDA time ratio, with values above one favouring CUDA. B: CUDA/CPU host-RSS increment ratio, with values below one indicating less process-memory growth. Values are medians of three processes and include initialization, transfers, fitting and prediction."
    )
    if any(p.text.strip().startswith("Figure S19 provides") for p in document.paragraphs):
        delete_paragraph(document, "Figure S19 provides")
    delete_caption_object(document, "Figure S19.", remove_following_table=False)

    set_paragraph(document, "S10. Cross-validation", "S5. Cross-validation workflow timing", "Heading 1")
    replace_picture_before_caption(document, "Figure S20.", FIGURES / "supp_gpu_cv_vs_fit_prediction_cmpb.png")
    set_paragraph(
        document, "Figure S20.",
        "Figure S20. Complete ten-fold cross-validation time divided by one full-training fit plus fixed-test prediction time on Linux CPU and NVIDIA CUDA. Columns show PLS-SVD, SIMPLS, OPLS and linear kernel PLS. Ratios above one indicate that cross-validation required more time."
    )
    set_paragraph(
        document, "All 192 paired workflows",
        "All Linux CPU and CUDA workflows shown in Figure S20 completed. The median ratios were 5.85 on Linux CPU and 2.67 on CUDA, with ranges of 2.27-16.90 and 1.11-34.30, respectively. These ratios are not expected to equal ten because a validation fold uses approximately 90% of the training observations, whereas the comparator fits the full training partition and predicts a separate test set; cross-validation also constructs complete out-of-fold predictions and scores."
    )

    set_paragraph(document, "S11. Large", "S6. Large embedding matrices", "Heading 1")
    set_paragraph(document, "S12. Computational", "S7. Computational environment and reproducibility", "Heading 1")
    set_paragraph(
        document, "BLAS (Basic",
        "BLAS (Basic Linear Algebra Subprograms) and LAPACK specify common matrix operations and decompositions. Linux performance builds used OpenBLAS. If OpenBLAS is absent, fastPLS can be built with R's BLAS/LAPACK, but such a build is not the CPU baseline reported here. The linked library affects speed and threading and can be inspected with fastPLS_blas()."
    )
    env_table = table_after(document, "Table S18.")
    for row in list(env_table.rows[1:]):
        text = " ".join(cell.text for cell in row.cells).lower()
        if any(token in text for token in ("apple", "mac", "metal")):
            env_table._tbl.remove(row._tr)
    set_paragraph(
        document, "The benchmark scripts record",
        "The benchmark scripts record method, component count, precision, randomized controls, implementation, elapsed time, predictive measures and completion status. The CMPB analyses use the Linux CPU and CUDA calculations identified with each figure or table. Interface checks and other platform-specific implementation studies are reported separately in the software article."
    )

    move_section_before(
        document, "S2. Independent implementations",
        "S5. Cross-validation workflow timing", "S3. CPU and CUDA comparisons"
    )
    move_section_before(
        document, "Argmax classification across accelerator backends",
        "S3. CPU and CUDA comparisons", "S5. Cross-validation workflow timing"
    )
    replace_references(document, {
        "Table S14": "Table S1", "Table S15": "Table S2",
        "Table S16": "Table S3", "Table S4": "Table S4",
        "Table S5": "Table S5", "Table S8": "Table S6",
        "Table S17": "Table S7", "Table S18": "Table S8",
        "Figure S17": "Figure S1", "Figure S18": "Figure S14",
        "Figure S20": "Figure S15", "Figure S12": "Figure S13",
        "Figure S11": "Figure S12", "Figure S10": "Figure S11",
        "Figure S9": "Figure S10", "Figure S8": "Figure S9",
        "Figure S7": "Figure S8", "Figure S6": "Figure S7",
        "Figure S5": "Figure S6", "Figure S4": "Figure S5",
        "Figure S3": "Figure S4", "Figure S2": "Figure S3",
        "Figure S1": "Figure S2", "[34]": "[33]"
    })
    for prefix, title in {
        "Table S4.": "Table S4. Fitting and prediction time and predictive performance for matched float32 CPU/CUDA pairs. A is the number of components; the predictive measure is accuracy for classification and RMSD for regression.",
        "Table S5.": "Table S5. Host-memory increase and CUDA device allocation for the CPU/CUDA comparisons in Table S4. RSS is reported in MiB.",
        "Table S6.": "Table S6. Component counts retained for the benchmark workflows. PLS-SVD counts for GTEx v8, MetRef and TCGA Pan-Cancer are constrained by response rank.",
    }.items():
        set_paragraph(document, prefix, title)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "fastPLS_CMPB_supplement_restructured.docx"
    document.save(path)
    return path


def build_cmpb_main():
    document = Document(INPUT_MAIN)
    find_paragraph(document, "2.2 Classification").style = "Heading 2"
    set_paragraph(
        document, "Here we describe fastPLS",
        "Here we describe fastPLS, an implementation that reduces repeated computation and model storage during SIMPLS fitting and prediction. Its main features are reusable matrix products, compact prediction factors and specialized calculations for many-response and multiclass problems. We evaluate the implementation against independent software and across CPU and GPU hardware [21], with NMR reconstruction as the principal biomedical application. The shared C++ core is available through R, Python and MATLAB."
    )
    set_paragraph(
        document, "SIMPLS extracts successive",
        "SIMPLS extracts successive directions while orthogonalizing predictor loadings and deflating the cross-covariance [8]. fastPLS reuses each deflation product and retains the predictor weights and response loadings needed for prediction (Algorithm 1). For responses with many columns, calculations in sample space avoid repeatedly multiplying by the full response matrix. For classification, class sums provide the required cross-products without constructing a dense indicator matrix. Some large problems use several candidate directions computed from one deflated state; these candidates then undergo sequential SIMPLS updates. This block calculation is an approximate variant and can differ from solving for a new leading direction after every component. The component-path and held-out results therefore evaluate the implemented execution directly."
    )
    set_paragraph(
        document, "OPLS first filters",
        "OPLS first filters response-orthogonal predictor variation and then fits a predictive PLS model. Linear kernel PLS uses the predictor representation directly, whereas nonlinear kernel PLS constructs and centres an n by n kernel matrix. The quadratic storage of the nonlinear Gram matrix limits the sample sizes that can be fitted. The principal rSVD benchmark setting used 32 oversampling directions, five power iterations and seed 123."
    )
    set_paragraph(
        document, "The backend experiment",
        "The backend experiment paired CPU and CUDA on an Intel Core i7-13700 workstation with 32 GiB RAM and an NVIDIA GeForce RTX 5060 Ti with 16 GiB device memory [21]. All four PLS families used float32 inputs; classification used argmax to avoid changing the prediction head between backends. Three fresh processes were requested per setting. Timing included initialization, transfers, fitting, prediction and synchronization, but excluded data loading and conversion."
    )
    delete_paragraph(document, "A separate precision comparison")
    set_paragraph(
        document, "Additional experiments compared",
        "Component-path experiments evaluated predictive performance, fitting and prediction time, and process memory across the retained component grids. The fixed workloads used for the independent-software comparison were kept separate from these path analyses."
    )
    set_paragraph(
        document, "NMR preprocessing",
        "NMR preprocessing and the supplied split of 1,200 training and 321 test spectra are described in Supplementary Section S1. Component paths were evaluated on the fixed held-out set. The principal comparison used 100 PLS-SVD components and 50 SIMPLS components. The deposited fastsimpls PLS-SVD implementation from the NMR study [7] was rerun at its published 165-component setting with IRLBA and float64 on the same prepared data."
    )
    delete_paragraph(document, "Numerical checks compared")
    delete_paragraph(document, "The dominant operations")
    delete_paragraph(document, "In the controlled solver comparison")
    delete_paragraph(document, "The effect of requesting four CPU cores")
    set_paragraph(
        document, "3.2 Solver and backend execution",
        "3.2 Component paths and CUDA execution",
        "Heading 2"
    )
    set_paragraph(
        document, "Figure 1.",
        "Figure 1. Prediction, runtime and memory for PLS software on an Intel Core i7-13700 workstation with one effective CPU thread. fastPLS uses SIMPLS-LDA for classification and SIMPLS regression. A and B: test accuracy and RMSD. C and D: fitting and prediction time. E and F: absolute peak process RSS. Time and memory share colour scales across classification and regression. fastPLS 0.99.65 and IKPLS used float32; independent-software settings and repetition counts are given in Tables S1-S3. Standard fastPLS and IKPLS tasks used ten processes; ImageNet used one. rSVD used oversampling 32, five power iterations and seed 123. Failed indicates an attempted calculation that did not complete; NE indicates not evaluated."
    )
    set_paragraph(
        document, "CUDA completed all 52",
        "CUDA completed all 52 paired calculations and was faster than the CPU in 15 (Figure 2): all four families on CIFAR-100, NMR and ImageNet, and SIMPLS, OPLS and linear kernel PLS on PRISM. CPU/CUDA runtime ratios ranged from 40.5 to 193.5 for ImageNet and from 3.4 to 19.2 for NMR. Supplementary Tables S4 and S5 give the paired performance and memory measurements. Supplementary Figures S2-S13 show the component paths and Table S6 records the retained counts. Figure S14 presents the argmax classification comparison separately. Ten-fold cross-validation relative to one fit and held-out prediction is summarized in Figure S15."
    )
    caption = find_paragraph(document, "Figure 2.")
    caption.clear()
    caption.add_run(
        "Figure 2. CPU and CUDA performance on the same Intel Core i7-13700/RTX 5060 Ti workstation. All 13 datasets and four PLS families used float32; classification used argmax. A: CPU time divided by CUDA time, with values above 1 indicating faster CUDA execution. B: CUDA host-RSS increment divided by the CPU increment, with values below 1 indicating less host-memory growth. Three processes were requested per setting using fastPLS 0.99.66. Timing includes initialization, transfers, synchronization, fitting and prediction. RSS increments include runtime allocations as well as numerical workspaces."
    )
    set_paragraph(
        document, "The NMR selection analysis",
        "The NMR component paths showed how held-out error and computational cost changed across the evaluated counts (Supplementary Figure S13). For the principal comparison, PLS-SVD used 100 components and SIMPLS used 50, matching the retained benchmark specification rather than claiming a unique optimum."
    )
    set_paragraph(
        document, "At 100 components",
        "At 100 components, PLS-SVD required 6.128 s on the Intel CPU and 0.500 s on CUDA, a 12.3-fold difference on the same workstation. Test RMSD was 0.0007195 and 0.0007194, respectively. At 50 components, SIMPLS required 1.619 s on the CPU and 0.463 s on CUDA, with RMSD 0.0007376 and 0.0007230 (Figure 3). The deposited 165-component PLS-SVD/IRLBA workflow required 457.511 s and achieved RMSD 0.000785806. Its different solver, precision and component count prevent attribution of the full timing difference to implementation alone."
    )
    replace_picture_before_caption(document, "Figure 3.", FIGURES / "figure3_nmr_family_components.png", 6.6)
    set_paragraph(
        document, "Figure 3.",
        "Figure 3. NMR spectral prediction with fastPLS and the deposited PLS-SVD implementation. fastPLS used float32 rSVD with 100 PLS-SVD or 50 SIMPLS components; the deposited fastsimpls model used float64 IRLBA with 165 components (grey). A: median fitting and prediction time and interquartile range from three processes. B: test RMSD over all 28,355 response intensities. C: peak host-RSS increment and CUDA device allocation. D: per-spectrum RMSD. E and F: observed and predicted spectra nearest the median CPU SIMPLS error, over 12-0 and 1.7-0.5 ppm. CPU and CUDA measurements were paired on the Intel/NVIDIA workstation. Conversion was excluded from timing, whereas CUDA initialization and transfers were included."
    )
    set_paragraph(
        document, "The separate scikit-learn",
        "The separate scikit-learn feasibility test completed a 50-component float64 NMR model in 107.5 s, with RMSD 0.000745 and peak process RSS of 9,180 MiB (Supplementary Table S1). This result uses a different software and precision setting from the fastPLS comparisons in Figure 3."
    )
    set_paragraph(
        document, "The ImageNet experiment fitted",
        find_paragraph(document, "The ImageNet experiment fitted").text
            .replace("Table S17", "Table S7")
            .replace("Table S18", "Table S8")
    )
    set_paragraph(
        document, "The benefit was limited",
        "The benefit was limited on several small or low-component tasks, where IKPLS was faster. Allocations, interface calls and small linear algebra operations can outweigh savings in larger matrix products. Calculating cross-covariance products without storing the matrix also requires repeated data passes, which can increase runtime. PLS-SVD uses one truncated decomposition, whereas SIMPLS supports a sequential path that can extend beyond response rank. Their different component counts in Figure 3 address practical prediction rather than isolating algorithmic speed."
    )
    set_paragraph(
        document, "GPU gains were largest",
        "GPU gains were largest when matrix products outweighed setup and transfer costs. CUDA was useful for the many-response NMR task and million-sample ImageNet task, while the CPU usually completed smaller fits sooner. These measurements support accelerator use for sufficiently large workloads; they do not establish a hardware-independent speed-up."
    )
    set_paragraph(
        document, "Single precision makes",
        "Single precision reduced the storage of the benchmark matrices and enabled the largest calculations, but its practical value depends on the complete workflow rather than element size alone. Numerical precision and platform capability are examined separately in the software study."
    )
    replace_references(document, {
        "Table S14": "Table S1", "Table S15": "Table S2",
        "Table S16": "Table S3", "Table S17": "Table S7",
        "Table S18": "Table S8", "Figure S17": "Figure S1"
    })
    delete_paragraph(document, "[22] Apple Inc.")
    for paragraph in document.paragraphs:
        text = paragraph.text
        updated = re.sub(
            r"\[(2[3-9]|3[0-4])\]",
            lambda match: f"[{int(match.group(1)) - 1}]",
            text
        )
        if updated != text:
            paragraph.clear()
            paragraph.add_run(updated)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "fastPLS_CMPB_manuscript_restructured.docx"
    document.save(path)
    return path


def build_jss():
    source = Document(INPUT_SUPP)
    document = Document(INPUT_JSS)
    anchor = find_paragraph_element(document, "8. Reproducibility")

    add_paragraph_before(document, anchor, "8. Extended algorithms and platform evaluation", "Heading 1")
    add_paragraph_before(document, anchor, "8.1 Computational requirements", "Heading 2")
    add_paragraph_before(
        document, anchor,
        "Let n denote training observations, p predictors, q responses, A retained components and l the randomized sketch width. An explicit route stores the cross-covariance X'Y, whereas an implicit route applies it to narrow matrices without storing it. The latter reduces storage but can require additional passes through the input. Table 7 summarizes dominant operations and residency."
    )
    add_paragraph_before(document, anchor, "Table 7. Dominant operations, storage and residency in the shared implementation.")
    add_table_copy_before(document, anchor, source.tables[0])
    for number, caption, table_index in [
        (3, "OPLS execution used by fastPLS.", 1),
        (4, "Linear and nonlinear kernel-PLS execution.", 2),
        (5, "Compiled cross-validation shared by PLS-SVD, SIMPLS, OPLS and kernel PLS.", 3),
    ]:
        add_paragraph_before(document, anchor, f"Algorithm {number}. {caption}")
        add_table_copy_before(document, anchor, source.tables[table_index], size=7.5)
    add_paragraph_before(
        document, anchor,
        "The compiled validation workflow constructs one maximal component path per fold and obtains exact fold statistics from full-data totals minus held-out contributions. This reduces repeated centering, cross-product construction, allocation and prefix refitting while preserving the supplied folds and grouping constraints."
    )

    add_paragraph_before(document, anchor, "8.2 Numerical precision and execution routes", "Heading 2")
    add_paragraph_before(document, anchor, "Table 8. Prediction differences across hardware and numerical precision, relative to Linux CPU float64.")
    add_table_copy_before(document, anchor, source.tables[8], size=6.2)
    add_paragraph_before(document, anchor, "Table 9. Numerical precision and execution routes used in the benchmark studies.")
    add_table_copy_before(document, anchor, source.tables[15], size=7.0)
    add_paragraph_before(
        document, anchor,
        "These measurements distinguish data representation from complete-process memory and runtime. Accelerator routes include initialization, data transfer and synchronization unless explicitly identified as repeated execution."
    )

    add_paragraph_before(document, anchor, "8.3 CPU threading and Metal execution", "Heading 2")
    add_paragraph_before(document, anchor, "Table 10. One- versus four-core CPU requests across the benchmark datasets and PLS families.")
    add_table_copy_before(document, anchor, source.tables[14], size=6.7)
    add_picture_before(document, anchor, ROOT / "fastPLS_results_local/multicore_0.99.65_20260913/macos/figureS15_multicore_runtime_ratio.png")
    add_paragraph_before(document, anchor, "Figure 2. Runtime with one- and four-core CPU requests on Linux/OpenBLAS and macOS/Accelerate. Values above one favour the four-core request.")
    add_picture_before(document, anchor, ROOT / "fastPLS_results_local/multicore_0.99.65_20260913/macos/figureS16_multicore_absolute_rss_ratio.png")
    add_paragraph_before(document, anchor, "Figure 3. Peak process memory with one- and four-core CPU requests. Values below one indicate lower memory with the four-core request.")
    add_picture_before(document, anchor, ROOT / "manuscript_revision_0.99.66/figures/figureS19_metal_backend_runtime.png")
    add_paragraph_before(document, anchor, "Figure 4. CPU and hybrid Metal fitting-and-prediction ratios on the Apple M3 across 13 datasets and four PLS families.")
    add_picture_before(document, anchor, FIGURES / "figure_jss_argmax_metal.png")
    add_paragraph_before(document, anchor, "Figure 5. Float32 argmax classification on Mac CPU and hybrid Metal, with matched component counts and outputs.")
    add_picture_before(document, anchor, FIGURES / "jss_metal_cv_vs_fit_prediction.png")
    add_paragraph_before(document, anchor, "Figure 6. Ten-fold cross-validation cost relative to one fit plus held-out prediction on Mac CPU and hybrid Metal.")

    add_paragraph_before(document, anchor, "8.4 NMR route measurements", "Heading 2")
    add_paragraph_before(document, anchor, "Table 11. NMR measurements across CPU, CUDA, Metal and the deposited PLS-SVD workflow.")
    add_table_copy_before(document, anchor, source.tables[13], size=6.5)
    add_paragraph_before(
        document, anchor,
        "The NMR table is retained here as a route-level software comparison. The CMPB paper uses the Linux CPU/CUDA subset for its biomedical analysis, while this software study also examines the hybrid Metal route and precision-specific execution."
    )

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        replacements = {
            "8. Reproducibility and extension": "9. Reproducibility and extension",
            "9. Discussion": "10. Discussion",
            "10. Availability": "11. Availability",
            "11. Conclusion": "12. Conclusion",
        }
        if text in replacements:
            paragraph.clear()
            paragraph.add_run(replacements[text])

    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "fastPLS_JSS_future_draft_extended.docx"
    document.save(path)
    return path


if __name__ == "__main__":
    print(build_cmpb_main())
    print(build_cmpb_supplement())
    print(build_jss())
