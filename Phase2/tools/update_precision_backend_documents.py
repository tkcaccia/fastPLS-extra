#!/usr/bin/env python3
"""Insert the cross-platform precision-concordance evidence into the paper."""

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


DATASETS = {
    "ccle": "CCLE",
    "cifar100": "CIFAR-100",
    "gtex_v8": "GTEx v8",
    "metref": "MetRef",
    "retina": "Retina",
    "tabula": "Tabula Muris",
    "tcga_brca": "TCGA-BRCA",
    "tcga_hnsc_methylation": "TCGA-HNSC methylation",
    "tcga_pan_cancer": "TCGA Pan-Cancer",
    "imagenet": "ImageNet/DINOv2",
    "cbmc_citeseq": "CBMC CITE-seq",
    "prism": "PRISM",
    "nmr": "NMR",
}
FAMILIES = {
    "plssvd": "PLS-SVD",
    "simpls": "SIMPLS",
    "opls": "OPLS",
    "kernelpls": "linear kernel PLS",
}
ROUTE_FIELDS = (
    "linux_cpu_float64", "linux_cpu_float32", "cuda_float64",
    "cuda_float32", "mac_cpu_float64", "mac_cpu_float32", "metal_float32",
)


def rows(path):
    with Path(path).open(newline="") as stream:
        return list(csv.DictReader(stream))


def set_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def find_table_after_caption(document, prefix):
    found = False
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip().startswith(prefix):
                found = True
        elif isinstance(child, CT_Tbl) and found:
            return Table(child, document)
    raise ValueError(f"Could not locate table after {prefix}")


def replace_figure_before_caption(document, prefix, image_path):
    paragraphs = list(document.paragraphs)
    caption_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip().startswith(prefix)
    )
    drawing = next(
        paragraph for paragraph in reversed(paragraphs[:caption_index])
        if paragraph._p.xpath(".//w:drawing")
    )
    drawing.clear()
    drawing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing.add_run().add_picture(str(image_path), width=Inches(6.45))


def format_metric(value, metric):
    if value in (None, ""):
        return "NE"
    number = float(value)
    if metric == "accuracy":
        return f"{number:.4f}"
    if abs(number) < 0.01:
        return f"{number:.3e}"
    if abs(number) < 100:
        return f"{number:.5f}"
    return f"{number:.2f}"


def table_rows(matrix):
    output = []
    for row in matrix:
        maximum = float(row["maximum_change"])
        if row["task_type"] == "classification":
            maximum_text = f"{100 * maximum:.2f} pp"
        else:
            maximum_text = f"{100 * maximum:.2f}%"
        output.append([
            DATASETS[row["dataset"]],
            FAMILIES[row["family"]],
            row["ncomp"],
            row["metric_name"],
            *(format_metric(row[field], row["metric_name"]) for field in ROUTE_FIELDS),
            maximum_text,
        ])
    return output


def replace_table(table, headers, values):
    for body_row in list(table.rows)[1:]:
        table._tbl.remove(body_row._tr)
    if len(table.rows[0].cells) != len(headers):
        raise ValueError("Table S9 column count changed unexpectedly")
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(5.6)
                run.bold = True
    for values_row in values:
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column, (cell, value) in enumerate(zip(row.cells, values_row)):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column < 2 else
                    WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.size = Pt(5.3)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def update_main(document):
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("A supporting float32-versus-float64 comparison"):
            set_text(
                paragraph,
                "A paired numerical audit compared Linux CPU, Mac CPU, "
                "CUDA and Metal results in float32 and, where supported, "
                "float64. It used all 13 datasets and all four PLS families "
                "at the component counts used for the backend benchmark. "
                "Classification used argmax so that the comparison isolated "
                "the PLS numerical route; regression used RMSD. Metal was "
                "evaluated in float32 only because Apple GPUs do not provide "
                "the native float64 arithmetic required by this route."
            )
        elif text.startswith("Accelerator benefit was conditional"):
            set_text(
                paragraph,
                text + " In the corresponding metric audit, all 236 "
                "nonreference classification comparisons differed from the "
                "Linux CPU float64 result by no more than 0.22 percentage "
                "points. Sixty-nine of 72 regression comparisons differed "
                "in RMSD by less than 1%, and all differed by less than "
                "2.91%. The three larger relative changes were confined to "
                "NMR: CUDA float32 SIMPLS and linear kernel PLS, and Mac CPU "
                "float32 PLS-SVD. Full metric values and signed differences "
                "are reported in Supplementary Table S9 and Figure S13."
            )
        elif text.startswith("float32 reduces the representation size"):
            set_text(
                paragraph,
                "float32 reduces the representation size of input and many "
                "intermediates, which was essential for the million-row and "
                "extreme-response workflows, but it is not a universal speed "
                "or process-memory guarantee. Classification accuracy was "
                "stable across the supported precision/backend routes, with "
                "a maximum observed difference of 0.22 percentage points. "
                "Regression RMSD was also stable in most comparisons, but "
                "three NMR routes showed relative differences of 2.83-2.90% "
                "against Linux CPU float64; this route-specific variation is "
                "reported rather than treated as exact numerical identity. "
                "Precision should therefore be selected according to "
                "feasibility and verified numerical behaviour."
            )


def update_supplement(document, matrix, figure):
    table = find_table_after_caption(document, "Table S9.")
    replace_table(table, [
        "Dataset", "Family", "A", "Metric", "L CPU64", "L CPU32",
        "CUDA64", "CUDA32", "M CPU64", "M CPU32", "Metal32", "Max change",
    ], table_rows(matrix))
    replace_figure_before_caption(document, "Figure S13.", figure)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Table S9."):
            set_text(
                paragraph,
                "Table S9. Accuracy and RMSD across CPU platforms, "
                "accelerator backends and numerical precision. Linux (L) "
                "CPU float64 is the common reference; Mac is abbreviated M. "
                "Maximum change is the largest absolute accuracy difference "
                "in percentage points or relative RMSD difference across the "
                "available routes."
            )
        elif text.startswith("Conversion to float32 was completed"):
            set_text(
                paragraph,
                "The audit comprised 360 successful family-dataset-route "
                "fits. All 236 nonreference classification comparisons were "
                "within 0.22 percentage points of Linux CPU float64. Of 72 "
                "regression comparisons, 69 were within 1% relative RMSD and "
                "all were within 2.91%. The exceptions were NMR SIMPLS and "
                "linear kernel PLS on CUDA float32 (-2.83%) and NMR PLS-SVD "
                "on the Mac CPU in float32 (+2.90%). For these exceptions, "
                "prediction correlations with the corresponding within-host "
                "CPU float64 result were 0.99987 or greater. Metal float32 "
                "itself was within 0.90% RMSD of the Linux CPU float64 value "
                "for every regression cell. Metal float64 is unavailable and "
                "returns an error rather than falling back to CPU. The same "
                "prepared split, component count and seed were used in every "
                "comparison; classification used argmax and regression used "
                "continuous predictions."
            )
        elif text.startswith("Figure S13."):
            set_text(
                paragraph,
                "Figure S13. Cross-platform precision and backend metric "
                "concordance across 13 datasets and four PLS families. Cell "
                "values are signed changes relative to Linux CPU float64. "
                "Panel A reports classification accuracy in percentage "
                "points; Panel B reports relative RMSD in percent. Gray cells "
                "were not evaluated because Metal has no native float64 route "
                "and the 8-GiB Mac could not hold the ImageNet float64 "
                "reference. Classification used argmax to isolate the PLS "
                "route from LDA fitting."
            )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--manuscript", required=True)
    parser.add_argument("--supplement", required=True)
    parser.add_argument("--matrix", required=True)
    parser.add_argument("--figure", required=True)
    parser.add_argument("--manuscript-output", required=True)
    parser.add_argument("--supplement-output", required=True)
    args = parser.parse_args()

    matrix = rows(args.matrix)
    manuscript = Document(args.manuscript)
    supplement = Document(args.supplement)
    update_main(manuscript)
    update_supplement(supplement, matrix, Path(args.figure))
    manuscript.save(args.manuscript_output)
    supplement.save(args.supplement_output)


if __name__ == "__main__":
    main()
